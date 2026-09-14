"""Recall against the modified EUSES corpus: does the auditor point at an injected fault?

    python benchmarks/corpus/fetch.py euses-modified
    python benchmarks/corpus/recall.py --convert                 # .xls -> .xlsx needs LibreOffice
    python benchmarks/corpus/recall.py --limit 200 --workers 4

Every seeded workbook comes with a .properties file naming the faulty cell(s)
and the fault type. For each fault:

* a **direct hit** is a finding whose location covers the faulty cell in the
  seeded workbook and does not also appear (same rule, same location) in the
  unmodified original, so the mutation is what the auditor reacted to;
* a **line hit** is such a finding on the same row or column of that sheet.

Results go to benchmarks/corpus/results/euses-modified/ as recall_summary.json
and recall_summary.md, broken down by fault type.
"""

from __future__ import annotations

import argparse
import json
import random
import re
import subprocess
import sys
from collections import Counter, defaultdict
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from corpuslib import DATA_DIR, RESULTS_DIR, ROOT, now_iso, run_one, write_json, write_text  # noqa: E402

sys.path.insert(0, str(ROOT))
from spreadsheet_auditor.locations import location_matches  # noqa: E402

FAULT_TYPE_NAMES = {
    "AOR": "arithmetic operator replaced",
    "ROR": "relational operator replaced",
    "RFR": "reference replaced",
    "CRR": "constant replaced by reference",
    "CRP": "constant replaced",
    "CRS": "constant replaced by shift",
    "FRC": "formula replaced by constant",
    "FFR": "function replaced",
}
_VERSION_RE = re.compile(r"_\d+FAULTS_FAULTVERSION\d+", re.IGNORECASE)
_ANCHOR_RE = re.compile(r"^(?P<sheet>.+?)!\$?(?P<col>[A-Za-z]{1,3})\$?(?P<row>\d{1,7})")


def parse_properties(text: str) -> dict:
    """Parse a Graz .properties file into the seeded path and its faults."""
    values: dict[str, str] = {}
    for raw in text.splitlines():
        line = raw.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        values[key.strip()] = value.strip()
    faults = []
    index = 1
    while f"FAULTY_CELLS_{index}" in values:
        cells = values[f"FAULTY_CELLS_{index}"]
        fault_type = values.get(f"FAULT_TYPE_{index}", "?")
        for cell in cells.split(","):
            parts = cell.strip().split("!")
            if len(parts) != 3:
                continue
            faults.append({"sheet_index": int(parts[0]), "col": parts[1].upper(), "row": int(parts[2]), "type": fault_type})
        index += 1
    return {"excel_sheet": values.get("EXCEL_SHEET", "").replace("\\\\", "/").replace("\\", "/"), "faults": faults}


def seeded_path(props_path: Path, props: dict) -> Path:
    return (props_path.parent / props["excel_sheet"]).resolve()


def original_for(seeded: Path) -> Path | None:
    stem = _VERSION_RE.sub("", seeded.name)
    original_dir = seeded.parent.parent / "original"
    for candidate in (original_dir / stem, original_dir / (Path(stem).stem + ".xlsx"), original_dir / (Path(stem).stem + ".xls")):
        if candidate.exists():
            return candidate
    return None


def converted_path(source: Path, data_root: Path, converted_root: Path) -> Path:
    relative = source.relative_to(data_root)
    if source.suffix.lower() == ".xlsx":
        return source
    return (converted_root / relative).with_suffix(".xlsx")


def convert_batch(sources: list[Path], data_root: Path, converted_root: Path, batch_size: int = 40) -> None:
    """Convert .xls files to .xlsx with LibreOffice, a batch per invocation."""
    from spreadsheet_auditor.recalc import seed_profile, soffice_path

    soffice = soffice_path()
    if not soffice:
        raise SystemExit("LibreOffice (soffice) not found; cannot convert .xls files.")
    by_dir: dict[Path, list[Path]] = defaultdict(list)
    for source in sources:
        target = converted_path(source, data_root, converted_root)
        if source.suffix.lower() == ".xlsx" or target.exists():
            continue
        by_dir[target.parent].append(source)
    profile = converted_root / "_profile"
    profile.mkdir(parents=True, exist_ok=True)
    seed_profile(profile)
    for out_dir, files in by_dir.items():
        out_dir.mkdir(parents=True, exist_ok=True)
        for start in range(0, len(files), batch_size):
            batch = files[start : start + batch_size]
            cmd = [
                soffice,
                "--headless",
                "--nologo",
                "--norestore",
                f"-env:UserInstallation=file:///{profile.as_posix()}",
                "--convert-to",
                "xlsx",
                "--outdir",
                str(out_dir),
                *[str(path) for path in batch],
            ]
            subprocess.run(cmd, capture_output=True, text=True, timeout=600 + 30 * len(batch))


def _covers(finding_location: str, target: str) -> bool:
    return any(location_matches(target, piece.strip()) for piece in finding_location.split(",") if piece.strip())


def _same_line(finding_location: str, sheet: str, col: str, row: int) -> bool:
    match = _ANCHOR_RE.match(finding_location.split(",", 1)[0].strip())
    if not match or match.group("sheet").strip("'").casefold() != sheet.casefold():
        return False
    return match.group("col").upper() == col or int(match.group("row")) == row


def evaluate(faults: list[dict], sheetnames: list[str], seeded_findings: list[dict], original_findings: list[dict]) -> list[dict]:
    """Per-fault verdicts: direct/line hits, and whether they are attributable to the mutation."""
    original_keys = {(f["rule_id"], f["location"].replace(" ", "")) for f in original_findings}
    verdicts = []
    for fault in faults:
        if fault["sheet_index"] >= len(sheetnames):
            verdicts.append({**fault, "target": None, "direct": [], "direct_attributable": [], "line": [], "line_attributable": []})
            continue
        sheet = sheetnames[fault["sheet_index"]]
        target = f"{sheet}!{fault['col']}{fault['row']}"
        direct = [f for f in seeded_findings if _covers(f["location"], target)]
        line = [f for f in seeded_findings if _same_line(f["location"], sheet, fault["col"], fault["row"])]

        def attributable(findings: list[dict]) -> list[str]:
            return sorted({f["rule_id"] for f in findings if (f["rule_id"], f["location"].replace(" ", "")) not in original_keys})

        verdicts.append(
            {
                **fault,
                "target": target,
                "direct": sorted({f["rule_id"] for f in direct}),
                "direct_attributable": attributable(direct),
                "line": sorted({f["rule_id"] for f in line}),
                "line_attributable": attributable(line),
            }
        )
    return verdicts


def _sheetnames(workbook: Path) -> list[str]:
    from openpyxl import load_workbook

    wb = load_workbook(workbook, read_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def _findings(payload_path: Path) -> list[dict]:
    if not payload_path.exists():
        return []
    payload = json.loads(payload_path.read_text(encoding="utf-8"))
    return [f for f in payload["findings"] if not f.get("suppressed")]


def summarize(verdicts: list[dict], per_workbook: list[dict]) -> dict:
    by_type: dict[str, Counter] = defaultdict(Counter)
    rules_direct: Counter = Counter()
    for verdict in verdicts:
        bucket = by_type[verdict["type"]]
        bucket["faults"] += 1
        bucket["direct_any"] += bool(verdict["direct"])
        bucket["direct_attributable"] += bool(verdict["direct_attributable"])
        bucket["line_attributable"] += bool(verdict["line_attributable"])
        for rule in verdict["direct_attributable"]:
            rules_direct[rule] += 1
    total = sum(b["faults"] for b in by_type.values())
    attributable = sum(b["direct_attributable"] for b in by_type.values())
    seeded_counts = [w["seeded_findings"] for w in per_workbook if w["status"] == "ok"]
    original_counts = [w["original_findings"] for w in per_workbook if w["status"] == "ok" and w["original_findings"] is not None]
    return {
        "faults": total,
        "direct_attributable": attributable,
        "recall": round(attributable / total, 3) if total else None,
        "line_attributable": sum(b["line_attributable"] for b in by_type.values()),
        "by_type": {t: dict(c) for t, c in sorted(by_type.items())},
        "rules_hitting_faults": dict(rules_direct.most_common()),
        "workbooks": len(per_workbook),
        "ok": sum(1 for w in per_workbook if w["status"] == "ok"),
        "errors": sum(1 for w in per_workbook if w["status"] != "ok"),
        "median_findings_seeded": sorted(seeded_counts)[len(seeded_counts) // 2] if seeded_counts else None,
        "median_findings_original": sorted(original_counts)[len(original_counts) // 2] if original_counts else None,
    }


def render_summary(summary: dict, generated: str) -> str:
    lines = [
        "# Recall on the modified EUSES corpus",
        "",
        f"Generated {generated}. A fault counts as found when a finding covers the faulty cell in the seeded workbook and not in the original.",
        "",
        f"- Faults evaluated: {summary['faults']} in {summary['ok']} workbooks ({summary['errors']} could not be audited)",
        f"- Found at the faulty cell: {summary['direct_attributable']} (**recall {summary['recall']}**); found on the same row or column: {summary['line_attributable']}",
        f"- Median findings per seeded workbook: {summary['median_findings_seeded']}; per original: {summary['median_findings_original']} (noise baseline on real spreadsheets)",
        "",
        "| Fault type | Meaning | Faults | Found at cell | Found on line | Any finding at cell (incl. pre-existing) |",
        "|---|---|---:|---:|---:|---:|",
    ]
    for fault_type, counts in summary["by_type"].items():
        lines.append(
            f"| {fault_type} | {FAULT_TYPE_NAMES.get(fault_type, '?')} | {counts['faults']} | {counts['direct_attributable']} | "
            f"{counts['line_attributable']} | {counts['direct_any']} |"
        )
    lines += ["", "Rules that hit injected faults:", ""]
    for rule, count in summary["rules_hitting_faults"].items():
        lines.append(f"- {rule}: {count}")
    return "\n".join(lines) + "\n"


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--data-dir", default=str(DATA_DIR))
    parser.add_argument("--out-dir", default=str(RESULTS_DIR / "euses-modified"))
    parser.add_argument("--convert", action="store_true", help="Convert missing .xls files with LibreOffice first.")
    parser.add_argument("--limit", type=int, default=None, help="Random sample of this many seeded workbooks (seeded).")
    parser.add_argument("--seed", type=int, default=7)
    parser.add_argument("--workers", type=int, default=4)
    parser.add_argument("--timeout", type=int, default=180)
    args = parser.parse_args(argv)

    data_root = Path(args.data_dir) / "euses-modified" / "files"
    converted_root = Path(args.data_dir) / "euses-modified" / "converted"
    out_dir = Path(args.out_dir)
    props_files = sorted(data_root.rglob("*.properties"))
    if not props_files:
        raise SystemExit(f"no .properties files under {data_root}; run fetch.py euses-modified first.")
    if args.limit is not None and args.limit < len(props_files):
        props_files = sorted(random.Random(args.seed).sample(props_files, args.limit))

    cases = []
    for props_path in props_files:
        props = parse_properties(props_path.read_text(encoding="utf-8", errors="replace"))
        seeded = seeded_path(props_path, props)
        if not seeded.exists() or not props["faults"]:
            continue
        cases.append({"props": props_path, "seeded": seeded, "original": original_for(seeded), "faults": props["faults"]})

    sources = {case["seeded"] for case in cases} | {case["original"] for case in cases if case["original"]}
    if args.convert:
        convert_batch(sorted(sources), data_root, converted_root)

    findings_dir = out_dir / "findings"
    findings_dir.mkdir(parents=True, exist_ok=True)
    # Converted workbooks already carry values LibreOffice computed with the
    # auditor's own forced-recalculation profile; auditing them again through
    # LibreOffice would only cost time.
    config = out_dir / "config.json"
    config.write_text(json.dumps({"recalc": {"enabled": False}}), encoding="utf-8")
    extra_args = ["--config", str(config)]
    audited: dict[Path, dict] = {}

    def audit(source: Path) -> tuple[Path, dict]:
        target = converted_path(source, data_root, converted_root)
        out_json = findings_dir / (re.sub(r"[^A-Za-z0-9_.-]+", "_", source.relative_to(data_root).as_posix()) + ".json")
        if not target.exists():
            return source, {"status": "missing", "elapsed": 0.0, "out": out_json}
        record = run_one(target, out_json, args.timeout, extra_args)
        record["out"] = out_json
        record["converted"] = target
        return source, record

    with ThreadPoolExecutor(max_workers=max(1, args.workers)) as pool:
        for source, record in pool.map(audit, sorted(sources)):
            audited[source] = record

    verdicts: list[dict] = []
    per_workbook: list[dict] = []
    for case in cases:
        seeded_record = audited[case["seeded"]]
        original_record = audited.get(case["original"]) if case["original"] else None
        if seeded_record["status"] != "ok":
            per_workbook.append({"seeded": case["seeded"].name, "status": seeded_record["status"], "seeded_findings": None, "original_findings": None})
            continue
        seeded_findings = _findings(seeded_record["out"])
        original_findings = _findings(original_record["out"]) if original_record and original_record["status"] == "ok" else []
        sheetnames = _sheetnames(seeded_record["converted"])
        case_verdicts = evaluate(case["faults"], sheetnames, seeded_findings, original_findings)
        for verdict in case_verdicts:
            verdict["seeded"] = case["seeded"].name
        verdicts.extend(case_verdicts)
        per_workbook.append(
            {
                "seeded": case["seeded"].name,
                "status": "ok",
                "seeded_findings": len(seeded_findings),
                "original_findings": len(original_findings) if original_record and original_record["status"] == "ok" else None,
            }
        )

    summary = summarize(verdicts, per_workbook)
    generated = now_iso()
    write_json(out_dir / "recall_summary.json", {"generated": generated, "summary": summary, "verdicts": verdicts, "workbooks": per_workbook})
    write_text(out_dir / "recall_summary.md", render_summary(summary, generated))
    print(render_summary(summary, generated))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
