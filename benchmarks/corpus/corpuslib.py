"""Shared code for the real-world corpus precision harness.

The pipeline is four scripts that share this module:

    fetch.py       download and extract a corpus into benchmarks/corpus/data/<source>/
    run_corpus.py  audit every workbook, recording findings, crashes, and timeouts
    sample.py      draw a stratified sample of findings and build context cards
    report.py      join sample + hand labels into benchmarks/real_world_precision.md

Workbooks and raw results stay out of git; sample cards, labels, and the
report are committed.
"""

from __future__ import annotations

import json
import math
import random
import re
import subprocess
import sys
import time
from collections import Counter, defaultdict
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime, timezone
from fnmatch import fnmatch
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
CORPUS_DIR = Path(__file__).resolve().parent
DATA_DIR = CORPUS_DIR / "data"
RESULTS_DIR = CORPUS_DIR / "results"
SAMPLES_DIR = CORPUS_DIR / "samples"
LABELS_DIR = CORPUS_DIR / "labels"
WORKBOOK_SUFFIXES = {".xlsx", ".xlsm"}
VALUE_DEPENDENT_RULES = {"TOTAL_MISMATCH", "CROSS_FOOT_FAILURE", "LIVE_ERROR"}


def load_sources() -> dict:
    return json.loads((CORPUS_DIR / "sources.json").read_text(encoding="utf-8"))["sources"]


def now_iso() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def write_json(path: Path, payload) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, indent=2, ensure_ascii=False) + "\n", encoding="utf-8", newline="\n")


def write_text(path: Path, text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8", newline="\n")


# --- statistics --------------------------------------------------------------


def wilson_interval(successes: int, n: int, z: float = 1.96) -> tuple[float, float]:
    """Wilson score interval for a binomial proportion; (0, 1) when n is 0."""
    if n <= 0:
        return (0.0, 1.0)
    p = successes / n
    denominator = 1 + z * z / n
    centre = (p + z * z / (2 * n)) / denominator
    margin = z * math.sqrt(p * (1 - p) / n + z * z / (4 * n * n)) / denominator
    return (max(0.0, centre - margin), min(1.0, centre + margin))


def percentile(values: list[float], fraction: float) -> float:
    if not values:
        return 0.0
    ordered = sorted(values)
    index = min(len(ordered) - 1, max(0, int(round(fraction * (len(ordered) - 1)))))
    return ordered[index]


# --- running the auditor -----------------------------------------------------


def list_workbooks(data_dir: Path, include_glob: str | None = None) -> list[Path]:
    workbooks = []
    for path in sorted(data_dir.rglob("*")):
        if path.suffix.lower() not in WORKBOOK_SUFFIXES or path.name.startswith("~$"):
            continue
        if include_glob and not fnmatch(path.name, include_glob):
            continue
        workbooks.append(path)
    return workbooks


def run_one(workbook: Path, out_json: Path, timeout: int, extra_args: list[str] | None = None) -> dict:
    """Audit one workbook in a subprocess and summarize the outcome."""
    out_json.parent.mkdir(parents=True, exist_ok=True)
    cmd = [
        sys.executable,
        "-m",
        "spreadsheet_auditor",
        str(workbook),
        "--json",
        str(out_json),
        "--fail-on",
        "None",
        "--quiet",
        *(extra_args or []),
    ]
    started = time.monotonic()
    try:
        proc = subprocess.run(cmd, capture_output=True, text=True, timeout=timeout, cwd=ROOT)
    except subprocess.TimeoutExpired:
        return {"status": "timeout", "exit_code": None, "elapsed": round(time.monotonic() - started, 3)}
    elapsed = round(time.monotonic() - started, 3)
    record: dict = {"status": "ok" if proc.returncode in (0, 1, 2) else "error", "exit_code": proc.returncode, "elapsed": elapsed}
    if record["status"] != "ok" or not out_json.exists():
        record["status"] = "error"
        record["stderr"] = (proc.stderr or "").strip()[-600:]
        return record
    payload = json.loads(out_json.read_text(encoding="utf-8"))
    findings = [f for f in payload["findings"] if not f.get("suppressed")]
    record.update(
        {
            "sha256": payload["workbook"]["sha256"],
            "formulas": payload["workbook"]["formulas_scanned"],
            "sheets": payload["workbook"]["sheets_analyzed"],
            "recalc": payload["workbook"]["recalc_status"],
            "rules": dict(Counter(f["rule_id"] for f in findings)),
            "severities": dict(Counter(f["severity"] for f in findings)),
            "limitations": len(payload["coverage"].get("limitations", [])),
            "unsupported": payload["coverage"].get("unsupported_features", []),
            "truncated": [k for k, v in (payload["coverage"].get("truncated") or {}).items() if v],
        }
    )
    return record


def run_corpus(
    source: str,
    data_dir: Path,
    out_dir: Path,
    limit: int | None = None,
    seed: int = 7,
    workers: int = 4,
    timeout: int = 180,
    include_glob: str | None = None,
    extra_args: list[str] | None = None,
    progress=None,
) -> dict:
    workbooks = list_workbooks(data_dir, include_glob)
    if limit is not None and limit < len(workbooks):
        rng = random.Random(seed)
        workbooks = sorted(rng.sample(workbooks, limit))
    findings_dir = out_dir / "findings"
    findings_dir.mkdir(parents=True, exist_ok=True)
    records: dict[str, dict] = {}
    started = time.monotonic()

    def task(index: int, workbook: Path) -> tuple[str, dict]:
        rel = workbook.relative_to(data_dir).as_posix()
        record = run_one(workbook, findings_dir / f"{index:05d}.json", timeout, extra_args)
        record["index"] = index
        record["workbook"] = rel
        return rel, record

    with ThreadPoolExecutor(max_workers=max(1, workers)) as pool:
        futures = [pool.submit(task, index, workbook) for index, workbook in enumerate(workbooks)]
        for done, future in enumerate(as_completed(futures), start=1):
            rel, record = future.result()
            records[rel] = record
            if progress and (done % 100 == 0 or done == len(futures)):
                progress(done, len(futures))

    summary = summarize(source, data_dir, include_glob, seed, limit, records)
    summary["wall_clock_seconds"] = round(time.monotonic() - started, 1)
    write_json(out_dir / "run_summary.json", summary)
    return summary


def summarize(source: str, data_dir: Path, include_glob: str | None, seed: int, limit: int | None, records: dict[str, dict]) -> dict:
    ok = [r for r in records.values() if r["status"] == "ok"]
    rule_findings: Counter = Counter()
    rule_workbooks: Counter = Counter()
    severities: Counter = Counter()
    for record in ok:
        for rule, count in record.get("rules", {}).items():
            rule_findings[rule] += count
            rule_workbooks[rule] += 1
        severities.update(record.get("severities", {}))
    formulas = sum(r.get("formulas", 0) for r in ok)
    elapsed = [r["elapsed"] for r in records.values()]
    recalc = Counter(r.get("recalc", "n/a") for r in ok)
    unsupported: Counter = Counter()
    for record in ok:
        unsupported.update(record.get("unsupported", []))
    return {
        "source": source,
        "data_dir": str(data_dir),
        "include_glob": include_glob,
        "seed": seed,
        "limit": limit,
        "generated": now_iso(),
        "auditor_version": _auditor_version(),
        "workbooks": len(records),
        "ok": len(ok),
        "errors": sum(1 for r in records.values() if r["status"] == "error"),
        "timeouts": sum(1 for r in records.values() if r["status"] == "timeout"),
        "formulas": formulas,
        "recalc": dict(recalc),
        "findings_total": sum(rule_findings.values()),
        "workbooks_with_findings": sum(1 for r in ok if r.get("rules")),
        "severities": dict(severities),
        "rules": {
            rule: {
                "findings": rule_findings[rule],
                "workbooks": rule_workbooks[rule],
                "per_1k_formulas": round(1000 * rule_findings[rule] / formulas, 2) if formulas else None,
            }
            for rule in sorted(rule_findings)
        },
        "unsupported_features": dict(unsupported),
        "elapsed": {
            "median": round(percentile(elapsed, 0.5), 3),
            "p95": round(percentile(elapsed, 0.95), 3),
            "max": round(max(elapsed), 3) if elapsed else 0.0,
        },
        "records": {rel: records[rel] for rel in sorted(records)},
    }


def _auditor_version() -> str:
    try:
        sys.path.insert(0, str(ROOT))
        import spreadsheet_auditor  # noqa: WPS433

        return spreadsheet_auditor.__version__
    except Exception:  # pragma: no cover - defensive
        return "unknown"


def render_run_table(summary: dict) -> str:
    lines = [
        f"source: {summary['source']}   workbooks: {summary['workbooks']}   ok: {summary['ok']}   "
        f"errors: {summary['errors']}   timeouts: {summary['timeouts']}   formulas: {summary['formulas']}",
        f"findings: {summary['findings_total']} in {summary['workbooks_with_findings']} workbooks   "
        f"elapsed median/p95/max: {summary['elapsed']['median']}/{summary['elapsed']['p95']}/{summary['elapsed']['max']} s",
        "",
        f"{'rule':28} {'findings':>9} {'workbooks':>10} {'per 1k formulas':>16}",
    ]
    for rule, stats in sorted(summary["rules"].items(), key=lambda item: -item[1]["findings"]):
        lines.append(f"{rule:28} {stats['findings']:>9} {stats['workbooks']:>10} {stats['per_1k_formulas'] if stats['per_1k_formulas'] is not None else '-':>16}")
    return "\n".join(lines)


# --- sampling and context cards ---------------------------------------------


def load_findings(results_dir: Path) -> list[dict]:
    """Every unsuppressed finding from a run, tagged with its workbook and index."""
    summary = json.loads((results_dir / "run_summary.json").read_text(encoding="utf-8"))
    items: list[dict] = []
    for rel, record in summary["records"].items():
        if record["status"] != "ok":
            continue
        payload_path = results_dir / "findings" / f"{record['index']:05d}.json"
        if not payload_path.exists():
            continue
        payload = json.loads(payload_path.read_text(encoding="utf-8"))
        for finding in payload["findings"]:
            if finding.get("suppressed"):
                continue
            items.append({"workbook": rel, "sha256": record["sha256"], "finding": finding})
    return items


def finding_key(sha256: str, rule_id: str, location: str) -> str:
    return f"{sha256[:12]}|{rule_id}|{location.replace(' ', '')}"


def stratified_sample(items: list[dict], per_rule: int, max_per_workbook: int, seed: int) -> list[dict]:
    """Up to ``per_rule`` findings per rule, at most ``max_per_workbook`` from one workbook."""
    by_rule: dict[str, list[dict]] = defaultdict(list)
    for item in items:
        by_rule[item["finding"]["rule_id"]].append(item)
    rng = random.Random(seed)
    chosen: list[dict] = []
    for rule in sorted(by_rule):
        pool = sorted(by_rule[rule], key=lambda it: (it["workbook"], it["finding"]["location"]))
        rng.shuffle(pool)
        taken_per_workbook: Counter = Counter()
        for item in pool:
            if len([c for c in chosen if c["finding"]["rule_id"] == rule]) >= per_rule:
                break
            if taken_per_workbook[item["workbook"]] >= max_per_workbook:
                continue
            taken_per_workbook[item["workbook"]] += 1
            chosen.append(item)
    return chosen


_LOCATION_RE = re.compile(r"^(?P<sheet>.+?)!(?P<ref>\$?[A-Za-z]{1,3}\$?\d{1,7})(?::\$?[A-Za-z]{1,3}\$?\d{1,7})?$")


def _anchor(location: str) -> tuple[str, str] | None:
    piece = location.split(",", 1)[0].strip()
    match = _LOCATION_RE.match(piece)
    if not match:
        return None
    return match.group("sheet"), match.group("ref").replace("$", "").upper()


class WorkbookCache:
    """Open each sampled workbook once, for both formulas and cached values."""

    def __init__(self) -> None:
        self._cache: dict[Path, tuple] = {}

    def get(self, path: Path):
        if path not in self._cache:
            from openpyxl import load_workbook

            keep_vba = path.suffix.lower() == ".xlsm"
            self._cache[path] = (
                load_workbook(path, data_only=False, keep_vba=keep_vba),
                load_workbook(path, data_only=True, keep_vba=keep_vba),
            )
        return self._cache[path]


def _short(value, limit: int = 60) -> str:
    text = repr(value) if isinstance(value, str) else str(value)
    return text if len(text) <= limit else text[: limit - 3] + "..."


def build_context(workbook_path: Path, finding: dict, cache: WorkbookCache, window: int = 2) -> dict:
    """A compact neighbourhood dump a labeler can judge the finding from."""
    from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter

    anchor = _anchor(finding.get("location", ""))
    if anchor is None:
        return {"note": "location is not a cell address"}
    sheet, coord = anchor
    try:
        formula_wb, value_wb = cache.get(workbook_path)
    except Exception as exc:  # pragma: no cover - corrupt workbook
        return {"note": f"could not open workbook: {exc!r}"}
    if sheet not in formula_wb.sheetnames:
        return {"note": f"sheet {sheet!r} not found"}
    fws, vws = formula_wb[sheet], value_wb[sheet]
    row, col = coordinate_to_tuple(coord)
    fcells = getattr(fws, "_cells", {})
    vcells = getattr(vws, "_cells", {})

    def fval(r, c):
        cell = fcells.get((r, c))
        return None if cell is None else cell.value

    def vval(r, c):
        cell = vcells.get((r, c))
        return None if cell is None else cell.value

    labels = [
        _short(fval(row, c)) for c in range(1, col) if isinstance(fval(row, c), str) and not str(fval(row, c)).startswith("=")
    ][-2:]
    header = ""
    for r in range(row - 1, max(0, row - 12), -1):
        value = fval(r, col)
        if value is None:
            continue
        if isinstance(value, str) and not value.startswith("="):
            header = _short(value)
        break
    neighbourhood = []
    for r in range(max(1, row - window), row + window + 1):
        for c in range(max(1, col - window), col + window + 1):
            f, v = fval(r, c), vval(r, c)
            if f is None and v is None:
                continue
            addr = f"{get_column_letter(c)}{r}"
            if isinstance(f, str) and f.startswith("="):
                neighbourhood.append(f"{addr}: {_short(f, 70)} -> {_short(v, 30)}")
            else:
                neighbourhood.append(f"{addr}: {_short(f if f is not None else v, 40)}")
    context = {
        "sheet": sheet,
        "cell": coord,
        "cached_value": _short(vval(row, col)),
        "row_labels": labels,
        "column_header": header,
        "neighbourhood": neighbourhood[:30],
        "sheet_used_range": f"A1:{get_column_letter(fws.max_column or 1)}{fws.max_row or 1}",
    }
    formula = finding.get("formula") or ""
    ranges = re.findall(r"(?:'[^']+'!|[A-Za-z0-9_]+!)?\$?[A-Z]{1,3}\$?\d+:\$?[A-Z]{1,3}\$?\d+", formula)
    if ranges:
        from openpyxl.utils.cell import range_boundaries

        text = ranges[0]
        rsheet = sheet
        if "!" in text:
            rsheet, text = text.rsplit("!", 1)
            rsheet = rsheet.strip("'")
        try:
            min_col, min_row, max_col, max_row = range_boundaries(text.replace("$", ""))
        except ValueError:
            min_col = None
        if min_col and rsheet in value_wb.sheetnames:
            rv = getattr(value_wb[rsheet], "_cells", {})
            rf = getattr(formula_wb[rsheet], "_cells", {})
            values = []
            for r in range(min_row, max_row + 1):
                for c in range(min_col, max_col + 1):
                    cell = rv.get((r, c))
                    values.append(_short(None if cell is None else cell.value, 20))
                    if len(values) >= 12:
                        break
                if len(values) >= 12:
                    break
            beyond = {}
            if min_col == max_col:
                for label, r in (("above", min_row - 1), ("below", max_row + 1)):
                    if r >= 1:
                        fcell, vcell = rf.get((r, min_col)), rv.get((r, min_col))
                        beyond[label] = f"{get_column_letter(min_col)}{r}: " + _short(
                            (fcell.value if fcell is not None and isinstance(fcell.value, str) and fcell.value.startswith('=') else (vcell.value if vcell is not None else None)), 40
                        )
            if min_row == max_row:
                for label, c in (("left", min_col - 1), ("right", max_col + 1)):
                    if c >= 1:
                        fcell, vcell = rf.get((min_row, c)), rv.get((min_row, c))
                        beyond[label] = f"{get_column_letter(c)}{min_row}: " + _short(
                            (fcell.value if fcell is not None and isinstance(fcell.value, str) and fcell.value.startswith('=') else (vcell.value if vcell is not None else None)), 40
                        )
            context["first_range"] = {"ref": ranges[0], "values": values, "beyond": beyond}
    return context


def render_cards(sample: dict) -> str:
    """Markdown cards for hand labeling, one section per sampled finding."""
    lines = [
        f"# Finding cards: {sample['source']}",
        "",
        f"Generated {sample['generated']} from {sample['results_dir']} with seed {sample['seed']}, "
        f"up to {sample['per_rule']} per rule and {sample['max_per_workbook']} per workbook.",
        "",
        "Label each card in `labels/<source>.json` as `TP` (the auditor is right), `FP` (it is wrong), or `unsure`, with a one-line reason.",
        "",
    ]
    current_rule = None
    for item in sample["items"]:
        finding = item["finding"]
        if finding["rule_id"] != current_rule:
            current_rule = finding["rule_id"]
            lines += [f"## {current_rule}", ""]
        lines.append(f"### `{item['key']}`")
        lines.append("")
        lines.append(f"- workbook: `{item['workbook']}`")
        lines.append(f"- location: `{finding['location']}`  severity: {finding['severity']}  confidence: {finding['error_confidence']}")
        if finding.get("formula"):
            lines.append(f"- formula: `{finding['formula']}`")
        for evidence in finding.get("evidence", [])[:3]:
            lines.append(f"- evidence: {evidence}")
        context = item.get("context", {})
        if context.get("note"):
            lines.append(f"- context: {context['note']}")
        else:
            lines.append(f"- cached value: {context.get('cached_value')}; row labels: {context.get('row_labels')}; column header: {context.get('column_header')!r}; used range: {context.get('sheet_used_range')}")
            if context.get("first_range"):
                fr = context["first_range"]
                lines.append(f"- range {fr['ref']}: values {fr['values']}; beyond: {fr['beyond']}")
            if context.get("neighbourhood"):
                lines.append("- neighbourhood: " + " | ".join(context["neighbourhood"]))
        lines.append("")
    return "\n".join(lines)


# --- report ------------------------------------------------------------------


def load_labels(path: Path) -> dict:
    if not path.exists():
        return {"labeler": "", "protocol": "", "rule_notes": {}, "labels": {}}
    return json.loads(path.read_text(encoding="utf-8"))


def precision_rows(sample: dict, labels: dict, run_summary: dict) -> list[dict]:
    per_rule: dict[str, Counter] = defaultdict(Counter)
    for item in sample["items"]:
        rule = item["finding"]["rule_id"]
        entry = labels.get("labels", {}).get(item["key"])
        per_rule[rule][(entry or {}).get("label", "unlabeled")] += 1
    rows = []
    for rule in sorted(set(per_rule) | set(run_summary["rules"])):
        counts = per_rule.get(rule, Counter())
        tp, fp, unsure, unlabeled = counts["TP"], counts["FP"], counts["unsure"], counts["unlabeled"]
        judged = tp + fp
        precision = tp / judged if judged else None
        low, high = wilson_interval(tp, judged) if judged else (None, None)
        stats = run_summary["rules"].get(rule, {"findings": 0, "workbooks": 0, "per_1k_formulas": None})
        rows.append(
            {
                "rule": rule,
                "findings": stats["findings"],
                "workbooks": stats["workbooks"],
                "per_1k_formulas": stats["per_1k_formulas"],
                "sampled": sum(counts.values()),
                "tp": tp,
                "fp": fp,
                "unsure": unsure,
                "unlabeled": unlabeled,
                "precision": precision,
                "ci_low": low,
                "ci_high": high,
            }
        )
    return rows


def _pct(value: float | None) -> str:
    return "-" if value is None else f"{100 * value:.0f}%"


def render_report(source_meta: dict, run_summary: dict, sample: dict, labels: dict, rows: list[dict]) -> str:
    judged = sum(r["tp"] + r["fp"] for r in rows)
    tp_total = sum(r["tp"] for r in rows)
    overall = tp_total / judged if judged else None
    low, high = wilson_interval(tp_total, judged) if judged else (None, None)
    lines = [
        "# Real-world precision",
        "",
        "Automatically generated by `python benchmarks/corpus/report.py`. Do not edit by hand; edit the labels file and regenerate.",
        "",
        f"**Corpus:** {source_meta['name']} ({source_meta['license']}). {source_meta['notes']}",
        "",
        f"**Run:** auditor {run_summary['auditor_version']}, {run_summary['generated']}, "
        f"{run_summary['ok']} of {run_summary['workbooks']} workbooks audited "
        f"({run_summary['errors']} errors, {run_summary['timeouts']} timeouts), "
        f"{run_summary['formulas']:,} formulas, recalculation {run_summary['recalc']}. "
        f"Median audit {run_summary['elapsed']['median']} s, p95 {run_summary['elapsed']['p95']} s, max {run_summary['elapsed']['max']} s.",
        "",
        f"**Labels:** {labels.get('labeler') or 'none yet'}. {labels.get('protocol', '')}",
        "",
        "## Headline",
        "",
        f"- Findings: {run_summary['findings_total']:,} across {run_summary['workbooks_with_findings']:,} workbooks "
        f"({100 * run_summary['workbooks_with_findings'] / max(run_summary['ok'], 1):.0f}% of audited workbooks have at least one finding).",
        f"- Sampled and judged: {judged} findings; overall precision **{_pct(overall)}** "
        f"(95% Wilson interval {_pct(low)} to {_pct(high)}). Unsure labels are excluded from precision and shown separately.",
        "",
        "## Precision by rule",
        "",
        "| Rule | Findings | Workbooks | Per 1k formulas | Sampled | TP | FP | Unsure | Precision | 95% CI |",
        "|---|---:|---:|---:|---:|---:|---:|---:|---:|---:|",
    ]
    for row in sorted(rows, key=lambda r: -r["findings"]):
        ci = "-" if row["ci_low"] is None else f"{_pct(row['ci_low'])} to {_pct(row['ci_high'])}"
        per_1k = "-" if row["per_1k_formulas"] is None else f"{row['per_1k_formulas']:.2f}"
        lines.append(
            f"| {row['rule']} | {row['findings']:,} | {row['workbooks']:,} | {per_1k} | {row['sampled']} | "
            f"{row['tp']} | {row['fp']} | {row['unsure']} | {_pct(row['precision'])} | {ci} |"
        )
    lines += ["", "## Notes by rule", ""]
    notes = labels.get("rule_notes", {})
    if notes:
        for rule in sorted(notes):
            lines.append(f"- **{rule}**: {notes[rule]}")
    else:
        lines.append("- No rule notes yet.")
    lines += [
        "",
        "## Coverage flags seen",
        "",
    ]
    if run_summary.get("unsupported_features"):
        for flag, count in sorted(run_summary["unsupported_features"].items(), key=lambda item: -item[1]):
            lines.append(f"- `{flag}`: {count} workbooks")
    else:
        lines.append("- none")
    lines += [
        "",
        "## Method and caveats",
        "",
        "- Every workbook is audited with the CLI in a subprocess with a per-file timeout; crashes and timeouts are counted, not hidden.",
        f"- The sample takes up to {sample['per_rule']} findings per rule, at most {sample['max_per_workbook']} from any one workbook, seeded, so a rule that fires thousands of times is judged on the same footing as one that fires ten times.",
        "- Each sampled finding is judged from a context card (formula, cached value, row labels, column header, a 5x5 neighbourhood, and the values inside and just outside the first referenced range). TP means the finding points at something a careful reviewer would want to look at for the reason the rule gives; FP means the rule misread the workbook; unsure means the card was not enough to tell.",
        "- Precision is TP / (TP + FP) with a Wilson 95% interval. Small samples have wide intervals; read the interval, not the point estimate.",
        "- This measures precision only. Recall needs ground truth; see the `euses-modified` source in `benchmarks/corpus/sources.json`.",
        "- Forum workbooks are small and question-shaped, so volumes per 1k formulas will not transfer to large financial models; the per-rule precision is the more portable number.",
        "",
    ]
    return "\n".join(lines)
