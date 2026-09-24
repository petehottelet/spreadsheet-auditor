"""Print the cells around audit findings so each one can be judged in context.

    python scripts/context.py WORKBOOK FINDINGS_JSON [ID_OR_CELL ...]

With no IDs, prints a card for every Critical or High finding that is not
already a certain defect (confidence other than `Defect`), up to 25. Pass
finding IDs (`FORMULA_DRIFT-002`) or cells (`P&L!F18`) to pick others.

Each card shows the row label and column header of the flagged cell, its
formula and cached value, what a one-cell link points at, and the formulas
and values two cells around it, which is what decides whether a flagged cell
is a mistake or a line that differs by design (a total, a net line, a summary
row, an input column).
"""

from __future__ import annotations

import json
import re
import sys
import warnings
from pathlib import Path

MAX_CARDS = 25  # about 100 lines of output; more is better read by rule from --summary
WINDOW = 2  # cells each way: the 5x5 neighbourhood the corpus labels were judged from


def _short(value, limit: int = 60) -> str:
    text = getattr(value, "text", value)  # array formulas carry their text on .text
    text = repr(text) if isinstance(text, str) else str(text)
    return text if len(text) <= limit else text[: limit - 3] + "..."


def _anchor(location: str) -> tuple[str, str] | None:
    first = location.split(",", 1)[0].strip()
    if "!" not in first:
        return None
    sheet, coord = first.rsplit("!", 1)
    return sheet.strip("'"), coord.split(":", 1)[0].replace("$", "").upper()


def _is_formula(value) -> bool:
    return (isinstance(value, str) and value.startswith("=")) or hasattr(value, "text")


_LINK_RE = re.compile(
    r"^=\+?\s*(?:(?:'(?P<quoted>(?:[^']|'')+)'|(?P<bare>[A-Za-z0-9_.]+))!)?"
    r"\$?(?P<col>[A-Za-z]{1,3})\$?(?P<row>\d+)\s*$"
)


def _link_target(formula, sheet: str) -> tuple[str, str] | None:
    """``(sheet, cell)`` when the formula is a single-cell link such as ``=C61`` or ``='Other'!M14``."""
    match = _LINK_RE.match(formula) if isinstance(formula, str) else None
    if match is None:
        return None
    target_sheet = match["quoted"].replace("''", "'") if match["quoted"] else match["bare"] or sheet
    return target_sheet, f"{match['col'].upper()}{match['row']}"


def card(formula_wb, value_wb, location: str) -> list[str]:
    from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter

    anchor = _anchor(location)
    if anchor is None or anchor[0] not in formula_wb.sheetnames:
        return [f"  (no cell context for {location!r})"]
    sheet, coord = anchor
    fcells = formula_wb[sheet]._cells  # existing cells only; never materializes new ones
    vcells = value_wb[sheet]._cells
    row, col = coordinate_to_tuple(coord)

    def f(r, c):
        cell = fcells.get((r, c))
        return None if cell is None else cell.value

    def v(r, c):
        cell = vcells.get((r, c))
        return None if cell is None else cell.value

    labels = [_short(f(row, c), 40) for c in range(1, col) if isinstance(f(row, c), str) and not _is_formula(f(row, c))][-2:]
    header = ""
    for r in range(row - 1, max(0, row - 200), -1):  # the nearest text above, past the numbers and formulas of the block
        above = f(r, col)
        if isinstance(above, str) and not _is_formula(above):
            header = _short(above, 40)
            break
    lines = [
        f"  row label: {', '.join(labels) or '-'} | column header: {header or '-'}",
        f"  cell: {coord} {_short(f(row, col), 90)} -> {_short(v(row, col), 30)}",
    ]
    link = _link_target(f(row, col), sheet)
    if link is not None and link[0] in formula_wb.sheetnames:
        # A one-cell link (=C61, ='Other sheet'!M14) is judged by what it points at.
        target_sheet, target = link
        trow, tcol = coordinate_to_tuple(target)
        tcells, tvalues = formula_wb[target_sheet]._cells, value_wb[target_sheet]._cells
        texts = [tcells[(trow, c)].value for c in range(1, tcol) if (trow, c) in tcells]
        label = next((t for t in reversed(texts) if isinstance(t, str) and not _is_formula(t)), None)
        value = tvalues[(trow, tcol)].value if (trow, tcol) in tvalues else None
        lines.append(f"  links to: {target_sheet}!{target} (row label: {_short(label, 40) if label else '-'}) -> {_short(value, 30)}")
    around = []
    for r in range(max(1, row - WINDOW), row + WINDOW + 1):
        for c in range(max(1, col - WINDOW), col + WINDOW + 1):
            if (r, c) == (row, col) or (f(r, c) is None and v(r, c) is None):
                continue
            addr = f"{get_column_letter(c)}{r}"
            if _is_formula(f(r, c)):
                around.append(f"{addr} {_short(f(r, c), 50)} -> {_short(v(r, c), 20)}")
            else:
                around.append(f"{addr} {_short(f(r, c), 30)}")
    lines.append("  around: " + (" | ".join(around[:16]) or "-"))
    return lines


def main(argv: list[str]) -> int:
    if len(argv) < 2:
        print(__doc__.strip(), file=sys.stderr)
        return 2
    workbook, findings_path, picks = Path(argv[0]), Path(argv[1]), argv[2:]
    for path in (workbook, findings_path):
        if not path.exists():
            print(f"Not found: {path}", file=sys.stderr)
            return 2
    findings = [f for f in json.loads(findings_path.read_text(encoding="utf-8"))["findings"] if not f.get("suppressed")]
    if picks:
        by_id = {f["id"]: f for f in findings}
        chosen = [by_id.get(p, {"id": None, "location": p}) for p in picks]
    else:
        chosen = [f for f in findings if f["severity"] in ("Critical", "High") and f["error_confidence"] != "Defect"]
    # One card per cell: a cell flagged by two rules is one thing to judge.
    by_cell: dict[str, list[dict]] = {}
    for finding in chosen:
        by_cell.setdefault(finding["location"], []).append(finding)
    from openpyxl import load_workbook

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")  # openpyxl warns about parts it drops; nothing is written here
        keep_vba = workbook.suffix.lower() == ".xlsm"
        formula_wb = load_workbook(workbook, data_only=False, keep_vba=keep_vba)
        value_wb = load_workbook(workbook, data_only=True, keep_vba=keep_vba)
    for location, group in list(by_cell.items())[:MAX_CARDS]:
        heads = [f"{f['id']} {f['rule_id']} {f['severity']} {f['error_confidence']}" for f in group if f.get("id")]
        print(f"{', '.join(heads)} at {location}" if heads else location)
        print("\n".join(card(formula_wb, value_wb, location)))
    if len(by_cell) > MAX_CARDS:
        print(f"... {len(by_cell) - MAX_CARDS} more cells; pass their IDs to see them.")
    if not chosen:
        print("No Critical or High finding needs a context check; pass IDs or cells to inspect others.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
