"""What an agent running the skill relies on: the source is never overwritten,
failures say what to do next, --summary names the rules, and context cards
show what a flagged cell sits among."""

from __future__ import annotations

import shutil
import subprocess
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
DEMO = ROOT / "examples" / "demo_bad_budget.xlsx"
AUDIT = ROOT / "scripts" / "audit.py"
CONTEXT = ROOT / "scripts" / "context.py"


def _run(script: Path, *args: str, cwd: Path | None = None) -> subprocess.CompletedProcess:
    return subprocess.run([sys.executable, str(script), *args], capture_output=True, text=True, cwd=cwd)


@pytest.mark.parametrize("flag", ["--out", "--json", "--annotated"])
def test_output_path_equal_to_source_is_refused(flag, tmp_path):
    workbook = tmp_path / "model.xlsx"
    shutil.copy2(DEMO, workbook)
    before = workbook.read_bytes()
    result = _run(AUDIT, str(workbook), flag, str(workbook))
    assert result.returncode == 4, result.stderr
    assert "workbook being audited" in result.stderr
    assert workbook.read_bytes() == before


def test_encrypted_workbook_gets_an_actionable_message(tmp_path):
    workbook = tmp_path / "locked.xlsx"
    workbook.write_bytes(bytes.fromhex("D0CF11E0A1B11AE1") + b"\x00" * 504)
    result = _run(AUDIT, str(workbook))
    assert result.returncode == 4
    assert "password-protected" in result.stderr


def test_legacy_extension_says_how_to_convert(tmp_path):
    workbook = tmp_path / "old.xls"
    workbook.write_bytes(bytes.fromhex("D0CF11E0A1B11AE1") + b"\x00" * 504)
    result = _run(AUDIT, str(workbook))
    assert result.returncode == 4
    assert "--convert-to xlsx" in result.stderr


def test_summary_counts_findings_by_rule():
    result = _run(AUDIT, str(DEMO), "--summary")
    assert result.returncode == 1, result.stderr
    assert "by rule" in result.stdout
    assert "FORMULA_DRIFT (High, Likely defect)" in result.stdout


def test_context_cards_show_labels_and_neighbours(tmp_path):
    findings = tmp_path / "findings.json"
    audit = _run(AUDIT, str(DEMO), "--json", str(findings), "--quiet")
    assert audit.returncode == 1, audit.stderr
    result = _run(CONTEXT, str(DEMO), str(findings))
    assert result.returncode == 0, result.stderr
    # Budget!B10 stops its SUM a row short: the card shows the row label and the Marketing row it skips.
    card = result.stdout.split("RANGE_EXCLUSION-001", 1)[1]
    assert "row label: 'COGS total'" in card
    assert "A9 'Marketing'" in card
    # Certain defects (confidence Defect) are not re-checked by default.
    assert "LIVE_ERROR-001" not in result.stdout


def test_annotated_copy_keeps_every_finding_on_a_cell(tmp_path):
    from openpyxl import load_workbook

    annotated = tmp_path / "annotated.xlsx"
    result = _run(AUDIT, str(DEMO), "--annotated", str(annotated), "--quiet")
    assert result.returncode == 1, result.stderr
    note = load_workbook(annotated)["Budget"]["B10"].comment.text
    # B10 carries a Critical RANGE_EXCLUSION and a High FORMULA_DRIFT; the Critical one leads.
    assert note.index("RANGE_EXCLUSION") < note.index("FORMULA_DRIFT")


def test_stray_far_cell_does_not_switch_off_grid_checks(tmp_path):
    import json

    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["Line", "Jan", "Feb", "Mar", "Total"])
    ws.append(["A", 1, 2, 3, "=SUM(B2:D2)"])
    ws.append(["B", 4, 5, 6, "=SUM(B3:D3)"])
    ws.append(["C", 7, 8, 9, 999])
    ws.append(["D", 1, 1, 1, "=SUM(B5:D5)"])
    ws["Z48000"] = "stray"  # the used range becomes 26 x 48,000 > max_cells; the sheet holds 26 cells
    path = tmp_path / "stray.xlsx"
    wb.save(path)
    result = _run(AUDIT, str(path), "--json", "-", "--fail-on", "Critical")
    payload = json.loads(result.stdout)
    assert payload["coverage"]["truncated"]["cells"] is False
    assert "HARDCODE_IN_FORMULA_BLOCK" in {f["rule_id"] for f in payload["findings"]}


def test_context_cards_accept_cells_and_report_unknown_ones(tmp_path):
    findings = tmp_path / "findings.json"
    _run(AUDIT, str(DEMO), "--json", str(findings), "--quiet")
    result = _run(CONTEXT, str(DEMO), str(findings), "Budget!B6", "Nowhere!A1")
    assert result.returncode == 0, result.stderr
    assert "row label: 'Revenue total'" in result.stdout
    assert "no cell context for 'Nowhere!A1'" in result.stdout


def test_context_card_shows_what_a_link_points_at(tmp_path):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    ws["A5"], ws["B5"] = "Income", "=B9"
    ws["A9"], ws["B9"] = "Total income", 100
    path = tmp_path / "links.xlsx"
    wb.save(path)
    findings = tmp_path / "findings.json"
    _run(AUDIT, str(path), "--json", str(findings), "--quiet")
    result = _run(CONTEXT, str(path), str(findings), "S!B5")
    assert result.returncode == 0, result.stderr
    assert "links to: S!B9 (row label: 'Total income') -> 100" in result.stdout


def test_drift_on_a_totals_row_points_at_its_own_column(tmp_path):
    import json

    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    ws.append(["Bill", None, None, None, "Paid", "Total spent", "Day 1", "Day 2", "Day 3"])
    for row in range(2, 7):
        ws.append([f"Bill {row}", None, None, None, row * 10, f"=SUM(G{row}:I{row})", 1, 2, 3])
    ws["A7"], ws["E7"], ws["F7"] = "Total", "=SUM(E2:E6)", "=SUM(G2:G6)"  # F7 adds column G, not F
    path = tmp_path / "totals.xlsx"
    wb.save(path)
    payload = json.loads(_run(AUDIT, str(path), "--json", "-").stdout)
    drift = next(f for f in payload["findings"] if f["rule_id"] == "FORMULA_DRIFT" and f["location"] == "S!F7")
    # The column's majority is a row sum; the other total on row 7 says what F7 should be.
    assert drift["suggested_fix"] == "Point the total at its own column: =SUM(F2:F6)"
    assert any("S!E7==SUM(E2:E6)" in line for line in drift["evidence"])


def test_live_error_finding_carries_its_formula():
    import json

    payload = json.loads(_run(AUDIT, str(DEMO), "--json", "-").stdout)
    live = next(f for f in payload["findings"] if f["rule_id"] == "LIVE_ERROR" and f["location"] == "Budget!B14")
    assert live["formula"] == "=SUM(#REF!)"


def test_libreoffice_profile_is_a_canonical_file_url(tmp_path, monkeypatch):
    from spreadsheet_auditor import recalc

    seen: dict = {}

    def fake_run(cmd, **kwargs):
        seen["cmd"] = cmd
        raise FileNotFoundError("no LibreOffice in this test")

    monkeypatch.setattr(recalc, "soffice_path", lambda: "soffice")
    monkeypatch.setattr(recalc.subprocess, "run", fake_run)
    result = recalc.recalc_if_available(DEMO, work_dir=tmp_path)
    assert result["status"] == "failed"
    arg = next(a for a in seen["cmd"] if a.startswith("-env:UserInstallation="))
    assert arg == "-env:UserInstallation=" + (tmp_path / "profile").resolve().as_uri()
    assert "file:////" not in arg
