"""Fingerprints follow the flagged cell's content, not its address."""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook

from spreadsheet_auditor import audit
from spreadsheet_auditor.finding import Finding
from spreadsheet_auditor.identity import assign_identities


def _model(path: Path, offset: int = 0, new_error_row: int | None = None) -> Path:
    """Nine labeled lines with an accepted #REF! on line 5, shifted down by ``offset`` rows."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Model"
    for line in range(1, 10):
        ws.cell(line + offset, 1, f"Line {line}")
        ws.cell(line + offset, 2, line * 100)
    ws.cell(5 + offset, 3, "=SUM(#REF!)")
    if new_error_row:
        ws.cell(new_error_row, 3, "=B1/#REF!")
    wb.save(path)
    return path


def _audit(workbook: Path, ignore: Path) -> dict:
    out = workbook.with_suffix(".json")
    audit.main([str(workbook), "--json", str(out), "--quiet", "--ignore", str(ignore), "--fail-on", "None"])
    return json.loads(out.read_text(encoding="utf-8"))


def _by_rule(payload: dict, rule: str) -> list[dict]:
    return [f for f in payload["findings"] if f["rule_id"] == rule]


def test_fingerprint_suppression_survives_a_row_insert_and_does_not_hide_a_new_error(tmp_path):
    ignore = tmp_path / ".audit-ignore"
    ignore.write_text("", encoding="utf-8")
    before = _audit(_model(tmp_path / "v1.xlsx"), ignore)
    [accepted] = _by_rule(before, "LIVE_ERROR")
    assert accepted["location"] == "Model!C5"
    ignore.write_text(f"fingerprint:{accepted['fingerprint']} accepted: legacy link, FIN-12\n", encoding="utf-8")

    # A title row goes in above the model; a new #REF! is typed at C5.
    after = _audit(_model(tmp_path / "v2.xlsx", offset=1, new_error_row=5), ignore)
    live = {f["location"]: f for f in _by_rule(after, "LIVE_ERROR")}
    assert live["Model!C6"]["suppressed"] is True
    assert live["Model!C6"]["fingerprint"] == accepted["fingerprint"]
    assert live["Model!C5"]["suppressed"] is False
    assert not [note for note in after["coverage"]["limitations"] if "matched no finding" in note]


def test_a_stale_suppression_is_reported(tmp_path):
    ignore = tmp_path / ".audit-ignore"
    ignore.write_text(
        "LIVE_ERROR Model!C5 accepted: legacy link\n"
        "LIVE_ERROR Model!C7 fixed long ago\n"
        "fingerprint:0123456789abcdef a finding that no longer exists\n",
        encoding="utf-8",
    )
    payload = _audit(_model(tmp_path / "v1.xlsx"), ignore)
    stale = [note for note in payload["coverage"]["limitations"] if "matched no finding" in note]
    assert len(stale) == 2
    assert ".audit-ignore:2) LIVE_ERROR Model!C7 matched no finding" in stale[0]
    assert ".audit-ignore:3) fingerprint:0123456789abcdef matched no finding" in stale[1]


def test_suppressions_that_could_not_match_stay_quiet(tmp_path):
    ignore = tmp_path / ".audit-ignore"
    ignore.write_text("LITERAL_CONSTANT Model!D5 house rate\n", encoding="utf-8")
    config = tmp_path / "config.json"
    config.write_text(json.dumps({"checks": {"LITERAL_CONSTANT": "off"}}), encoding="utf-8")
    workbook = _model(tmp_path / "v1.xlsx")
    out = tmp_path / "f.json"
    audit.main([str(workbook), "--json", str(out), "--quiet", "--ignore", str(ignore), "--config", str(config)])
    limitations = json.loads(out.read_text(encoding="utf-8"))["coverage"]["limitations"]
    assert not [note for note in limitations if "matched no finding" in note]


def _finding(location: str) -> Finding:
    return Finding("FORMULA_DRIFT", "High", "Likely defect", "DET", location, "t", ["e"], "fix")


def test_identity_is_content_plus_labels_numbered_in_sheet_order():
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    ws.cell(1, 3, "Amount")
    for row in (2, 3, 4):
        ws.cell(row, 1, "Same label")
        ws.cell(row, 3, f"=A{row}*B{row}")
    ws.cell(5, 1, "Other label")
    ws.cell(5, 3, "=A5*B5")
    findings = [_finding("S!C4"), _finding("S!C2"), _finding("S!C5"), _finding("S!C3")]
    assign_identities(findings, wb)
    by_location = {f.location: f for f in findings}
    assert by_location["S!C2"].identity == "s|Same label|Amount|=RC[-2]*RC[-1]|1/3"
    assert by_location["S!C3"].identity.endswith("|2/3") and by_location["S!C4"].identity.endswith("|3/3")
    assert by_location["S!C5"].identity == "s|Other label|Amount|=RC[-2]*RC[-1]|1"
    # Only the three alike in everything but their order are marked as sharing.
    assert [by_location[f"S!C{row}"].identity_shared for row in (2, 3, 4, 5)] == [True, True, True, False]
    # A location outside the workbook keeps the location-based fingerprint.
    elsewhere = _finding("Gone!A1")
    assign_identities([elsewhere], wb)
    assert elsewhere.identity is None and elsewhere.fingerprint == _finding("Gone!A1").fingerprint


# --- pinned location lines ----------------------------------------------------


def _notes(payload: dict, text: str) -> list[str]:
    return [note for note in payload["coverage"]["limitations"] if text in note]


def test_a_pinned_line_follows_its_finding_and_never_hides_a_new_one(tmp_path):
    ignore = tmp_path / ".audit-ignore"
    ignore.write_text("", encoding="utf-8")
    before = _audit(_model(tmp_path / "v1.xlsx"), ignore)
    [accepted] = _by_rule(before, "LIVE_ERROR")
    # The line the report hands out for this finding.
    ignore.write_text(f"LIVE_ERROR Model!C5 fingerprint:{accepted['fingerprint']} accepted: legacy link\n", encoding="utf-8")
    assert _by_rule(_audit(tmp_path / "v1.xlsx", ignore), "LIVE_ERROR")[0]["suppressed"] is True

    # A title row goes in above the model and a new #REF! is typed at C5:
    # the line keeps the accepted error suppressed at C6, leaves the new one
    # reported, and says the address is out of date.
    after = _audit(_model(tmp_path / "v2.xlsx", offset=1, new_error_row=5), ignore)
    live = {f["location"]: f["suppressed"] for f in _by_rule(after, "LIVE_ERROR")}
    assert live == {"Model!C5": False, "Model!C6": True}
    [moved] = _notes(after, "follows its finding")
    assert "now at Model!C6" in moved
    assert not _notes(after, "matched no finding")


def test_a_pinned_line_whose_finding_is_gone_names_what_is_there_now(tmp_path):
    ignore = tmp_path / ".audit-ignore"
    ignore.write_text("LIVE_ERROR Model!C5 fingerprint:0123456789abcdef accepted long ago\n", encoding="utf-8")
    payload = _audit(_model(tmp_path / "v1.xlsx"), ignore)
    [live] = _by_rule(payload, "LIVE_ERROR")
    assert live["location"] == "Model!C5" and live["suppressed"] is False
    [stale] = _notes(payload, "matched no finding")
    assert f"Model!C5 now holds a different LIVE_ERROR finding, which is reported (fingerprint {live['fingerprint']})" in stale


def test_an_unpinned_one_cell_line_names_the_pinned_line_to_use(tmp_path):
    ignore = tmp_path / ".audit-ignore"
    ignore.write_text(
        "LIVE_ERROR Model!C5 accepted: legacy link\n"
        "BROKEN_REFERENCE Model!C1:C9 legacy block\n",
        encoding="utf-8",
    )
    payload = _audit(_model(tmp_path / "v1.xlsx"), ignore)
    [live] = _by_rule(payload, "LIVE_ERROR")
    assert live["suppressed"] is True
    [hint] = _notes(payload, "follows the address")
    assert hint.endswith(f"LIVE_ERROR Model!C5 fingerprint:{live['fingerprint']} accepted: legacy link")
    assert "Model!C1:C9" not in hint  # an area is meant to follow its address


def test_pinned_lines_parse_and_match_only_their_rule(tmp_path):
    from spreadsheet_auditor.suppressions import apply_suppressions, load_suppressions

    ignore = tmp_path / ".audit-ignore"
    finding = _finding("S!C2")
    ignore.write_text(
        f"FORMULA_DRIFT 'Q3 Budget'!C2 FINGERPRINT:{finding.fingerprint.upper()} approved override\n"
        f"LITERAL_CONSTANT S!C2 fingerprint:{finding.fingerprint} wrong rule\n"
        "FORMULA_DRIFT S!C2 fingerprint: no pin\n",
        encoding="utf-8",
    )
    warnings: list[str] = []
    pinned, wrong_rule = load_suppressions(None, str(ignore), warnings=warnings)
    assert pinned["range"] == "'Q3 Budget'!C2"
    assert pinned["fingerprint"] == finding.fingerprint and pinned["reason"] == "approved override"
    assert "fingerprint:<fp>" in warnings[0]
    apply_suppressions([finding], [wrong_rule])
    assert not finding.suppressed
    # The fingerprint decides, not the address: the line names another sheet.
    apply_suppressions([finding], [pinned])
    assert finding.suppressed


def test_report_hands_out_the_pinned_line_with_the_sheet_quoted():
    from spreadsheet_auditor.locations import format_target
    from spreadsheet_auditor.report import _render_finding

    finding = _finding("Revenue Detail!B7, Revenue Detail!B9").to_dict()
    finding["id"] = "FORMULA_DRIFT-001"
    lines = _render_finding(finding)
    expected = f"FORMULA_DRIFT 'Revenue Detail'!B7 fingerprint:{finding['fingerprint']} <reason>"
    assert f"- If intended, accept it in `.audit-ignore`: `{expected}`" in lines
    assert format_target("Budget!B14") == "Budget!B14"
    assert format_target("O'Brien!A1") == "'O''Brien'!A1"
    assert format_target("Q1!B2") == "'Q1'!B2"


# --- findings alike in all but their order -----------------------------------


def _plugs(path: Path, plugs: list[int], headers: bool = True) -> Path:
    """A Revenue row growing 10% a month, with hard-coded plugs in the given columns."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Model"
    if headers:
        for col in range(2, 15):
            ws.cell(4, col, f"Month {col - 1}")
    ws.cell(5, 1, "Revenue")
    ws.cell(5, 2, 1000)
    for col in range(3, 15):
        ws.cell(5, col, f"={ws.cell(5, col - 1).column_letter}5*1.1")
    for col in plugs:
        ws.cell(5, col, 500)
    wb.save(path)
    return path


def _plug_findings(payload: dict) -> dict[str, dict]:
    return {f["location"]: f for f in _by_rule(payload, "HARDCODE_IN_FORMULA_BLOCK")}


def test_column_headers_tell_plugs_in_one_row_apart(tmp_path):
    ignore = tmp_path / ".audit-ignore"
    ignore.write_text("", encoding="utf-8")
    before = _plug_findings(_audit(_plugs(tmp_path / "v1.xlsx", [6, 10]), ignore))
    accepted = before["Model!F5"]
    assert accepted["fingerprint"] != before["Model!J5"]["fingerprint"]
    ignore.write_text(f"HARDCODE_IN_FORMULA_BLOCK Model!F5 fingerprint:{accepted['fingerprint']} agreed plug\n", encoding="utf-8")

    # A new plug typed at D5 is reported; the accepted one stays hidden.
    after = _plug_findings(_audit(_plugs(tmp_path / "v2.xlsx", [4, 6, 10]), ignore))
    assert {loc: f["suppressed"] for loc, f in after.items()} == {"Model!D5": False, "Model!F5": True, "Model!J5": False}

    # The accepted plug is replaced by its formula: J5 is not hidden in its place.
    fixed = _audit(_plugs(tmp_path / "v3.xlsx", [10]), ignore)
    assert {loc: f["suppressed"] for loc, f in _plug_findings(fixed).items()} == {"Model!J5": False}
    assert _notes(fixed, "matched no finding")


def test_a_pinned_line_never_passes_to_a_finding_alike_but_for_its_order(tmp_path):
    # No headers: the plugs share row label and value, so only their order tells them apart.
    ignore = tmp_path / ".audit-ignore"
    ignore.write_text("", encoding="utf-8")
    before = _plug_findings(_audit(_plugs(tmp_path / "v1.xlsx", [6, 10], headers=False), ignore))
    accepted = before["Model!F5"]
    ignore.write_text(f"HARDCODE_IN_FORMULA_BLOCK Model!F5 fingerprint:{accepted['fingerprint']} agreed plug\n", encoding="utf-8")
    assert _plug_findings(_audit(tmp_path / "v1.xlsx", ignore))["Model!F5"]["suppressed"] is True

    # A new plug at D5 would take F5's number, as would J5 once F5 is fixed, or
    # D5 when it replaces F5: the line hides none of them, and says it matched
    # nothing, so the reader re-pins what they still accept.
    for name, plugs in (("v2.xlsx", [4, 6, 10]), ("v3.xlsx", [10]), ("v4.xlsx", [4, 10])):
        payload = _audit(_plugs(tmp_path / name, plugs, headers=False), ignore)
        assert not [loc for loc, f in _plug_findings(payload).items() if f["suppressed"]], name
        assert _notes(payload, "matched no finding"), name


def test_identity_ignores_what_an_insert_elsewhere_moves():
    from spreadsheet_auditor.identity import assign_identities

    def identity(offset: int) -> str:
        wb = Workbook()
        ws = wb.active
        ws.title = "Model"
        inputs = wb.create_sheet("Inputs")
        inputs.cell(5 + offset, 2, 0.2)
        ws.cell(1 + offset, 2, 1.07)
        ws.cell(5 + offset, 1, "Revenue")
        ws.cell(5 + offset, 3, f"=A{5 + offset}*$B${1 + offset}*Inputs!B{5 + offset}")
        finding = _finding(f"Model!C{5 + offset}")
        assign_identities([finding], wb)
        return finding.identity

    assert identity(0) == identity(3)


def test_a_data_table_cell_has_the_same_identity_every_run():
    from openpyxl.worksheet.formula import DataTableFormula

    from spreadsheet_auditor.identity import assign_identities

    identities = set()
    for _run in range(2):
        wb = Workbook()
        ws = wb.active
        ws.title = "S"
        ws["B5"] = DataTableFormula(ref="B5:B9", r1="A1")
        finding = _finding("S!B5")
        assign_identities([finding], wb)
        identities.add(finding.identity)
    [identity] = identities
    assert " at 0x" not in identity


def test_a_repeated_key_list_is_a_new_finding_when_another_key_repeats(tmp_path):
    def keys(path: Path, values: list[str]) -> Path:
        wb = Workbook()
        ws = wb.active
        ws.title = "S"
        for row, value in enumerate(values, start=1):
            ws.cell(row, 1, value)
            ws.cell(row, 2, row)
        ws["D1"] = "=VLOOKUP(\"apple\",A1:B20,2,FALSE)"
        wb.save(path)
        return path

    ignore = tmp_path / ".audit-ignore"
    ignore.write_text("", encoding="utf-8")
    [before] = _by_rule(_audit(keys(tmp_path / "v1.xlsx", ["apple", "pear", "apple", "pear"]), ignore), "DUPLICATE_KEY")
    assert before["members"] == ["S!A1", "S!A3", "S!A2", "S!A4"]
    ignore.write_text(f"DUPLICATE_KEY S!A1 fingerprint:{before['fingerprint']} first match is intended\n", encoding="utf-8")
    assert _by_rule(_audit(tmp_path / "v1.xlsx", ignore), "DUPLICATE_KEY")[0]["suppressed"] is True

    # A plum duplicate added later is not hidden by the line accepting apple and pear.
    grown = keys(tmp_path / "v2.xlsx", ["apple", "pear", "apple", "pear", "plum", "plum"])
    [after] = _by_rule(_audit(grown, ignore), "DUPLICATE_KEY")
    assert after["suppressed"] is False


def test_a_quoted_sheet_name_may_hold_an_exclamation_mark():
    from spreadsheet_auditor.suppressions import apply_suppressions

    finding = _finding("Wow!!C2")
    apply_suppressions([finding], [{"rule_id": "formula_drift", "range": "'Wow!'", "reason": "r"}])
    assert finding.suppressed
    finding.suppressed = False
    apply_suppressions([finding], [{"rule_id": "FORMULA_DRIFT", "range": "'Wow!'!C2", "reason": "r"}])
    assert finding.suppressed
