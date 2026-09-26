"""Regressions for truncated audit status, additive cross-footing, and output aliases."""
from __future__ import annotations

import argparse
import json
import os

import pytest
from openpyxl import Workbook

from spreadsheet_auditor.audit import audit_workbook, main
from spreadsheet_auditor.reconcile import detect_cross_foot_failures


def test_suppressed_findings_cannot_hide_failure_after_output_cap(tmp_path):
    wb = Workbook()
    wb.active.title = "S"
    wb.active["A1"] = "=SUM(#REF!)"
    workbook = tmp_path / "broken.xlsx"
    wb.save(workbook)
    config = tmp_path / "config.json"
    config.write_text(json.dumps({
        "recalc": {"enabled": False},
        "limits": {"max_reported_findings": 1},
        "suppressions": [{"rule_id": "LIVE_ERROR", "range": "S", "reason": "Known issue"}],
    }), encoding="utf-8")
    payload, code = audit_workbook(argparse.Namespace(
        workbook=str(workbook), config=str(config), ignore=str(tmp_path / "none"),
        fail_on="High", strict=False, recalc_timeout=None,
    ))
    assert code == 1
    assert payload["coverage"]["truncated"]["findings"]
    assert [f["rule_id"] for f in payload["findings"] if not f["suppressed"]] == ["BROKEN_REFERENCE"]


@pytest.mark.parametrize("style", ["AVERAGE", "scaled SUM"])
def test_cross_foot_does_not_treat_averages_as_additive_totals(style):
    formula_wb, value_wb = Workbook(), Workbook()
    ws, values = formula_wb.active, value_wb.active
    ws.title = values.title = "Survey"
    for row, first, second in ((2, 10, 20), (3, 30, 40), (4, 50, 60)):
        ws[f"B{row}"] = values[f"B{row}"] = first
        ws[f"C{row}"] = values[f"C{row}"] = second
        ws[f"D{row}"] = f"=AVERAGE(B{row}:C{row})" if style == "AVERAGE" else f"=SUM(B{row}:C{row})/2"
        values[f"D{row}"] = (first + second) / 2
    for col, average in (("B", 30), ("C", 40)):
        ws[f"{col}5"] = f"=AVERAGE({col}2:{col}4)" if style == "AVERAGE" else f"=SUM({col}2:{col}4)/3"
        values[f"{col}5"] = average
    assert detect_cross_foot_failures(formula_wb, value_wb) == []


@pytest.mark.parametrize("flags", [("--out", "--json"), ("--out", "--annotated"), ("--json", "--annotated")])
def test_output_destinations_must_be_distinct_before_any_write(tmp_path, flags):
    source = tmp_path / "source.xlsx"
    wb = Workbook()
    wb.active["A1"] = 1
    wb.save(source)
    config = tmp_path / "config.json"
    config.write_text('{"recalc": {"enabled": false}}', encoding="utf-8")
    target = tmp_path / "report.xlsx"
    assert main([str(source), "--config", str(config), flags[0], str(target), flags[1], str(target)]) == 4
    assert not target.exists()


def test_output_hardlink_aliases_must_be_distinct(tmp_path):
    source = tmp_path / "source.xlsx"
    wb = Workbook()
    wb.active["A1"] = 1
    wb.save(source)
    first, second = tmp_path / "first.md", tmp_path / "second.json"
    first.write_text("keep both", encoding="utf-8")
    try:
        os.link(first, second)
    except OSError:
        pytest.skip("hardlinks unavailable")
    assert main([str(source), "--out", str(first), "--json", str(second)]) == 4
    assert first.read_text(encoding="utf-8") == "keep both"
    assert second.read_text(encoding="utf-8") == "keep both"
