"""A mistake repeated down a column is one finding, and the report cap never hides a whole rule."""

from __future__ import annotations

import json
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook

from spreadsheet_auditor import audit
from spreadsheet_auditor.finding import Finding


def _no_suppressions(tmp_path: Path) -> Path:
    """An empty suppression file, so no `.audit-ignore` in the working folder applies."""
    path = tmp_path / "no-suppressions"
    path.touch()
    return path


def _run(path: Path, tmp_path: Path, config: dict | None = None) -> dict:
    out = tmp_path / "findings.json"
    args = [str(path), "--json", str(out), "--quiet", "--fail-on", "None", "--ignore", str(_no_suppressions(tmp_path))]
    if config is not None:
        (tmp_path / "config.json").write_text(json.dumps(config), encoding="utf-8")
        args += ["--config", str(tmp_path / "config.json")]
    audit.main(args)
    return json.loads(out.read_text(encoding="utf-8"))


def _by_rule(payload: dict) -> dict[str, list[dict]]:
    found: dict[str, list[dict]] = defaultdict(list)
    for finding in payload["findings"]:
        found[finding["rule_id"]].append(finding)
    return found


def test_an_error_filled_down_a_column_is_one_finding(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    for row in range(2, 22):
        ws.cell(row, 1, row)
        ws.cell(row, 2, f"=A{row}+#N/A")  # 20 cells, one relative formula
    ws["B22"] = "=A21*#DIV/0!"  # a different formula and error: its own finding
    for row in (30, 31, 32):
        ws.cell(row, 4, "#N/A")  # the same error typed as a value
    path = tmp_path / "filled.xlsx"
    wb.save(path)
    live = {f["location"]: f for f in _by_rule(_run(path, tmp_path))["LIVE_ERROR"]}
    assert sorted(live) == ["S!B2", "S!B22", "S!D30"]
    assert "evaluates to #N/A in 20 cells on this sheet; the others are S!B3, S!B4" in live["S!B2"]["evidence"][1]
    assert live["S!B2"]["evidence"][1].endswith("S!B10, and 11 more.")
    assert "#N/A is typed as a value in 3 cells" in live["S!D30"]["evidence"][1]
    assert len(live["S!B22"]["evidence"]) == 1


def test_errors_in_a_lookup_that_reads_its_own_column_are_one_finding(tmp_path):
    # Every row's lookup table spans column E, so each erroring cell reads the
    # others: a cycle with no root error. SpreadsheetBench 54196 reported 278
    # findings this way.
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet two"
    for row in range(2, 7):
        ws.cell(row, 4, f"code {row}")
        ws.cell(row, 5, f"=VLOOKUP(G{row},$D$2:$E$6,2,FALSE)+#N/A")
    path = tmp_path / "lookup.xlsx"
    wb.save(path)
    found = _by_rule(_run(path, tmp_path))
    [live] = found["LIVE_ERROR"]
    assert live["location"] == "Sheet two!E2"
    assert "evaluates to #N/A in 5 cells on this sheet" in live["evidence"][1]
    assert len(found["CIRCULAR_REFERENCE"]) == 1


def test_broken_and_self_referencing_formulas_group_by_pattern(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    for row in range(2, 12):
        ws.cell(row, 1, row)
        ws.cell(row, 2, f"=SUM(#REF!)+A{row}")
        ws.cell(row, 3, f"=A{row}+C{row}")  # each cell reads itself
    path = tmp_path / "broken.xlsx"
    wb.save(path)
    found = _by_rule(_run(path, tmp_path))
    [broken] = found["BROKEN_REFERENCE"]
    assert broken["location"] == "S!B2" and "contains #REF! in 10 cells" in broken["evidence"][1]
    [cycle] = found["CIRCULAR_REFERENCE"]
    assert cycle["location"] == "S!C2" and cycle["formula"] == "=A2+C2"
    assert "references its own cell in 10 cells" in cycle["evidence"][1]


def test_duplicate_keys_are_one_finding_per_searched_column(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    keys = ["apple", "pear", "apple", "plum", "pear", "fig", "pear"]
    for row, key in enumerate(keys, start=1):
        ws.cell(row, 1, key)
        ws.cell(row, 2, row * 10)
    ws["D1"] = "=VLOOKUP(\"pear\",A1:B7,2,FALSE)"
    path = tmp_path / "keys.xlsx"
    wb.save(path)
    [dup] = _by_rule(_run(path, tmp_path))["DUPLICATE_KEY"]
    assert dup["location"] == "S!A1, S!A3" and dup["title"] == "Duplicate keys in lookup range"
    assert "2 keys repeat in column A; the others are 'pear' 3 times (A2, A5, A7)." in dup["evidence"][1]


def _finding(rule: str, severity: str, location: str) -> Finding:
    return Finding(rule, severity, "Review", "DET", location, "t", ["e"], "fix")


def test_the_cap_keeps_every_rule_and_counts_what_it_trims():
    findings = [_finding("LIVE_ERROR", "Critical", f"S!A{row}") for row in range(1, 8)]
    findings += [_finding("VOLATILE_FUNCTION", "Low", "S!Z1"), _finding("WHITESPACE_KEY", "Medium", "S!Y1")]
    truncated = {"findings": False}
    notes: list[str] = []
    kept, counts = audit._cap_findings(audit.sort_findings(findings), 4, truncated, notes)
    rules = [f.rule_id for f in kept]
    # Each rule keeps its most severe finding; the rest go by severity.
    assert rules.count("VOLATILE_FUNCTION") == 1 and rules.count("WHITESPACE_KEY") == 1
    assert rules.count("LIVE_ERROR") == 2 and len(kept) == 4
    assert truncated["findings"] is True
    assert counts["total"] == 9 and counts["shown"] == 4
    assert counts["by_severity"] == {"Critical": 7, "Medium": 1, "Low": 1}
    assert notes == [
        "Findings output capped at 4 of 9 findings by config; not shown: 5 LIVE_ERROR. "
        "Raise limits.max_reported_findings to see them all."
    ]


def test_summary_and_reports_count_before_the_cap(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    ops = ["+", "-", "*", "/", "^"]
    for row, op in enumerate(ops, start=1):
        ws.cell(row, 1, row)
        ws.cell(row, 2, f"=A{row}{op}#N/A")  # five different formulas, five findings
    ws["C1"] = "=TODAY()"
    path = tmp_path / "capped.xlsx"
    wb.save(path)
    payload = _run(path, tmp_path, {"limits": {"max_reported_findings": 3}})
    assert len(payload["findings"]) == 3
    assert "VOLATILE_FUNCTION" in {f["rule_id"] for f in payload["findings"]}
    counts = payload["coverage"]["finding_counts"]
    assert (counts["total"], counts["shown"], counts["by_severity"]["Critical"]) == (6, 3, 5)
    summary = audit._summary_lines(payload, "Critical")
    assert "findings   : 5 Critical, 0 High, 0 Medium, 1 Low, 0 Info (3 of 6 shown)" in summary
    assert "     5  LIVE_ERROR (Critical, Defect)" in summary
    from spreadsheet_auditor.report import render_markdown

    assert "5 Critical, 0 High, 0 Medium, 1 Low, 0 Info (3 of 6 shown; see limitations)" in render_markdown(payload)


def test_csv_whitespace_is_one_finding_per_column_and_capped(tmp_path):
    rows = ["id,name,notes"] + [f"{i},Item {i},see memo " for i in range(1, 1501)]
    rows[40] = "40,Item 40 ,see memo "  # one stray padded key among clean names
    path = tmp_path / "export.csv"
    path.write_text("\n".join(rows), encoding="utf-8")
    payload = _run(path, tmp_path, {"limits": {"max_reported_findings": 1}})
    counts = payload["coverage"]["finding_counts"]
    assert (counts["total"], counts["shown"]) == (2, 1)
    [shown] = payload["findings"]
    assert (shown["location"], shown["severity"]) == ("CSV!R41C2", "Medium")
    payload = _run(path, tmp_path)
    notes = next(f for f in payload["findings"] if f["location"] == "CSV!R2C3")
    assert (notes["severity"], notes["error_confidence"]) == ("Low", "Info")
    assert "1500 of 1501 values in column 3 carry leading or trailing whitespace" in notes["evidence"][0]


def test_hidden_columns_of_one_sheet_are_one_finding(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    for row in range(2, 6):
        for col in range(1, 7):
            ws.cell(row, col, row * col)
    # Each hidden column feeds its own formula; they used to be one finding each.
    ws["H2"], ws["H3"], ws["H4"] = "=C2*2", "=D3*2", "=SUM(E2:E5)"
    for letter in ("C", "D", "E"):
        ws.column_dimensions[letter].hidden = True
    ws.row_dimensions[5].hidden = True
    path = tmp_path / "hidden.xlsx"
    wb.save(path)
    hidden = {f["evidence"][0] for f in _by_rule(_run(path, tmp_path))["HIDDEN_STRUCTURE_IN_TOTAL"]}
    assert hidden == {
        "Hidden columns C to E on S feed 3 visible formula(s): S!H2, S!H3, S!H4.",
        "Hidden row 5 on S feeds 1 visible formula(s): S!H4.",
    }


# --- a grouped finding stands for all of its cells ----------------------------


def _broken_column(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    for row in range(2, 22):
        ws.cell(row, 1, row)
        ws.cell(row, 2, f"=A{row}+SUM(#REF!)")
    wb.save(path)
    return path


def _run_with(path: Path, tmp_path: Path, ignore_text: str, config: dict | None = None) -> dict:
    ignore = tmp_path / "suppressions"
    ignore.write_text(ignore_text, encoding="utf-8")
    out = tmp_path / "findings.json"
    args = [str(path), "--json", str(out), "--quiet", "--fail-on", "None", "--ignore", str(ignore)]
    if config is not None:
        (tmp_path / "config.json").write_text(json.dumps(config), encoding="utf-8")
        args += ["--config", str(tmp_path / "config.json")]
    audit.main(args)
    return json.loads(out.read_text(encoding="utf-8"))


def test_a_suppression_hides_a_grouped_finding_only_when_it_covers_every_cell(tmp_path):
    path = _broken_column(tmp_path / "broken.xlsx")
    [finding] = _by_rule(_run(path, tmp_path))["BROKEN_REFERENCE"]
    assert finding["members"] == [f"S!B{row}" for row in range(2, 22)]

    for target in ("S!B2:B3", "S!B10:B21"):
        payload = _run_with(path, tmp_path, f"BROKEN_REFERENCE {target} only these rows\n")
        [finding] = _by_rule(payload)["BROKEN_REFERENCE"]
        assert finding["suppressed"] is False, target
        notes = payload["coverage"]["limitations"]
        [partial] = [note for note in notes if "so the finding stays reported" in note]
        assert "cover all of them (S!B2:B21)" in partial and "fingerprint:" in partial
        assert not [note for note in notes if "matched no finding" in note], target

    for target in ("S!B2:B21", "S!B:B", "S"):
        payload = _run_with(path, tmp_path, f"BROKEN_REFERENCE {target} known broken block\n")
        assert _by_rule(payload)["BROKEN_REFERENCE"][0]["suppressed"] is True, target


def test_any_cell_of_a_grouped_finding_can_be_the_headline_output(tmp_path):
    path = _broken_column(tmp_path / "broken.xlsx")
    for headline in ("S!B2", "S!B10"):
        payload = _run(path, tmp_path, {"scope": {"headline_outputs": [headline]}})
        [finding] = _by_rule(payload)["BROKEN_REFERENCE"]
        assert finding["severity"] == "Critical", headline
        assert finding["impact"]["feeds_headline_output"] is True


def test_annotated_copy_and_sarif_mark_every_cell_of_a_grouped_finding(tmp_path):
    from openpyxl import load_workbook

    from spreadsheet_auditor.sarif import render_sarif

    path = _broken_column(tmp_path / "broken.xlsx")
    annotated = tmp_path / "annotated.xlsx"
    audit.main([str(path), "--quiet", "--fail-on", "None", "--ignore", str(_no_suppressions(tmp_path)),
                "--annotated", str(annotated)])
    ws = load_workbook(annotated)["S"]
    assert "same issue is in 19 other cell(s)" in ws["B2"].comment.text
    assert "Same issue as S!B2" in ws["B21"].comment.text

    payload = _run(path, tmp_path)
    [result] = [r for r in render_sarif(payload)["runs"][0]["results"] if r["ruleId"] == "BROKEN_REFERENCE"]
    related = [loc["logicalLocations"][0]["name"] for loc in result["relatedLocations"]]
    assert related == [f"S!B{row}" for row in range(3, 22)]
    assert result["properties"]["members"][0] == "S!B2"
