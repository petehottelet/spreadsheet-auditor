"""A mistake in the input exits 4 and says what to fix; a bug in the auditor exits 5 with its traceback."""

from __future__ import annotations

import codecs
import csv
import json
import os
import stat
import zipfile
from pathlib import Path

import pytest
from openpyxl import Workbook

from spreadsheet_auditor import audit, workbook_inventory
from spreadsheet_auditor.config_loader import DEFAULT_CONFIG, load_config


def _workbook(tmp_path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Model"
    ws["A1"], ws["B1"] = "Revenue", 100
    ws["A2"], ws["B2"] = "Error", "=SUM(#REF!)"
    path = tmp_path / "model.xlsx"
    wb.save(path)
    return path


def _config(tmp_path: Path, text: str, name: str = "config.json") -> Path:
    path = tmp_path / name
    path.write_text(text, encoding="utf-8")
    return path


def _audit(tmp_path: Path, *args: str) -> int:
    ignore = tmp_path / "no-suppressions"
    ignore.touch()
    return audit.main([*args, "--quiet", "--ignore", str(ignore)])


def _json_audit(tmp_path: Path, source: Path, *args: str) -> dict:
    out = tmp_path / "findings.json"
    _audit(tmp_path, str(source), "--json", str(out), "--fail-on", "None", *args)
    return json.loads(out.read_text(encoding="utf-8"))


@pytest.mark.parametrize(
    ("config", "message"),
    [
        ('{"checks": {"LIVE_EROR": "off"}}', "Unknown rule 'LIVE_EROR' in config; did you mean 'LIVE_ERROR'?"),
        ('{"checks": {"Volatile_Functon": "off"}}', "did you mean 'VOLATILE_FUNCTION'?"),
        ('{"limit": {"max_formulas": 10}}', "Unknown section 'limit' in config; did you mean 'limits'?"),
        ('{"limits": {"max_finding": 10}}', "did you mean 'max_reported_findings'?"),
        ('{"limits": {"max_reported_findings": "200"}}', "limits.max_reported_findings must be a whole number"),
        ('{"limits": {"max_formulas": 50000.5}}', "limits.max_formulas must be a whole number of 0 or more, not 50000.5"),
        ('{"scope": {"include_sheets": [2024]}}', 'quote a sheet name that looks like a number, as in ["2024"]'),
        ('{"recalc": {"timeout_seconds": 0}}', "recalc.timeout_seconds must be a whole number of 1 or more"),
        ('{"recalc": {"enabled": "no"}}', "recalc.enabled must be true or false, not 'no'"),
        ('{"checks": {"LIVE_ERROR": "maybe"}}', "checks.LIVE_ERROR must be true, false, or one of"),
        ('{"checks": ["LIVE_ERROR"]}', "'checks' must map rule names to settings"),
        ('{"suppressions": {"rule_id": "LIVE_ERROR"}}', "'suppressions' must be a list"),
        ('["checks"]', "must hold a mapping of sections"),
        ('{"checks": {"LIVE_ERROR": "off",}}', "is not valid JSON: Expecting property name enclosed in double quotes at line 1"),
        ("", "is not valid JSON"),
    ],
)
def test_a_config_mistake_exits_4_and_names_the_fix(tmp_path, capsys, config, message):
    code = _audit(tmp_path, str(_workbook(tmp_path)), "--config", str(_config(tmp_path, config)))
    err = capsys.readouterr().err
    assert code == 4, err
    assert err.startswith("Config error: ")
    assert message in err
    assert "Traceback" not in err


def test_a_json_syntax_error_names_its_line(tmp_path, capsys):
    config = _config(tmp_path, '{\n  "recalc": {"enabled": false}\n  "limits": {}\n}')
    assert _audit(tmp_path, str(_workbook(tmp_path)), "--config", str(config)) == 4
    assert "Expecting ',' delimiter at line 3, column 3" in capsys.readouterr().err


def test_a_missing_or_misnamed_config_exits_4(tmp_path, capsys):
    workbook = str(_workbook(tmp_path))
    assert _audit(tmp_path, workbook, "--config", str(tmp_path / "missing.json")) == 4
    assert "Config file not found:" in capsys.readouterr().err
    assert _audit(tmp_path, workbook, "--config", str(_config(tmp_path, "{}", "config.toml"))) == 4
    assert "must be .json, .yml or .yaml" in capsys.readouterr().err


def test_a_config_that_is_not_utf8_exits_4(tmp_path, capsys):
    config = tmp_path / "config.json"
    config.write_bytes('{"mode": "café"}'.encode("cp1252"))
    assert _audit(tmp_path, str(_workbook(tmp_path)), "--config", str(config)) == 4
    assert "is not UTF-8 text" in capsys.readouterr().err


def test_a_yaml_syntax_error_exits_4(tmp_path, capsys):
    pytest.importorskip("yaml")
    config = _config(tmp_path, "checks:\n  LIVE_ERROR: off\n limits: [\n", "config.yml")
    assert _audit(tmp_path, str(_workbook(tmp_path)), "--config", str(config)) == 4
    assert "is not valid YAML" in capsys.readouterr().err


def test_check_names_and_settings_match_in_any_case(tmp_path):
    workbook = _workbook(tmp_path)
    for checks in ({"live_error": "OFF"}, {"Live_Errors": False}):
        config = _config(tmp_path, json.dumps({"recalc": {"enabled": False}, "checks": checks}))
        payload = _json_audit(tmp_path, workbook, "--config", str(config))
        assert "LIVE_ERROR" not in {finding["rule_id"] for finding in payload["findings"]}, checks


def test_a_complete_valid_config_is_accepted(tmp_path):
    config = {
        "mode": "financial_model",
        "scope": {"include_sheets": ["Model"], "exclude_sheets": [], "headline_outputs": ["Model!B1"]},
        "materiality": {"absolute": 1000, "relative": 0.001, "percent_points": 0.1},
        "limits": {"max_cells": 0, "max_formulas": 100, "max_range_expansion_cells": 1000,
                   "max_reported_findings": 50, "timeout_seconds": 0},
        "recalc": {"enabled": False, "timeout_seconds": 30},
        "finance": {"enabled": True},
        "checks": {"LIVE_ERROR": "warn", "blank_precedent": True, "volatile_function": "off"},
        "suppressions": [],
    }
    path = _config(tmp_path, json.dumps(config))
    assert _audit(tmp_path, str(_workbook(tmp_path)), "--config", str(path)) == 0  # LIVE_ERROR is only a warning


def test_a_windows_1252_csv_is_audited_with_a_note(tmp_path, capsys):
    source = tmp_path / "export.csv"
    source.write_bytes("Name,Amount\r\ncafé ,10\r\nthé,20\r\n".encode("cp1252"))
    payload = _json_audit(tmp_path, source)
    assert any("read as Windows-1252" in note for note in payload["coverage"]["limitations"])
    [finding] = payload["findings"]
    assert finding["location"] == "CSV!R2C1"
    assert "'café '" in finding["evidence"][0]


def test_a_csv_in_neither_encoding_exits_4(tmp_path, capsys):
    source = tmp_path / "export.csv"
    source.write_bytes(b"Name,Amount\r\n\x81\x8d,10\r\n")  # bytes undefined in both UTF-8 and Windows-1252
    assert _audit(tmp_path, str(source)) == 4
    assert "Preflight failed: The CSV is neither UTF-8 nor Windows-1252" in capsys.readouterr().err


def test_a_suppression_file_that_is_not_utf8_exits_4(tmp_path, capsys):
    ignore = tmp_path / ".audit-ignore"
    ignore.write_bytes("LIVE_ERROR Model!B2 accepted by José\n".encode("cp1252"))
    code = audit.main([str(_workbook(tmp_path)), "--quiet", "--ignore", str(ignore)])
    assert code == 4
    assert "is not UTF-8 text; save it as UTF-8" in capsys.readouterr().err


def test_a_workbook_whose_contents_do_not_parse_exits_4(tmp_path, capsys):
    source = tmp_path / "broken.xlsx"
    with zipfile.ZipFile(source, "w") as archive:
        archive.writestr("[Content_Types].xml", "not xml")
    assert _audit(tmp_path, str(source)) == 4
    err = capsys.readouterr().err
    assert err.startswith("Preflight failed: The workbook could not be read (")
    assert "save a copy, and audit the copy" in err


def test_a_folder_named_like_a_workbook_exits_4(tmp_path, capsys):
    folder = tmp_path / "model.xlsx"
    folder.mkdir()
    assert _audit(tmp_path, str(folder)) == 4
    assert f"Preflight failed: Could not read {folder}" in capsys.readouterr().err


def test_an_output_path_in_a_missing_folder_exits_4(tmp_path, capsys):
    report = tmp_path / "no-such-folder" / "report.md"
    assert _audit(tmp_path, str(_workbook(tmp_path)), "--out", str(report), "--fail-on", "None") == 4
    err = capsys.readouterr().err
    assert err.startswith("Could not write output: ")
    assert "Traceback" not in err


def test_a_bug_while_auditing_exits_5_with_its_traceback(tmp_path, monkeypatch, capsys):
    def broken(args):
        raise KeyError("sheet_index")

    monkeypatch.setattr(audit, "audit_workbook", broken)
    assert _audit(tmp_path, str(_workbook(tmp_path))) == 5
    err = capsys.readouterr().err
    assert "Internal audit error while auditing: KeyError: 'sheet_index'" in err
    assert "This is a bug in spreadsheet-auditor" in err
    assert "Traceback (most recent call last):" in err
    assert "in broken" in err


def test_a_bug_while_writing_output_exits_5_with_its_traceback(tmp_path, monkeypatch, capsys):
    def broken(*args, **kwargs):
        raise TypeError("render_markdown() got an unexpected keyword argument")

    monkeypatch.setattr(audit, "_write_report", broken)
    code = _audit(tmp_path, str(_workbook(tmp_path)), "--out", str(tmp_path / "report.md"), "--fail-on", "None")
    err = capsys.readouterr().err
    assert code == 5
    assert "Internal audit error while writing output: TypeError: render_markdown()" in err
    assert "Traceback (most recent call last):" in err


def _findings(payload: dict) -> list[tuple]:
    return [(f["rule_id"], f["location"], f["evidence"]) for f in payload["findings"]]


@pytest.mark.parametrize(
    ("name", "text"),
    [
        ("config.json", '{"checks": null, "scope": null, "limits": null, "suppressions": null, "recalc": null}'),
        ("config.yml", "checks:\nsuppressions:\n  # none yet\nscope:\nlimits:\n"),
    ],
)
def test_an_empty_config_section_is_the_same_as_leaving_it_out(tmp_path, name, text):
    if name.endswith(".yml"):
        pytest.importorskip("yaml")
    workbook = _workbook(tmp_path)
    config = _config(tmp_path, text, name)
    assert load_config(str(config), known_rules=audit._known_rules())["limits"] == DEFAULT_CONFIG["limits"]
    assert _audit(tmp_path, str(workbook), "--config", str(config)) == _audit(tmp_path, str(workbook)) == 1
    with_config = _json_audit(tmp_path, workbook, "--config", str(config))
    assert _findings(with_config) == _findings(_json_audit(tmp_path, workbook)) != []


def test_a_whole_number_written_as_a_float_counts_as_that_whole_number(tmp_path):
    config = _config(
        tmp_path, '{"limits": {"max_formulas": 50000.0, "max_reported_findings": 1.0}, "recalc": {"timeout_seconds": 30.0}}'
    )
    loaded = load_config(str(config), known_rules=audit._known_rules())
    assert [type(value) for value in (loaded["limits"]["max_formulas"], loaded["recalc"]["timeout_seconds"])] == [int, int]
    payload = _json_audit(tmp_path, _workbook(tmp_path), "--config", str(config))
    assert len(payload["findings"]) == 1  # the cap slices the findings list


def test_scope_include_sheets_naming_no_sheet_exits_4_with_the_sheet_names(tmp_path, capsys):
    config = _config(tmp_path, '{"scope": {"include_sheets": ["Modle", "Summary"]}}')
    assert _audit(tmp_path, str(_workbook(tmp_path)), "--config", str(config)) == 4
    assert capsys.readouterr().err.startswith(
        "Config error: scope.include_sheets names no sheet of this workbook: 'Modle' (did you mean 'Model'?), "
        "'Summary'. Its sheets are 'Model'"
    )


def test_scope_sheets_the_workbook_lacks_are_coverage_limitations(tmp_path):
    config = _config(tmp_path, '{"scope": {"include_sheets": ["Model", "Q4"], "exclude_sheets": ["model"]}}')
    payload = _json_audit(tmp_path, _workbook(tmp_path), "--config", str(config))
    assert payload["workbook"]["sheets_analyzed"] == 1
    notes = payload["coverage"]["limitations"]
    assert "scope.include_sheets names sheets this workbook does not have: 'Q4'; only the sheets it does have were audited." in notes
    assert (
        "scope.exclude_sheets names sheets this workbook does not have, so they excluded nothing: "
        "'model' (did you mean 'Model'?)."
    ) in notes


def test_a_yaml_sheet_name_that_reads_as_a_number_gets_a_quoting_hint(tmp_path, capsys):
    pytest.importorskip("yaml")
    config = _config(tmp_path, "scope:\n  include_sheets: [2024]\n", "config.yml")
    assert _audit(tmp_path, str(_workbook(tmp_path)), "--config", str(config)) == 4
    assert 'not [2024]; quote a sheet name that looks like a number, as in ["2024"]' in capsys.readouterr().err


def test_a_folder_given_as_the_config_or_suppression_file_exits_4(tmp_path, monkeypatch, capsys):
    workbook = str(_workbook(tmp_path))
    folder = tmp_path / "settings.json"
    folder.mkdir()
    assert _audit(tmp_path, workbook, "--config", str(folder)) == 4
    assert f"Config error: Config file {folder} is a folder, not a file" in capsys.readouterr().err
    assert audit.main([workbook, "--quiet", "--ignore", str(folder)]) == 4
    assert f"Preflight failed: suppression file {folder} is a folder, not a file" in capsys.readouterr().err
    monkeypatch.chdir(tmp_path)
    (tmp_path / ".audit-ignore").mkdir()  # the default file
    assert audit.main([workbook, "--quiet"]) == 4
    assert "Config error: Suppression file .audit-ignore is a folder, not a file" in capsys.readouterr().err


def test_an_unreadable_config_or_suppression_file_exits_4(tmp_path, monkeypatch, capsys):
    workbook = str(_workbook(tmp_path))
    config = _config(tmp_path, "{}")
    ignore = tmp_path / ".audit-ignore"
    ignore.write_text("# nothing accepted yet\n", encoding="utf-8")
    read_text = Path.read_text

    def locked(self, *args, **kwargs):
        if self.name in {config.name, ignore.name}:
            raise PermissionError(13, "Permission denied", str(self))
        return read_text(self, *args, **kwargs)

    monkeypatch.setattr(Path, "read_text", locked)
    assert _audit(tmp_path, workbook, "--config", str(config)) == 4
    assert f"Config error: Could not read config file {config} (Permission denied)" in capsys.readouterr().err
    assert audit.main([workbook, "--quiet", "--ignore", str(ignore)]) == 4
    err = capsys.readouterr().err
    assert f"Config error: Could not read suppression file {ignore} (Permission denied)" in err
    assert "Traceback" not in err


@pytest.mark.parametrize(
    "entry", [{"reason": "accepted"}, {"range": "Model!B2", "reason": "accepted"}, {"rule_id": "LIVE_ERROR", "reason": "x"}]
)
def test_a_config_suppression_that_names_no_finding_is_dropped_with_a_warning(tmp_path, entry):
    config = _config(tmp_path, json.dumps({"suppressions": [entry]}))
    payload = _json_audit(tmp_path, _workbook(tmp_path), "--config", str(config))
    assert f"Suppression ignored (needs 'rule_id' and 'range', or 'fingerprint'): {entry!r}" in payload["coverage"]["limitations"]
    assert not any(finding["suppressed"] for finding in payload["findings"])


def _pin(workbook: Path, ignore: str) -> int:
    return audit.main([str(workbook), "--quiet", "--fail-on", "None", "--ignore", ignore, "--pin-suppressions"])


def test_pinning_finds_the_suppression_file_however_its_path_is_spelled(tmp_path, monkeypatch, capsys):
    workbook = _workbook(tmp_path)
    (tmp_path / "audit").mkdir()
    ignore = tmp_path / "audit" / ".audit-ignore"
    monkeypatch.chdir(tmp_path)
    for spelling in ("./audit/.audit-ignore", ignore.as_posix()):
        ignore.write_text("BROKEN_REFERENCE Model!B2 legacy link\n", encoding="utf-8")
        _pin(workbook, spelling)
        assert "BROKEN_REFERENCE Model!B2 fingerprint:" in ignore.read_text(encoding="utf-8"), spelling
        assert "Pinned 1 line(s)" in capsys.readouterr().err


@pytest.mark.skipif(hasattr(os, "geteuid") and os.geteuid() == 0, reason="root may write a read-only file")
def test_pinning_a_read_only_suppression_file_exits_4_and_leaves_it_alone(tmp_path, capsys):
    ignore = tmp_path / ".audit-ignore"
    ignore.write_text("BROKEN_REFERENCE Model!B2 legacy link\n", encoding="utf-8")
    os.chmod(ignore, stat.S_IREAD)
    try:
        code = _pin(_workbook(tmp_path), str(ignore))
    finally:
        os.chmod(ignore, stat.S_IREAD | stat.S_IWRITE)
    assert code == 4
    err = capsys.readouterr().err
    assert f"Config error: Could not rewrite suppression file {ignore} (it is read-only)" in err
    assert ignore.read_text(encoding="utf-8") == "BROKEN_REFERENCE Model!B2 legacy link\n"
    assert not (tmp_path / ".audit-ignore.pinning").exists()


def test_a_failed_pinning_write_exits_4_and_removes_its_scratch_file(tmp_path, monkeypatch, capsys):
    ignore = tmp_path / ".audit-ignore"
    ignore.write_text("BROKEN_REFERENCE Model!B2 legacy link\n", encoding="utf-8")

    def locked(source, target):
        raise PermissionError(13, "Permission denied", str(target))

    monkeypatch.setattr(os, "replace", locked)  # as when Windows or a sync client holds the file open
    assert _pin(_workbook(tmp_path), str(ignore)) == 4
    assert f"Could not rewrite suppression file {ignore} (Permission denied)" in capsys.readouterr().err
    assert not (tmp_path / ".audit-ignore.pinning").exists()


def test_a_utf16_csv_with_a_byte_order_mark_is_audited_like_its_utf8_twin(tmp_path):
    text = "Name,Amount\r\ncafé ,10\r\nthé,20\r\n"
    utf8, utf16 = tmp_path / "utf8.csv", tmp_path / "utf16.csv"
    utf8.write_bytes(text.encode("utf-8"))
    utf16.write_bytes(codecs.BOM_UTF16_LE + text.encode("utf-16-le"))
    assert _findings(_json_audit(tmp_path, utf16)) == _findings(_json_audit(tmp_path, utf8)) != []


@pytest.mark.parametrize("text", ["Name,Amount\r\nAda ,10\r\n", "Name,Amount\r\ncafé ,10\r\n"])
def test_a_utf16_csv_without_a_byte_order_mark_exits_4(tmp_path, capsys, text):
    # ASCII-only UTF-16 is valid UTF-8; the rest falls back to Windows-1252.
    source = tmp_path / "export.csv"
    source.write_bytes(text.encode("utf-16-le"))
    assert _audit(tmp_path, str(source)) == 4
    err = capsys.readouterr().err
    assert err.startswith("Preflight failed: The CSV holds NUL characters")
    assert "save it as 'CSV UTF-8'" in err


def test_a_csv_field_over_128_kib_is_audited(tmp_path):
    source = tmp_path / "notes.csv"
    source.write_bytes(("Name,Notes\r\nAda," + "x" * 200_000 + "\r\n Bob,short\r\n").encode("utf-8"))
    previous = csv.field_size_limit(128 * 1024)  # the csv module's default, whatever ran before
    try:
        payload = _json_audit(tmp_path, source)
    finally:
        csv.field_size_limit(previous)
    assert [finding["location"] for finding in payload["findings"]] == ["CSV!R3C1"]


def test_a_csv_the_csv_module_cannot_parse_exits_4(tmp_path, monkeypatch, capsys):
    source = tmp_path / "notes.csv"
    source.write_bytes(b"Name,Notes\r\nAda,far too long for the limit\r\n")
    monkeypatch.setattr(audit, "_allow_long_csv_fields", lambda: None)
    previous = csv.field_size_limit(10)
    try:
        code = _audit(tmp_path, str(source))
    finally:
        csv.field_size_limit(previous)
    assert code == 4
    assert "Preflight failed: The CSV could not be parsed (field larger than field limit" in capsys.readouterr().err


def test_a_bug_while_loading_the_workbook_exits_5_not_4(tmp_path, monkeypatch, capsys):
    def broken(path):
        raise TypeError("load_workbook() got an unexpected keyword argument")

    monkeypatch.setattr(workbook_inventory, "load_workbook_formulas", broken)
    assert _audit(tmp_path, str(_workbook(tmp_path))) == 5
    err = capsys.readouterr().err
    assert "Internal audit error while auditing: TypeError: load_workbook()" in err
    assert "Traceback (most recent call last):" in err


def test_a_value_the_workbook_reader_rejects_exits_4(tmp_path, monkeypatch, capsys):
    # openpyxl raises TypeError for a malformed attribute; raised inside the
    # reader, that is a workbook it cannot read, not a bug in the auditor.
    import builtins
    import types

    def reject(path):
        raise TypeError("expected <class 'int'>")

    # The same function, run as if it were openpyxl's own code.
    reader = {"__name__": "openpyxl.descriptors.base", "__builtins__": builtins}
    monkeypatch.setattr(
        workbook_inventory, "load_workbook_formulas", types.FunctionType(reject.__code__, reader, "load")
    )
    assert _audit(tmp_path, str(_workbook(tmp_path))) == 4
    err = capsys.readouterr().err
    assert err.startswith("Preflight failed: The workbook could not be read (TypeError: expected <class 'int'>)")
    assert "Traceback" not in err


@pytest.mark.parametrize(
    ("source", "copy", "message"),
    [
        ("export.csv", "annotated.xlsx", "--annotated copies an .xlsx or .xlsm workbook, and "),
        ("model.xlsx", "annotated.csv", "must end in .xlsx, like the workbook it copies."),
        ("model.xlsx", "annotated.xlsm", "must end in .xlsx, like the workbook it copies."),
    ],
)
def test_annotated_needs_a_workbook_and_a_path_of_its_type(tmp_path, capsys, source, copy, message):
    path = _workbook(tmp_path) if source.endswith(".xlsx") else tmp_path / source
    if source.endswith(".csv"):
        path.write_bytes(b"Name,Amount\r\nAda,10\r\n")
    assert _audit(tmp_path, str(path), "--annotated", str(tmp_path / copy)) == 4
    err = capsys.readouterr().err
    assert err.startswith("Preflight failed: --annotated") and message in err
    assert not (tmp_path / copy).exists()


def _stdout_audit(tmp_path: Path, capsys, *args: str) -> tuple[int, str, str]:
    ignore = tmp_path / "no-suppressions"
    ignore.touch()
    # Recalculation off: a coverage limitation on every machine, so --fail-on
    # None exits 2 whether LibreOffice is installed or not.
    config = _config(tmp_path, '{"recalc": {"enabled": false}}', "stdout-config.json")
    code = audit.main(
        [str(_workbook(tmp_path)), "--fail-on", "None", "--ignore", str(ignore), "--config", str(config), *args]
    )
    captured = capsys.readouterr()
    return code, captured.out, captured.err


@pytest.mark.parametrize("extra", [(), ("--json", "findings.json")])
def test_format_without_out_prints_that_format_on_stdout(tmp_path, capsys, extra):
    extra = tuple(str(tmp_path / arg) if arg.endswith(".json") else arg for arg in extra)
    _code, out, _err = _stdout_audit(tmp_path, capsys, "--format", "sarif", *extra)
    sarif = json.loads(out)
    assert sarif["version"] == "2.1.0" and sarif["runs"][0]["results"]
    _code, out, _err = _stdout_audit(tmp_path, capsys, "--format", "html", *extra)
    assert out.lstrip().lower().startswith("<!doctype html")


def test_a_sarif_json_report_path_gets_sarif(tmp_path):
    report = tmp_path / "report.sarif.json"
    _audit(tmp_path, str(_workbook(tmp_path)), "--out", str(report), "--fail-on", "None")
    assert json.loads(report.read_text(encoding="utf-8"))["version"] == "2.1.0"


@pytest.mark.parametrize("extra", [("--summary",), ("--out", "report.md"), ("--summary", "--out", "report.md")])
def test_json_on_stdout_stays_parseable_next_to_other_output(tmp_path, capsys, extra):
    extra = tuple(str(tmp_path / arg) if arg.endswith(".md") else arg for arg in extra)
    code, out, err = _stdout_audit(tmp_path, capsys, "--json", "-", *extra)
    assert code == 2
    assert json.loads(out)["findings"]
    if "--summary" in extra:
        assert "fail_on    : None" in err
    else:
        assert "Audit complete: 1 Critical, 1 High" in err


def test_json_on_stdout_and_a_report_format_on_stdout_exit_4(tmp_path, capsys):
    code, out, err = _stdout_audit(tmp_path, capsys, "--json", "-", "--format", "sarif")
    assert (code, out) == (4, "")
    assert "--json - and --format sarif would both write to stdout" in err


@pytest.mark.parametrize("value", ["0", "-5", "soon"])
def test_a_recalc_timeout_below_one_second_exits_4(tmp_path, capsys, value):
    with pytest.raises(SystemExit) as raised:
        audit.main([str(_workbook(tmp_path)), f"--recalc-timeout={value}"])
    assert raised.value.code == 4
    assert "argument --recalc-timeout: must be" in capsys.readouterr().err


def test_quiet_still_says_on_stderr_why_it_exits_1_or_2(tmp_path, capsys):
    workbook = str(_workbook(tmp_path))  # one Critical and one High finding
    assert _audit(tmp_path, workbook, "--fail-on", "High") == 1
    captured = capsys.readouterr()
    assert (captured.out, captured.err) == ("", "Exit 1: 2 finding(s) at or above --fail-on High.\n")
    config = _config(tmp_path, '{"recalc": {"enabled": false}, "checks": {"LIVE_ERROR": "off"}}')
    for flag, where in (("--strict", "under --strict;"), ("--fail-on=None", "under --fail-on None;")):
        assert _audit(tmp_path, workbook, "--config", str(config), flag) == 2
        captured = capsys.readouterr()
        assert captured.out == ""
        assert captured.err.startswith("Exit 2: ") and where in captured.err, captured.err
