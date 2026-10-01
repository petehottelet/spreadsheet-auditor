"""IFERROR/IFNA wrappers are reported only when they pass off a failure as data."""

from __future__ import annotations

import json

import pytest
from openpyxl import Workbook

from spreadsheet_auditor import audit
from spreadsheet_auditor.error_masks import masks


@pytest.mark.parametrize(
    "formula, sources, fallback",
    [
        # An average of nothing shown as an average of 0.
        ('=IFERROR(AVERAGEIFS(T[P],T[G],">=75%"),0)', ("AVERAGEIFS",), "0"),
        # A missing record read as an age of 0.
        ("=IFERROR(INDEX($S$2:$S$124,MATCH($A9,$Q$2:$Q124,0)),0)", ("INDEX", "MATCH"), "0"),
        # IFNA catches only the #N/A of the MATCH; INDEX and the division fail with other errors.
        ("=_xlfn.IFNA(INDEX(S!C3:G20,MATCH(A14,S!A3:A20,0)),0)", ("MATCH",), "0"),
        ("=_xlfn.IFNA(INDEX(C:C,MATCH(A1,B:B,0))/D1,0)", ("MATCH",), "0"),
        # A bad date or a blank frequency becomes a zero payment.
        ('=IFERROR((MOD(DATEDIF(L$4+1,EOMONTH($D10,0)+1,"m"),12/$C10)=0)*$B10,0)', ("MOD", "DATEDIF", "arithmetic on a cell", "EOMONTH", "a division"), "0"),
        # A division by zero written as the text "0", which totals skip.
        ('=IFERROR(((C12*1.3)+(D12*1))/((B12)),"0")', ("arithmetic on a cell", "a division"), '"0"'),
        ("=IFERROR(B2*C2,)", ("arithmetic on a cell",), "0"),
        ("=IFERROR(B2/C2,D2)", ("a division",), "D2"),
        # Falling back to an input is deliberate only for a text conversion.
        ("=IFERROR(A1/B1,A1)", ("a division",), "A1"),
        # A negated or parenthesised cell fails on text like any arithmetic.
        ("=IFERROR(-A1,0)", ("arithmetic on a cell",), "0"),
        ("=IFERROR((A1)*(B1),0)", ("arithmetic on a cell",), "0"),
        # The tokenizer glues the range start to the call in A1:OFFSET(...).
        ("=IFERROR(A1:OFFSET(A1,2,0),0)", ("OFFSET",), "0"),
        ("=IFERROR(A1/B1,{0,0})", ("a division",), "{0,0}"),
        # An inner IFNA lets a division error through to the outer IFERROR.
        ('=IFERROR(_xlfn.IFNA(A1/B1,""),0)', ("a division",), "0"),
        # An inner IFERROR that turns the miss into #N/A for the IFNA around it.
        ("=_xlfn.IFNA(IFERROR(VLOOKUP(A1,B:C,2,0),NA()),0)", ("NA()",), "0"),
        # ISEVEN and ISODD pass an error on instead of absorbing it.
        ("=IFERROR(ISEVEN(VALUE(A1)),FALSE)", ("VALUE",), "FALSE"),
    ],
)
def test_failures_passed_off_as_data_are_reported(formula, sources, fallback):
    [mask] = masks(formula)
    assert mask.sources == sources
    assert mask.fallback == fallback


@pytest.mark.parametrize(
    "formula, source",
    [
        ('=IFERROR(_xlfn.TEXTBEFORE(A1,"-"),0)', "TEXTBEFORE"),
        ('=_xlfn.IFNA(_xlfn.TEXTAFTER(A1,"-"),0)', "TEXTAFTER"),
        ("=IFERROR(SLOPE(B1:B5,A1:A5),0)", "SLOPE"),
        ("=IFERROR(INTERCEPT(B1:B5,A1:A5),0)", "INTERCEPT"),
        ("=IFERROR(CORREL(B1:B5,A1:A5),0)", "CORREL"),
        ("=IFERROR(PEARSON(B1:B5,A1:A5),0)", "PEARSON"),
        ("=IFERROR(RSQ(B1:B5,A1:A5),0)", "RSQ"),
        ("=IFERROR(FORECAST(6,B1:B5,A1:A5),0)", "FORECAST"),
        ("=IFERROR(_xlfn.FORECAST.LINEAR(6,B1:B5,A1:A5),0)", "FORECAST.LINEAR"),
        ("=IFERROR(STDEVP(A1:A5),0)", "STDEVP"),
        ("=IFERROR(_xlfn.STDEV.P(A1:A5),0)", "STDEV.P"),
        ("=IFERROR(VARP(A1:A5),0)", "VARP"),
        ("=IFERROR(_xlfn.VAR.P(A1:A5),0)", "VAR.P"),
        ("=IFERROR(SUBTOTAL(1,A1:A5),0)", "SUBTOTAL"),
        ("=IFERROR(SUBTOTAL(101,A1:A5),0)", "SUBTOTAL"),
        ("=IFERROR(_xlfn.AGGREGATE(1,6,A1:A5),0)", "AGGREGATE"),
    ],
)
def test_functions_that_can_fail(formula, source):
    [mask] = masks(formula)
    assert mask.sources == (source,)


@pytest.mark.parametrize("text", ["0", "0%", "$0", "(1)", "1,000", "-1.5", "1E3", " 0 "])
def test_numbers_written_as_text_read_as_data(text):
    [mask] = masks(f'=IFERROR(A1/B1,"{text}")')
    assert mask.fallback == f'"{text}"'


@pytest.mark.parametrize("text", ["NaN", "inf", "-inf", "Infinity", "1_000", "n/a"])
def test_words_python_reads_as_numbers_are_messages(text):
    assert masks(f'=IFERROR(A1/B1,"{text}")') == []


@pytest.mark.parametrize(
    "formula",
    [
        '=IFERROR(VLOOKUP(C3,DATA!$A:$E,2,0),"")',  # a lookup miss shown as blank
        '=IFERROR(VLOOKUP(I19,$I$6:$J$8,2,FALSE),"Select freq")',  # an explicit message
        '=IFERROR(VLOOKUP(A10,S1!A2:D10,3,FALSE),IFERROR(VLOOKUP(A10,S2!A2:D10,3,FALSE),""))',  # a chain of lists
        "=IFERROR(VLOOKUP(A1,B:C,2,FALSE),NA())",  # turns any error into #N/A
        '=IFERROR(IF(COUNTIFS(T[C],"<1")>=10,AVERAGEIFS(T[P],T[C],"<1"),0),0)',  # the IF already guards it
        '=IFERROR(COUNTIFS(T[C],"<1",T[F],">3"),0)',  # COUNTIFS cannot error
        "=IFERROR(MID($B5,COLUMNS($D7:P7),1),\"\")",  # MID cannot error here
        '=IFERROR(INDEX(R!M:M,SMALL(IF(R!$A$2:$A$70=F!$B$3,ROW(R!$A$2:$A$70)),ROWS($1:1))),"")',  # end of list
        '=IFERROR(LOOKUP(1,0/COUNTIFS(O5:AA5,O$1),{1,0,-1}),"")',
        "=IFERROR(SUM(A1:A5),0)",
        # A sum, count, max or min of nothing is 0, not an error.
        "=IFERROR(SUBTOTAL(109,A1:A5),0)",
        "=IFERROR(_xlfn.AGGREGATE(9,6,A1:A5),0)",
    ],
)
def test_deliberate_wrappers_are_left_alone(formula):
    assert masks(formula) == []


@pytest.mark.parametrize(
    "formula",
    [
        "=_xlfn.IFNA(A1/B1,0)",
        "=_xlfn.IFNA(AVERAGE(A1:A5),0)",
        "=_xlfn.IFNA(VALUE(A1),0)",
    ],
)
def test_ifna_ignores_errors_it_cannot_catch(formula):
    assert masks(formula) == []


@pytest.mark.parametrize(
    "formula",
    [
        "=IFERROR(ISNUMBER(MATCH(A1,B:B,0)),FALSE)",
        '=IFERROR(SUMPRODUCT(--ISNUMBER(SEARCH("x",A1:A5))),0)',
        '=IFERROR(IFERROR(VLOOKUP(A1,B:C,2,0),"")&"",0)',
        "=IFERROR(ERROR.TYPE(A1/B1),0)",
    ],
)
def test_errors_absorbed_before_the_wrapper_do_not_count(formula):
    assert masks(formula) == []


@pytest.mark.parametrize(
    "formula",
    [
        '=IFERROR(LEFT(A1,FIND(" ",A1)-1),A1)',
        '=IFERROR(MID(A1,SEARCH("-",A1)+1,99),A1)',
        "=IFERROR(VALUE(A1),A1)",
        "=IFERROR(DATEVALUE($A1),A1)",
    ],
)
def test_a_conversion_that_falls_back_to_its_input_is_deliberate(formula):
    assert masks(formula) == []


@pytest.mark.parametrize(
    "fallback",
    [
        'INDIRECT("C1")',
        "OFFSET(C1,1,0)",
        "_xlfn.XMATCH(A1,C:C)",
        "_xlfn._xlws.FILTER(C:C,D:D=A1)",
        '"Missing: "&C1',
        'CONCATENATE("Missing: ",C1)',
    ],
)
def test_a_second_lookup_or_a_message_fallback_is_left_alone(fallback):
    assert masks(f"=IFERROR(VLOOKUP(A1,B:C,2,0),{fallback})") == []


def test_google_sheets_placeholders_are_always_reported():
    [mask] = masks('=IFERROR(__xludf.DUMMYFUNCTION("""COMPUTED_VALUE"""),"60*90")')
    assert mask.sources == ("a Google Sheets placeholder",)
    assert mask.fallback == '"60*90"'


def test_each_wrapper_in_a_formula_is_judged_on_its_own():
    formula = '=IFERROR(A1/B1,0)+IFERROR(VLOOKUP(C1,D:E,2,FALSE),"")'
    assert [(m.function, m.sources) for m in masks(formula)] == [("IFERROR", ("a division",))]


def test_placeholders_get_their_own_suggested_fix(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    ws["A1"] = '=IFERROR(__xludf.DUMMYFUNCTION("""COMPUTED_VALUE"""),"60*90")'
    ws["B1"] = "=IFERROR(C1/D1,0)"
    path = tmp_path / "gsheets.xlsx"
    wb.save(path)
    out = tmp_path / "findings.json"
    ignore = tmp_path / "no-suppressions"
    ignore.touch()
    audit.main([str(path), "--json", str(out), "--quiet", "--fail-on", "None", "--ignore", str(ignore)])
    fixes = {
        f["location"]: f["suggested_fix"]
        for f in json.loads(out.read_text(encoding="utf-8"))["findings"]
        if f["rule_id"] == "IFERROR_MASK"
    }
    assert fixes["S!A1"].startswith("Replace the frozen placeholder with a formula Excel evaluates")
    assert fixes["S!B1"].startswith("Return a visible marker for the failure")
