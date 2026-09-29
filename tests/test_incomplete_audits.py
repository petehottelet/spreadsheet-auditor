"""An audit that skipped part of the workbook must never pass as a complete one."""

from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import Workbook

from spreadsheet_auditor import audit
from spreadsheet_auditor.checks import checks
from spreadsheet_auditor.checks.formula_integrity import LiveErrorCheck
from spreadsheet_auditor.report import render_html, render_markdown
from spreadsheet_auditor.sarif import render_sarif

ROOT = Path(__file__).resolve().parents[1]


def _workbook_with_ref_error(tmp_path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Model"
    for row in range(1, 6):
        ws.cell(row, 1, f"Line {row}")
        ws.cell(row, 2, row * 100)
    ws["C5"] = "=SUM(#REF!)"
    path = tmp_path / "model.xlsx"
    wb.save(path)
    return path


def _crash(self, ctx):
    raise RuntimeError("simulated bug")


def test_a_crashed_check_fails_the_gate_instead_of_passing(tmp_path, monkeypatch, capsys):
    workbook = _workbook_with_ref_error(tmp_path)
    out = tmp_path / "findings.json"
    # Without the crash the Critical LIVE_ERROR at C5 fails the default gate.
    assert audit.main([str(workbook), "--json", str(out), "--ignore", str(tmp_path / "none")]) == 1

    # With it, that finding is lost; the run used to exit 0 as if the workbook were clean.
    monkeypatch.setattr(LiveErrorCheck, "run", _crash)
    code = audit.main([str(workbook), "--json", str(out), "--ignore", str(tmp_path / "none")])
    assert code == 6
    payload = json.loads(out.read_text(encoding="utf-8"))
    assert payload["coverage"]["complete"] is False
    [entry] = payload["coverage"]["incomplete"]
    assert entry["reason"] == "check_failed"
    assert entry["rules"] == ["LIVE_ERROR"]
    assert "simulated bug" in entry["message"]
    assert entry["message"] in payload["coverage"]["limitations"]
    stderr = capsys.readouterr().err
    assert "Audit incomplete: Check 'live_errors' raised an exception" in stderr


def test_incomplete_outranks_fail_on_none(tmp_path, monkeypatch):
    workbook = _workbook_with_ref_error(tmp_path)
    monkeypatch.setattr(LiveErrorCheck, "run", _crash)
    code = audit.main([str(workbook), "--quiet", "--fail-on", "None", "--ignore", str(tmp_path / "none")])
    assert code == 6


def test_usage_errors_do_not_share_the_limitations_exit_code(capsys):
    # argparse exits 2, which a CI script accepting "completed with
    # limitations" would read as a finished audit.
    with pytest.raises(SystemExit) as raised:
        audit.main(["model.xlsx", "--fail-on", "Severe"])
    assert raised.value.code == 4
    assert "invalid choice" in capsys.readouterr().err
    with pytest.raises(SystemExit) as raised:
        audit.main([])
    assert raised.value.code == 4


def test_checks_that_are_right_every_time_run_first():
    # When the time budget runs out, later checks are skipped; live errors and
    # broken references must not be the ones lost.
    names = [cls.__name__ for cls in checks()]
    assert names[:2] == ["LiveErrorCheck", "ReferenceIntegrityCheck"]
    assert names.index("DataHygieneCheck") > names.index("TotalMismatchCheck")


def _incomplete_payload(tmp_path, monkeypatch) -> dict:
    workbook = _workbook_with_ref_error(tmp_path)
    out = tmp_path / "findings.json"
    monkeypatch.setattr(LiveErrorCheck, "run", _crash)
    audit.main([str(workbook), "--json", str(out), "--quiet", "--ignore", str(tmp_path / "none")])
    return json.loads(out.read_text(encoding="utf-8"))


def test_sarif_reports_the_run_as_unsuccessful_with_the_reason(tmp_path, monkeypatch):
    sarif = render_sarif(_incomplete_payload(tmp_path, monkeypatch))
    invocation = sarif["runs"][0]["invocations"][0]
    assert invocation["executionSuccessful"] is False
    errors = [n for n in invocation["toolExecutionNotifications"] if n["level"] == "error"]
    assert [n["descriptor"]["id"] for n in errors] == ["check_failed"]
    jsonschema = pytest.importorskip("jsonschema")
    schema = json.loads((ROOT / "schemas" / "sarif-2.1.0.schema.json").read_text(encoding="utf-8"))
    jsonschema.validate(sarif, schema, format_checker=jsonschema.FormatChecker())


def test_payload_validates_and_reports_say_the_audit_is_partial(tmp_path, monkeypatch):
    payload = _incomplete_payload(tmp_path, monkeypatch)
    jsonschema = pytest.importorskip("jsonschema")
    schema = json.loads((ROOT / "schemas" / "findings.schema.json").read_text(encoding="utf-8"))
    jsonschema.validate(payload, schema)
    markdown = render_markdown(payload)
    assert markdown.splitlines()[2].startswith("> **Incomplete audit.**")
    assert "- Audit status: incomplete" in markdown
    assert "Incomplete audit." in render_html(payload)
    summary = audit._summary_lines(payload, "Critical")
    assert summary[1].startswith("audit      : INCOMPLETE")


def test_complete_audit_reports_success(tmp_path):
    workbook = _workbook_with_ref_error(tmp_path)
    out = tmp_path / "findings.json"
    audit.main([str(workbook), "--json", str(out), "--quiet", "--ignore", str(tmp_path / "none")])
    payload = json.loads(out.read_text(encoding="utf-8"))
    assert payload["coverage"]["complete"] is True and payload["coverage"]["incomplete"] == []
    assert render_sarif(payload)["runs"][0]["invocations"][0]["executionSuccessful"] is True
    assert "Incomplete audit" not in render_markdown(payload)
