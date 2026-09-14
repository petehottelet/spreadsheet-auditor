"""The real-world corpus harness: statistics, sampling, an end-to-end run on a
tiny synthetic corpus, and report rendering."""

from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "benchmarks" / "corpus"))

import corpuslib  # noqa: E402


def test_wilson_interval_behaves():
    low, high = corpuslib.wilson_interval(0, 0)
    assert (low, high) == (0.0, 1.0)
    low, high = corpuslib.wilson_interval(18, 20)
    assert 0.68 < low < 0.72 and 0.96 < high < 0.99
    low, high = corpuslib.wilson_interval(20, 20)
    assert low > 0.8 and high == 1.0


def test_stratified_sample_is_seeded_and_caps_per_workbook():
    items = []
    for workbook in ("a.xlsx", "b.xlsx", "c.xlsx"):
        for row in range(1, 11):
            items.append({"workbook": workbook, "sha256": "x" * 64, "finding": {"rule_id": "R1", "location": f"S!A{row}"}})
    for row in range(1, 4):
        items.append({"workbook": "d.xlsx", "sha256": "y" * 64, "finding": {"rule_id": "R2", "location": f"S!B{row}"}})
    first = corpuslib.stratified_sample(items, per_rule=5, max_per_workbook=2, seed=3)
    second = corpuslib.stratified_sample(items, per_rule=5, max_per_workbook=2, seed=3)
    assert [i["finding"]["location"] for i in first] == [i["finding"]["location"] for i in second]
    r1 = [i for i in first if i["finding"]["rule_id"] == "R1"]
    assert len(r1) == 5
    assert max(sum(1 for i in r1 if i["workbook"] == w) for w in ("a.xlsx", "b.xlsx", "c.xlsx")) <= 2
    assert len([i for i in first if i["finding"]["rule_id"] == "R2"]) == 2  # capped by max_per_workbook


def _tiny_corpus(root: Path) -> Path:
    data = root / "files"
    data.mkdir(parents=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    ws.append(["Item", "Amount"])
    for row, value in enumerate((10, 20, 30, 40), start=2):
        ws.cell(row=row, column=1, value=f"Item {row}")
        ws.cell(row=row, column=2, value=value)
    ws["A6"] = "Total"
    ws["B6"] = "=SUM(B2:B4)"  # stops short of B5
    wb.save(data / "one_input.xlsx")
    clean = Workbook()
    cs = clean.active
    cs.title = "S"
    cs["A1"] = "x"
    cs["B1"] = 1
    clean.save(data / "two_input.xlsx")
    (data / "skip_answer.xlsx").write_bytes((data / "two_input.xlsx").read_bytes())
    return data


def test_run_sample_and_report_end_to_end(tmp_path):
    data = _tiny_corpus(tmp_path)
    out = tmp_path / "results"
    summary = corpuslib.run_corpus("spreadsheetbench", data, out, workers=1, timeout=120, include_glob="*_input.*")
    assert summary["workbooks"] == 2
    assert summary["ok"] == 2 and summary["errors"] == 0 and summary["timeouts"] == 0
    assert summary["rules"]["RANGE_EXCLUSION"]["findings"] == 1
    assert (out / "run_summary.json").exists()

    items = corpuslib.load_findings(out)
    assert {i["finding"]["rule_id"] for i in items} == {"RANGE_EXCLUSION"}
    cache = corpuslib.WorkbookCache()
    context = corpuslib.build_context(data / items[0]["workbook"], items[0]["finding"], cache)
    assert context["cell"] == "B6"
    assert context["first_range"]["ref"] == "B2:B4"
    assert context["first_range"]["beyond"]["below"].startswith("B5: 40")

    key = corpuslib.finding_key(items[0]["sha256"], "RANGE_EXCLUSION", "S!B6")
    sample = {
        "source": "spreadsheetbench",
        "generated": "now",
        "results_dir": str(out),
        "seed": 7,
        "per_rule": 20,
        "max_per_workbook": 2,
        "total_findings": 1,
        "items": [{"key": key, "workbook": items[0]["workbook"], "finding": items[0]["finding"], "context": context}],
    }
    labels = {"labeler": "test", "protocol": "unit test", "rule_notes": {"RANGE_EXCLUSION": "n/a"}, "labels": {key: {"label": "TP", "reason": "B5 is a number between the range and the total"}}}
    rows = corpuslib.precision_rows(sample, labels, summary)
    row = next(r for r in rows if r["rule"] == "RANGE_EXCLUSION")
    assert (row["tp"], row["fp"], row["unsure"], row["unlabeled"]) == (1, 0, 0, 0)
    assert row["precision"] == 1.0
    report = corpuslib.render_report(corpuslib.load_sources()["spreadsheetbench"], summary, sample, labels, rows)
    assert "| RANGE_EXCLUSION | 1 | 1 |" in report
    assert "overall precision **100%**" in report
    cards = corpuslib.render_cards(sample)
    assert key in cards and "=SUM(B2:B4)" in cards


def test_run_records_crashes_instead_of_dying(tmp_path):
    data = tmp_path / "files"
    data.mkdir()
    (data / "broken_input.xlsx").write_bytes(b"not a workbook")
    summary = corpuslib.run_corpus("spreadsheetbench", data, tmp_path / "results", workers=1, timeout=60)
    assert summary["workbooks"] == 1 and summary["errors"] == 1
    record = next(iter(summary["records"].values()))
    assert record["exit_code"] == 4
    assert "not a valid" in record["stderr"].lower() or "preflight" in record["stderr"].lower()


def test_assign_keys_disambiguates_findings_that_share_a_location():
    sha = "a" * 64
    items = [
        {"sha256": sha, "finding": {"rule_id": "HIDDEN_STRUCTURE_IN_TOTAL", "location": "S!B1", "evidence": ["Hidden row 3 on S feeds 1 visible formula(s): S!B1."]}},
        {"sha256": sha, "finding": {"rule_id": "HIDDEN_STRUCTURE_IN_TOTAL", "location": "S!B1", "evidence": ["Hidden row 7 on S feeds 1 visible formula(s): S!B1."]}},
        {"sha256": sha, "finding": {"rule_id": "LITERAL_CONSTANT", "location": "S!B1", "evidence": ["Non-trivial numeric literal(s) found: 9."]}},
    ]
    corpuslib.assign_keys(items)
    keys = [item["key"] for item in items]
    assert len(set(keys)) == 3
    assert keys[2] == corpuslib.finding_key(sha, "LITERAL_CONSTANT", "S!B1")
    base = corpuslib.finding_key(sha, "HIDDEN_STRUCTURE_IN_TOTAL", "S!B1")
    assert keys[0].startswith(base + "|") and keys[1].startswith(base + "|")
    corpuslib.assign_keys(items)
    assert [item["key"] for item in items] == keys
