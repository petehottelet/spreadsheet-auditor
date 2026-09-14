"""Recalculation profile, static error inference, original-formula loading,
LF output, the max_cells cap, deprecated config, and the annotated-copy warning."""

from __future__ import annotations

import argparse
import json
import subprocess
import sys
import zipfile
from pathlib import Path

from openpyxl import Workbook

from spreadsheet_auditor import audit as audit_module
from spreadsheet_auditor.finding import Finding
from spreadsheet_auditor.formula_parser import parse_formula
from spreadsheet_auditor.recalc import PROFILE_SETTINGS, seed_profile
from spreadsheet_auditor.report import build_payload, write_html, write_json, write_markdown
from spreadsheet_auditor.sarif import write_sarif


def _run(path: Path, *extra: str) -> tuple[dict, subprocess.CompletedProcess]:
    result = subprocess.run(
        [sys.executable, "-m", "spreadsheet_auditor", str(path), "--json", "-", "--fail-on", "None", *extra],
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
    )
    assert result.returncode in (0, 1, 2), result.stderr
    return json.loads(result.stdout), result


def test_lowercase_error_literals_still_parse():
    parsed = parse_formula("=SUM(#ref!)")
    assert parsed.parse_error is None
    assert parsed.deleted_reference


def test_error_literal_outside_a_handler_is_a_live_error_without_recalc(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    ws["A1"] = "=SUM(#REF!)"
    ws["A2"] = "=IFERROR(#REF!,0)"
    path = tmp_path / "err.xlsx"
    wb.save(path)
    payload, _ = _run(path)
    live = [f for f in payload["findings"] if f["rule_id"] == "LIVE_ERROR"]
    assert [f["location"] for f in live] == ["S!A1"]
    assert "whatever its inputs" in live[0]["evidence"][0]


def test_static_checks_read_the_original_not_the_recalculated_copy(tmp_path, monkeypatch):
    original = Workbook()
    ws = original.active
    ws.title = "S"
    ws["A1"] = "Label"
    ws["B1"] = 10
    ws["B2"] = 20
    ws["B3"] = 30
    ws["A4"] = "Total"
    ws["B4"] = "=SUM(B1:B2)"  # stops short of B3
    source = tmp_path / "orig.xlsx"
    original.save(source)

    # A converter that re-serializes the formula differently must not change
    # what the static checks see.
    mangled = Workbook()
    ms = mangled.active
    ms.title = "S"
    for coord, value in (("A1", "Label"), ("B1", 10), ("B2", 20), ("B3", 30), ("A4", "Total"), ("B4", "=B1+B2")):
        ms[coord] = value
    copy = tmp_path / "recalculated.xlsx"
    mangled.save(copy)
    monkeypatch.setattr(
        audit_module,
        "recalc_if_available",
        lambda *args, **kwargs: {"status": "completed", "path": str(copy), "limitations": []},
    )
    args = argparse.Namespace(
        workbook=str(source),
        config=None,
        ignore=str(tmp_path / "no-ignore"),
        fail_on="None",
        strict=False,
        recalc_timeout=None,
    )
    payload, _ = audit_module.audit_workbook(args)
    exclusions = [f for f in payload["findings"] if f["rule_id"] == "RANGE_EXCLUSION"]
    assert [f["formula"] for f in exclusions] == ["=SUM(B1:B2)"]
    assert payload["workbook"]["recalc_status"] == "completed"


def test_generated_files_use_lf_newlines(tmp_path):
    payload = build_payload(
        "0.1.0",
        {"path": "x.xlsx", "sha256": "abc", "sheets_analyzed": 1, "formulas_scanned": 1, "recalc_status": "completed"},
        {
            "macros_present": False,
            "macros_executed": False,
            "external_links_present": False,
            "unsupported_features": [],
            "limitations": [],
        },
        [Finding("LIVE_ERROR", "Critical", "Defect", "DET", "S!A1", "t", ["e"], "f")],
    )
    for writer, name in (
        (write_json, "f.json"),
        (write_markdown, "r.md"),
        (write_html, "r.html"),
        (write_sarif, "r.sarif"),
    ):
        target = tmp_path / name
        writer(payload, target)
        assert b"\r\n" not in target.read_bytes(), name


def test_max_cells_cap_skips_grid_checks_and_deprecated_limit_is_noted(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    ws.append(["Line", "Jan", "Feb", "Mar", "Total"])
    ws.append(["A", 1, 2, 3, "=SUM(B2:D2)"])
    ws.append(["B", 4, 5, 6, "=SUM(B3:D3)"])
    ws.append(["C", 7, 8, 9, 999])
    ws.append(["D", 1, 1, 1, "=SUM(B5:D5)"])
    path = tmp_path / "cap.xlsx"
    wb.save(path)
    config = tmp_path / "config.json"
    config.write_text(
        json.dumps({"limits": {"max_cells": 5, "max_range_expansion_cells": 10}}), encoding="utf-8"
    )

    capped, _ = _run(path, "--config", str(config))
    assert "HARDCODE_IN_FORMULA_BLOCK" not in {f["rule_id"] for f in capped["findings"]}
    assert capped["coverage"]["truncated"]["cells"] is True
    notes = capped["coverage"]["limitations"]
    assert any("cell-grid checks" in note for note in notes)
    assert any("max_range_expansion_cells is deprecated" in note for note in notes)

    uncapped, _ = _run(path)
    assert "HARDCODE_IN_FORMULA_BLOCK" in {f["rule_id"] for f in uncapped["findings"]}


def test_annotated_copy_warns_when_drawings_would_be_dropped(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    ws["A1"] = 1
    ws["B1"] = "=A1*2"
    path = tmp_path / "drawn.xlsx"
    wb.save(path)
    with zipfile.ZipFile(path, "a") as archive:
        archive.writestr("xl/media/image1.png", b"\x89PNG stub")

    payload, result = _run(path, "--annotated", str(tmp_path / "annotated.xlsx"))
    assert any(note.startswith("The annotated copy") for note in payload["coverage"]["limitations"])
    assert "Warning: The annotated copy" in result.stderr
    assert (tmp_path / "annotated.xlsx").exists()

    quiet, _ = _run(path)
    assert not any(note.startswith("The annotated copy") for note in quiet["coverage"]["limitations"])


def test_recalc_profile_forces_recalculation_and_disables_macros(tmp_path):
    settings = seed_profile(tmp_path / "profile")
    assert settings == tmp_path / "profile" / "user" / "registrymodifications.xcu"
    text = settings.read_text(encoding="utf-8")
    assert PROFILE_SETTINGS.startswith("<?xml")
    assert 'oor:name="OOXMLRecalcMode"' in text
    assert 'oor:name="MacroSecurityLevel"' in text
    assert text.count("<value>0</value>") == 2
    assert "<value>3</value>" in text
