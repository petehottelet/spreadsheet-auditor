"""Precision: correct workbooks stay quiet, seeded defects still fire, and every
finding on the seeded corpora maps to a catalogued defect."""

from __future__ import annotations

import datetime as dt
import json
import subprocess
import sys
from collections import defaultdict
from pathlib import Path

import pytest
from openpyxl import Workbook

from spreadsheet_auditor.checks import CheckContext
from spreadsheet_auditor.checks.formula_integrity import LiveErrorCheck
from spreadsheet_auditor.reconcile import double_counted_cells
from spreadsheet_auditor.workbook_inventory import formula_cells

ROOT = Path(__file__).resolve().parents[1]
CATALOG = ROOT / "benchmarks" / "seeded_defects.json"
VALUE_DEPENDENT = {"TOTAL_MISMATCH", "CROSS_FOOT_FAILURE", "LIVE_ERROR"}


def _audit(path: Path) -> dict:
    result = subprocess.run(
        [sys.executable, "-m", "spreadsheet_auditor", str(path), "--json", "-", "--fail-on", "None"],
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
    )
    assert result.returncode in (0, 1, 2), result.stderr
    return json.loads(result.stdout)


def _rules(payload: dict) -> dict[str, list[str]]:
    out: dict[str, list[str]] = defaultdict(list)
    for finding in payload["findings"]:
        out[finding["rule_id"]].append(finding["location"])
    return out


def _severe(payload: dict) -> list[tuple[str, str]]:
    return [
        (f["rule_id"], f["location"]) for f in payload["findings"] if f["severity"] in ("Critical", "High")
    ]


def _save(wb: Workbook, tmp_path: Path, name: str) -> Path:
    path = tmp_path / name
    wb.save(path)
    return path


# --- correct workbooks stay quiet -------------------------------------------


def test_correct_budget_with_subtotals_and_total_column_is_quiet(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Budget"
    ws.merge_cells("A1:E1")
    ws["A1"] = "Operating Budget"
    ws.append(["Line item", "Jan", "Feb", "Mar", "Q1 total"])
    ws.append(["Subscriptions", 1200, 1320, 1450, "=SUM(B3:D3)"])
    ws.append(["Services", 800, 820, 900, "=SUM(B4:D4)"])
    ws.append(["Revenue subtotal", "=SUM(B3:B4)", "=SUM(C3:C4)", "=SUM(D3:D4)", "=SUM(B5:D5)"])
    ws.append(["Hosting", 200, 210, 220, "=SUM(B6:D6)"])
    ws.append(["Staff", 600, 620, 640, "=SUM(B7:D7)"])
    ws.append(["Cost subtotal", "=SUM(B6:B7)", "=SUM(C6:C7)", "=SUM(D6:D7)", "=SUM(B8:D8)"])
    ws.append(["Net", "=B5-B8", "=C5-C8", "=D5-D8", "=SUM(B9:D9)"])
    ws.append(["Avg monthly subscriptions", "=AVERAGE(B3:D3)"])
    ws.append(["Other income", 30])
    ws.append(["Revenue incl. other (subtotal + other)", "=B5+B11"])
    payload = _audit(_save(wb, tmp_path, "budget.xlsx"))
    assert payload["findings"] == []


def test_projection_with_absolute_anchor_and_date_headers_is_quiet(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Proj"
    ws["A1"] = "Growth"
    ws["B1"] = 0.05
    ws["A3"] = "Period"
    for offset, col in enumerate("BCDEFG"):
        ws[f"{col}3"] = dt.date(2024, offset + 1, 1)
    ws["H3"] = "Total"
    ws["A4"] = "Revenue"
    ws["B4"] = 1000
    ws["A5"] = "Cost"
    ws["B5"] = 400
    for prev, col in zip("BCDEF", "CDEFG"):
        ws[f"{col}4"] = f"={prev}4*(1+$B$1)"
        ws[f"{col}5"] = f"={prev}5*(1+$B$1)"
    ws["A6"] = "Margin"
    for col in "BCDEFG":
        ws[f"{col}6"] = f"={col}4-{col}5"
    for row in (4, 5, 6):
        ws[f"H{row}"] = f"=SUM(B{row}:G{row})"
    payload = _audit(_save(wb, tmp_path, "proj.xlsx"))
    assert payload["findings"] == []


def test_lookup_sheet_with_iferror_reports_only_review_items(tmp_path):
    wb = Workbook()
    rates = wb.active
    rates.title = "Rates"
    rates.append(["Code", "Rate"])
    rates.append(["A", 0.1])
    rates.append(["B", 0.2])
    rates.append(["C", 0.3])
    calc = wb.create_sheet("Calc")
    calc.append(["Code", "Amount", "Rate", "Tax"])
    for row, (code, amount) in enumerate([("A", 100), ("B", 200), ("C", 300)], start=2):
        calc.append(
            [code, amount, f"=IFERROR(VLOOKUP(A{row},Rates!$A$2:$B$4,2,FALSE),0)", f"=B{row}*C{row}"]
        )
    calc["A6"] = "Total"
    calc["D6"] = "=SUM(D2:D4)"
    payload = _audit(_save(wb, tmp_path, "lookup.xlsx"))
    assert set(_rules(payload)) <= {"IFERROR_MASK"}
    assert _severe(payload) == []


def test_top_placed_total_and_quarterly_sums_are_quiet(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Months"
    ws["A1"] = "Year total"
    ws["B1"] = "=SUM(B3:B14)"
    ws["A2"] = "Month"
    ws["B2"] = "Sales"
    for offset in range(12):
        ws.cell(row=3 + offset, column=1, value=f"M{offset + 1}")
        ws.cell(row=3 + offset, column=2, value=100 + offset)
    ws["D2"] = "Quarter"
    ws["E2"] = "Sales"
    for offset, quarter in enumerate(("Q1", "Q2", "Q3", "Q4")):
        first = 3 + 3 * offset
        ws.cell(row=3 + offset, column=4, value=quarter)
        ws.cell(row=3 + offset, column=5, value=f"=SUM(B{first}:B{first + 2})")
    payload = _audit(_save(wb, tmp_path, "quarters.xlsx"))
    assert payload["findings"] == []


def test_balance_sheet_style_totals_are_quiet(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "BS"
    ws.append(["Cash", 100])
    ws.append(["Receivables", 50])
    ws.append(["Total current assets", "=SUM(B1:B2)"])
    ws.append(["PP&E", 300])
    ws.append(["Total assets", "=B3+B4"])
    ws.append(["Payables", 80])
    ws.append(["Equity", 370])
    ws.append(["Total liabilities and equity", "=SUM(B6:B7)"])
    ws.append(["Check", "=B5-B8"])
    payload = _audit(_save(wb, tmp_path, "bs.xlsx"))
    assert payload["findings"] == []


# --- seeded defects still fire ----------------------------------------------


def test_off_by_one_directly_below_is_flagged_but_remote_total_is_not(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    for row, value in enumerate((10, 20, 30, 40), start=1):
        ws.cell(row=row, column=1, value=f"Item {row}")
        ws.cell(row=row, column=2, value=value)
    ws["A5"] = "Total"
    ws["B5"] = "=SUM(B1:B3)"  # stops short of B4, which sits between range and total
    ws["D5"] = "=SUM(B1:B3)"  # same range, but the total lives elsewhere
    payload = _audit(_save(wb, tmp_path, "offbyone.xlsx"))
    assert _rules(payload)["RANGE_EXCLUSION"] == ["S!B5"]


def test_plug_inside_a_formula_run_is_flagged(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    ws.append(["Line", "Jan", "Feb", "Mar", "Total"])
    ws.append(["A", 1, 2, 3, "=SUM(B2:D2)"])
    ws.append(["B", 4, 5, 6, "=SUM(B3:D3)"])
    ws.append(["C", 7, 8, 9, 999])
    ws.append(["D", 1, 1, 1, "=SUM(B5:D5)"])
    payload = _audit(_save(wb, tmp_path, "plug.xlsx"))
    assert _rules(payload)["HARDCODE_IN_FORMULA_BLOCK"] == ["S!E4"]


def test_double_count_is_static_and_other_aggregates_are_not_totals():
    assert double_counted_cells("=SUM(B2:B5)+B5", "S") == [("B5", "B2:B5")]
    assert double_counted_cells("=B3+SUM(B2:B5)", "S") == [("B3", "B2:B5")]
    assert double_counted_cells("=SUM(B2:B5,B3)", "S") == [("B3", "B2:B5")]
    assert double_counted_cells("=SUM(B2:B5)-B3", "S") == []
    assert double_counted_cells("=SUM(B2:B5)+B7", "S") == []
    assert double_counted_cells("=AVERAGE(B2:B5)", "S") == []
    assert double_counted_cells("=SUM(B2:B5)/1000", "S") == []


def test_double_count_in_workbook_is_reported_without_recalculation(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    for row in range(1, 5):
        ws.cell(row=row, column=2, value=row * 10)
    ws["B5"] = "=SUM(B1:B4)+B4"
    payload = _audit(_save(wb, tmp_path, "double.xlsx"))
    findings = [f for f in payload["findings"] if f["rule_id"] == "TOTAL_MISMATCH"]
    assert [f["location"] for f in findings] == ["S!B5"]
    assert findings[0]["severity"] == "High"


def test_subtotal_double_count_is_flagged_but_subtotal_plus_other_is_not(tmp_path):
    wb = Workbook()
    bad = wb.active
    bad.title = "Bad"
    bad.append(["Item A", 5])
    bad.append(["Item B", 7])
    bad.append(["Subtotal", "=SUM(B1:B2)"])
    bad.append(["Item C", 9])
    bad.append(["Grand total", "=SUM(B1:B4)"])  # counts A and B twice
    good = wb.create_sheet("Good")
    good.append(["Item A", 5])
    good.append(["Item B", 7])
    good.append(["Subtotal", "=SUM(B1:B2)"])
    good.append(["Other", 9])
    good.append(["Total", "=SUM(B3:B4)"])  # subtotal + other
    payload = _audit(_save(wb, tmp_path, "subtotals.xlsx"))
    assert _rules(payload)["RANGE_INCLUDES_SUBTOTAL"] == ["Bad!B5"]


def test_hidden_row_is_reported_once_with_its_dependents(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    ws.append(["Visible", 1, 2, 3])
    ws.append(["Hidden", 4, 5, 6])
    ws.row_dimensions[2].hidden = True
    ws.append(["Total", "=B1+B2", "=C1+C2", "=D1+D2"])
    payload = _audit(_save(wb, tmp_path, "hidden.xlsx"))
    hidden = [f for f in payload["findings"] if f["rule_id"] == "HIDDEN_STRUCTURE_IN_TOTAL"]
    assert [f["location"] for f in hidden] == ["S!B3"]
    assert "3 visible formula(s)" in hidden[0]["evidence"][0]


def test_live_errors_collapse_to_their_root():
    formula_wb = Workbook()
    ws = formula_wb.active
    ws.title = "S"
    ws["A1"] = "=1/0"
    ws["B1"] = "=A1*2"
    ws["C1"] = "=B1+1"
    ws["D1"] = "=SUM(A1:C1)"
    value_wb = Workbook()
    vs = value_wb.active
    vs.title = "S"
    for coord in ("A1", "B1", "C1", "D1"):
        vs[coord] = "#DIV/0!"
    ctx = CheckContext(
        workbook_path=None,
        formula_wb=formula_wb,
        value_wb=value_wb,
        allowed_sheet_names={"S"},
        formulas=formula_cells(formula_wb),
        config={},
        inventory={},
    )
    findings = LiveErrorCheck().run(ctx)
    assert [f.location for f in findings] == ["S!A1"]
    assert "3 dependent cell(s)" in findings[0].evidence[1]


def test_numbers_stored_as_text_are_gated_on_use(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    ws.append(["Summed", 10, 20, "1,250", "=SUM(B1:D1)"])  # text inside a summed range -> High
    ws.append(["Numeric column", 1, 2, 3])
    ws.append(["More numbers", 4, 5, 6])
    ws.append(["Text number in numeric column", 7, "8", 9])  # Medium
    ws.append(["Notes", None, None, None, None, "42"])  # "42" in a column with no numbers -> nothing
    payload = _audit(_save(wb, tmp_path, "textnums.xlsx"))
    by_location = {f["location"]: f["severity"] for f in payload["findings"] if f["rule_id"] == "NUMBERS_STORED_AS_TEXT"}
    assert by_location == {"S!D1": "High", "S!C4": "Medium"}


def test_duplicate_keys_only_matter_inside_lookup_ranges(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    ws.append(["Total", 1])
    ws.append(["Other", 2])
    ws.append(["Total", 3])  # repeated label, nothing looks it up
    ws.append([])
    ws.append(["Key", "Value"])
    ws.append(["k1", 10])
    ws.append([" K1 ", 11])
    ws.append(["k2", 12])
    ws["D6"] = '=VLOOKUP("k1",A6:B8,2,FALSE)'
    payload = _audit(_save(wb, tmp_path, "dups.xlsx"))
    assert _rules(payload)["DUPLICATE_KEY"] == ["S!A6, S!A7"]


def test_merged_title_is_quiet_but_merge_inside_summed_block_is_flagged(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    ws.merge_cells("A1:D1")
    ws["A1"] = "Quarterly report"
    ws.append(["Item", 1, 2, 3])
    ws.append(["Item", 4, 5, 6])
    ws.merge_cells("B4:C4")
    ws["B4"] = 7
    ws["A5"] = "Total"
    ws["B5"] = "=SUM(B2:B4)"
    payload = _audit(_save(wb, tmp_path, "merged.xlsx"))
    assert _rules(payload)["MERGED_CELL_IN_DATA_RANGE"] == ["S!B4:C4"]


def test_blank_precedent_in_empty_row_is_quiet_but_anomalous_blank_is_flagged(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    ws.append(["Qty", "Price", "Amount"])
    ws.append([2, 5, "=A2*B2"])
    ws.append([3, None, "=A3*B3"])  # B3 blank where its row and column hold data
    ws.append([4, 7, "=A4*B4"])
    ws["E2"] = "=A2/A99"  # row 99 is empty space, not a broken link
    payload = _audit(_save(wb, tmp_path, "blanks.xlsx"))
    assert _rules(payload)["BLANK_PRECEDENT"] == ["S!C3"]


# --- the seeded corpora are fully catalogued ---------------------------------


def _normalize(location: str) -> str:
    return location.replace(" ", "")


@pytest.mark.parametrize("entry", json.loads(CATALOG.read_text(encoding="utf-8"))["workbooks"], ids=lambda e: e["path"])
def test_every_static_finding_on_seeded_workbooks_is_catalogued(entry):
    workbook = ROOT / entry["path"]
    if not workbook.exists():
        pytest.skip(f"{workbook} not generated")
    payload = _audit(workbook)
    expected = {(seed["rule_id"], _normalize(seed["location"])) for seed in entry["seeded"]}
    static_expected = {
        (seed["rule_id"], _normalize(seed["location"]))
        for seed in entry["seeded"]
        if not seed["requires_recalc"]
    }
    produced = {(f["rule_id"], _normalize(f["location"])) for f in payload["findings"]}

    uncatalogued = sorted(pair for pair in produced if pair[0] not in VALUE_DEPENDENT and pair not in expected)
    assert not uncatalogued, f"Findings without a seed entry (false positives?): {uncatalogued}"

    missed = sorted(pair for pair in static_expected if pair not in produced)
    assert not missed, f"Seeded static defects not detected: {missed}"
