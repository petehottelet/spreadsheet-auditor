"""Tokenizer-based formula parser: references, names, tables, literals, normalization."""

from __future__ import annotations

from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.worksheet.table import Table

from spreadsheet_auditor.formula_parser import (
    extract_functions,
    extract_numeric_literals,
    extract_references,
    normalize_formula,
    parse_formula,
)
from spreadsheet_auditor.names import NameTable
from spreadsheet_auditor.workbook_inventory import formula_cells


def _refs(formula: str, **kwargs) -> list[tuple[str | None, str]]:
    return [(ref.sheet, ref.ref) for ref in extract_references(formula, **kwargs)]


def test_function_names_with_digits_are_not_references():
    assert _refs("=LOG10(B5)") == [(None, "B5")]
    assert _refs("=DAYS360(B2,C2)") == [(None, "B2"), (None, "C2")]
    assert _refs("=ATAN2(1,2)") == []
    assert extract_functions("=LOG10(B5)+DAYS360(B2,C2)") == {"LOG10", "DAYS360"}


def test_names_with_more_than_three_leading_letters_are_not_cells():
    parsed = parse_formula("=EBITDA2025*2")
    assert parsed.references == []
    assert parsed.unresolved_names == ["EBITDA2025"]


def test_defined_name_resolves_to_its_range():
    wb = Workbook()
    ws = wb.active
    ws.title = "Budget"
    wb.defined_names["Revenue_FY"] = DefinedName("Revenue_FY", attr_text="Budget!$B$3:$D$3")
    names = NameTable.from_workbook(wb)
    refs = extract_references("=SUM(Revenue_FY)", names=names, origin=("Budget", 5, 2))
    assert [(r.sheet, r.ref, r.is_range, r.via_name) for r in refs] == [
        ("Budget", "B3:D3", True, "Revenue_FY")
    ]


def _table_workbook() -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Region", "Sales", "Double"])
    ws.append(["N", 10, "=[@Sales]*2"])
    ws.append(["S", 20, "=[@Sales]*2"])
    ws.append(["E", 30, "=[@Sales]*2"])
    ws.add_table(Table(displayName="Table1", ref="A1:C4"))
    return wb


def test_structured_references_resolve_against_the_table():
    names = NameTable.from_workbook(_table_workbook())
    column = parse_formula("=SUM(Table1[Sales])", names=names, origin=("Data", 1, 5))
    assert column.external_references == []
    assert [(r.sheet, r.ref) for r in column.references] == [("Data", "B2:B4")]

    # Unqualified [@Column] is only legal inside the table (calculated column C).
    this_row = parse_formula("=[@Sales]*2", names=names, origin=("Data", 3, 3))
    assert this_row.external_references == []
    assert [(r.sheet, r.ref) for r in this_row.references] == [("Data", "B3")]

    headers = parse_formula("=SUM(Table1[[#Headers],[Sales]])", names=names)
    assert [(r.sheet, r.ref) for r in headers.references] == [("Data", "B1")]

    whole = parse_formula("=COUNTA(Table1[])", names=names)
    assert [(r.sheet, r.ref) for r in whole.references] == [("Data", "A2:C4")]


def test_unknown_table_is_unresolved_not_an_external_link():
    parsed = parse_formula("=SUM(Table9[Sales])")
    assert parsed.external_references == []
    assert parsed.unresolved_structured == ["Table9[Sales]"]
    assert parsed.references == []


def test_external_references_are_classified():
    for formula in ("=[1]Sheet1!A1", "='[Book.xlsx]Sheet1'!A1*2"):
        parsed = parse_formula(formula)
        assert len(parsed.external_references) == 1, formula
        assert parsed.references == [], formula


def test_error_literals_and_deleted_references():
    assert parse_formula("=SUM(#REF!)").deleted_reference
    assert parse_formula("=Sheet1!#REF!").deleted_reference
    assert not parse_formula('=IF(A1="#REF!",1,0)').deleted_reference


def test_whole_column_and_row_references_are_unbounded():
    refs = extract_references("=SUM(A:A)+SUM(Sheet2!2:2)")
    assert [(r.ref, r.bounded) for r in refs] == [("A:A", False), ("2:2", False)]


def test_literals_ignore_references_that_share_digits():
    assert extract_numeric_literals("=C3*C35") == []
    assert extract_numeric_literals("=D5+D55") == []
    assert extract_numeric_literals("=A1*1.05") == ["1.05"]
    assert extract_numeric_literals("=A1*10%") == []
    assert extract_numeric_literals("=A1*2025") == []
    assert extract_numeric_literals('=IF(A1>0,"1.05",0)') == []


def test_feature_flags_for_unsupported_constructs():
    assert parse_formula("=SUM(Sheet1:Sheet3!A1)").three_d_references == ["Sheet1:Sheet3!A1"]
    assert parse_formula("=_xlfn.SINGLE(A1:A10)").implicit_intersection
    assert parse_formula("=@A1:A10").implicit_intersection
    assert parse_formula("=_xlfn.ANCHORARRAY(A1)").spill
    assert "XLOOKUP" in parse_formula("=_xlfn.XLOOKUP(A1,B:B,C:C)").functions


def test_unparseable_formula_is_flagged_not_fatal():
    parsed = parse_formula('="unterminated')
    assert parsed.parse_error
    assert parsed.references == []


def test_normalize_keeps_absolute_anchors_stable_across_a_row():
    assert normalize_formula("=B2*$B$1", 2, 3) == normalize_formula("=C2*$B$1", 2, 4)
    assert normalize_formula("=B2*$B$1", 2, 3) != normalize_formula("=B2*$C$1", 2, 3)
    assert normalize_formula("=SUM(B2:B4)", 5, 2) == normalize_formula("=SUM(C2:C4)", 5, 3)
    assert normalize_formula("=SUM(B2:B4)", 5, 2) != normalize_formula("=SUM(B2:B5)", 5, 2)


def test_quoted_and_lowercase_sheet_names():
    refs = extract_references("='My Sheet'!A1+summary!B2")
    assert [(r.sheet, r.ref) for r in refs] == [("My Sheet", "A1"), ("summary", "B2")]


def test_array_formulas_are_visible_to_the_inventory():
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    ws["A1"] = 1
    ws["A2"] = 2
    ws["B1"] = ArrayFormula("B1:B2", "=A1:A2*2")
    ws["C1"] = "=SUM(A1:A2)"
    cells = formula_cells(wb)
    assert {(cell["location"], cell["array"]) for cell in cells} == {("S!B1", True), ("S!C1", False)}
    assert all(cell["formula"].startswith("=") for cell in cells)
