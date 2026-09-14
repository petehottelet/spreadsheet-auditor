"""Hardening: bounded scans, bounded cycle detection, real suppression semantics,
a cooperative time budget, and non-ASCII output that never crashes."""

from __future__ import annotations

import argparse
import json
import subprocess
import sys
import time
from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.utils.cell import get_column_letter
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.worksheet.table import Table

from spreadsheet_auditor import audit as audit_module
from spreadsheet_auditor.audit import apply_impact_escalation, detect_cycles
from spreadsheet_auditor.budget import AuditTimeout, Budget
from spreadsheet_auditor.config_loader import load_config
from spreadsheet_auditor.finding import Finding
from spreadsheet_auditor.formula_drift import detect_hardcode_breaks
from spreadsheet_auditor.suppressions import apply_suppressions


def _run(workbook: Path) -> dict:
    result = subprocess.run(
        [sys.executable, "-m", "spreadsheet_auditor", str(workbook), "--json", "-", "--fail-on", "None"],
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
    )
    assert result.returncode in (0, 1, 2), result.stderr
    return json.loads(result.stdout)


def _formula(row: int, col: int, formula: str, sheet: str = "S") -> dict:
    coord = f"{get_column_letter(col)}{row}"
    return {
        "sheet": sheet,
        "row": row,
        "col": col,
        "coord": coord,
        "location": f"{sheet}!{coord}",
        "formula": formula,
    }


def _finding(location: str, rule_id: str = "BROKEN_REFERENCE") -> Finding:
    return Finding(
        rule_id=rule_id,
        severity="High",
        error_confidence="Review",
        detection_mode="DET",
        location=location,
        title="t",
        evidence=[],
        suggested_fix="",
    )


# --- used-range inflation ---------------------------------------------------


def test_function_names_and_far_names_do_not_inflate_the_scan(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    ws["A1"] = 45000
    ws["A2"] = 45090
    ws["B1"] = "=LOG10(A1)"
    ws["B2"] = "=DAYS360(A1,A2)"
    ws["B3"] = "=EBITDA2025*2"
    ws["B4"] = "=ATAN2(1,2)"
    path = tmp_path / "funcs.xlsx"
    wb.save(path)

    started = time.monotonic()
    payload = _run(path)
    assert time.monotonic() - started < 15
    assert payload["workbook"]["formulas_scanned"] == 4
    assert not [f for f in payload["findings"] if f["rule_id"] == "BLANK_PRECEDENT"]
    assert "unresolved_defined_names" in payload["coverage"]["unsupported_features"]


# --- circular references ----------------------------------------------------


def test_self_inclusive_total_is_reported_once():
    findings = detect_cycles([_formula(3, 1, "=SUM(A1:A3)")])
    assert [f.location for f in findings] == ["S!A3"]
    assert "own cell" in findings[0].title


def test_filled_down_self_inclusive_total_stays_bounded():
    n = 40
    cells = [_formula(r, 1, f"=SUM(A1:A{n})") for r in range(1, n + 1)]
    started = time.monotonic()
    findings = detect_cycles(cells)
    assert time.monotonic() - started < 5
    assert len(findings) == 1
    assert findings[0].location == "S!A1"
    assert f"{n} cells" in findings[0].evidence[0]


def test_whole_column_self_reference_is_a_cycle():
    findings = detect_cycles([_formula(5, 1, "=SUM(A:A)")], extents={"S": (5, 1)})
    assert [f.location for f in findings] == ["S!A5"]


def test_long_dependency_chain_does_not_hit_the_recursion_limit():
    cells = [_formula(r, 1, f"=A{r - 1}+1") for r in range(2, 3002)]
    assert detect_cycles(cells) == []


def test_two_cell_cycle_is_one_finding():
    findings = detect_cycles([_formula(1, 1, "=B1+1"), _formula(1, 2, "=A1+1")])
    assert len(findings) == 1
    assert findings[0].location == "S!A1"
    assert "2 cells" in findings[0].evidence[0]


# --- suppressions -----------------------------------------------------------


def test_cell_suppression_does_not_match_by_prefix():
    findings = [
        _finding("Imports!A1"),
        _finding("Imports!A10"),
        _finding("Imports!A100"),
        _finding("Imports!AA1"),
    ]
    apply_suppressions(findings, [{"rule_id": "BROKEN_REFERENCE", "range": "Imports!A1", "reason": "r"}])
    assert [f.location for f in findings if f.suppressed] == ["Imports!A1"]


def test_range_suppression_covers_cells_inside_it():
    findings = [
        _finding("Imports!A5"),
        _finding("Imports!A50"),
        _finding("Imports!A500"),
        _finding("Other!A5"),
    ]
    apply_suppressions(
        findings, [{"rule_id": "BROKEN_REFERENCE", "range": "imports!$A$1:$A$100", "reason": "r"}]
    )
    assert [f.location for f in findings if f.suppressed] == ["Imports!A5", "Imports!A50"]


def test_whole_column_sheet_and_multi_location_suppressions():
    column = [_finding("Imports!B7"), _finding("Imports!C7")]
    apply_suppressions(column, [{"rule_id": "BROKEN_REFERENCE", "range": "Imports!B:B", "reason": "r"}])
    assert [f.location for f in column if f.suppressed] == ["Imports!B7"]

    sheet = [_finding("Imports!Z99"), _finding("Model!Z99")]
    apply_suppressions(sheet, [{"rule_id": "BROKEN_REFERENCE", "range": "Imports", "reason": "r"}])
    assert [f.location for f in sheet if f.suppressed] == ["Imports!Z99"]

    multi = [_finding("Model!A37, Model!A38", rule_id="DUPLICATE_KEY")]
    apply_suppressions(multi, [{"rule_id": "DUPLICATE_KEY", "range": "Model!A38", "reason": "r"}])
    assert multi[0].suppressed


def test_headline_output_escalation_is_exact():
    hit = _finding("Summary!C1")
    miss = _finding("Summary!C10")
    apply_impact_escalation([hit, miss], {"scope": {"headline_outputs": ["Summary!C1"]}})
    assert hit.impact.get("feeds_headline_output") is True
    assert hit.severity == "Critical"
    assert "feeds_headline_output" not in miss.impact
    assert miss.severity == "High"


# --- time budget ------------------------------------------------------------


def test_default_config_has_a_time_budget():
    assert load_config(None)["limits"]["timeout_seconds"] == 120


def test_budget_interrupts_a_check_and_zero_disables_it():
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    for row in range(1, 50):
        ws.cell(row=row, column=1, value=row)
    ws["B1"] = "=A1"
    ws["B49"] = "=A49"

    expired = Budget(seconds=1e-9)
    time.sleep(0.01)
    with pytest.raises(AuditTimeout):
        detect_hardcode_breaks(wb, budget=expired)
    assert detect_hardcode_breaks(wb, budget=Budget(seconds=0)) == []


def test_orchestrator_reports_timeout_as_limitation(tmp_path, monkeypatch):
    class ExpiredBudget(audit_module.Budget):
        def tick(self) -> None:
            raise AuditTimeout("expired")

    monkeypatch.setattr(audit_module, "Budget", ExpiredBudget)
    wb = Workbook()
    ws = wb.active
    ws["A1"] = 1
    ws["B1"] = "=A1*2"
    path = tmp_path / "t.xlsx"
    wb.save(path)
    args = argparse.Namespace(
        workbook=str(path),
        config=None,
        ignore=str(tmp_path / "no-ignore-file"),
        fail_on="None",
        strict=False,
        recalc_timeout=None,
    )
    payload, code = audit_module.audit_workbook(args)
    assert payload["coverage"]["truncated"]["timeout"] is True
    assert any("timeout" in note.lower() for note in payload["coverage"]["limitations"])
    assert code == 2


# --- tables, array formulas, output encoding ---------------------------------


def test_structured_reference_is_not_an_external_link(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Region", "Sales"])
    ws.append(["N", 10])
    ws.append(["S", 20])
    ws.append(["E", 30])
    ws.add_table(Table(displayName="Table1", ref="A1:B4"))
    ws["D1"] = "=SUM(Table1[Sales])"
    ws["E1"] = ArrayFormula("E1:E3", "=B2:B4*2")
    path = tmp_path / "table.xlsx"
    wb.save(path)

    payload = _run(path)
    rules = {f["rule_id"] for f in payload["findings"]}
    assert "BROKEN_REFERENCE" not in rules
    assert "BLANK_PRECEDENT" not in rules
    assert payload["coverage"]["external_links_present"] is False
    assert payload["workbook"]["formulas_scanned"] == 2
    assert "array_formulas" in payload["coverage"]["unsupported_features"]
    assert any("array formula" in note for note in payload["coverage"]["limitations"])


def test_non_ascii_labels_do_not_crash_redirected_output(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    ws["A1"] = " 总计 Σ "  # whitespace-padded first-column label -> WHITESPACE_KEY evidence
    ws["B1"] = 5
    ws["C1"] = "=SUM(B1:B1)"
    path = tmp_path / "unicode.xlsx"
    wb.save(path)

    result = subprocess.run(
        [sys.executable, "-m", "spreadsheet_auditor", str(path), "--fail-on", "None"],
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
    )
    assert result.returncode in (0, 1, 2), result.stderr
    assert "Traceback" not in result.stderr
    assert "总计" in result.stdout
