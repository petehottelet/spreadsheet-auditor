"""Fixes driven by the SpreadsheetBench precision measurement.

Each test reproduces a false-positive pattern that the real-world corpus
surfaced, with a control case next to it so the rule still fires when it
should.
"""

from __future__ import annotations

import json
import subprocess
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook


def _audit(path: Path) -> dict[str, list[dict]]:
    result = subprocess.run(
        [sys.executable, "-m", "spreadsheet_auditor", str(path), "--json", "-", "--fail-on", "None"],
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
    )
    assert result.returncode in (0, 1, 2), result.stderr
    by_rule: dict[str, list[dict]] = defaultdict(list)
    for finding in json.loads(result.stdout)["findings"]:
        by_rule[finding["rule_id"]].append(finding)
    return by_rule


def test_positional_references_are_not_value_dependencies(tmp_path):
    """ROWS(A$1:A1) filled down names a range that includes the cell itself, but
    ROWS only looks at the shape of the reference, so Excel does not treat it
    as circular and neither should we. The same goes for a blank cell handed
    to ROW or COLUMN."""
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    for row in range(1, 6):
        ws.cell(row=row, column=1, value=f"=ROWS(A$1:A{row})")
        ws.cell(row=row, column=2, value=f'=IF(ROWS(B$1:B{row})>3,"",ROWS(B$1:B{row}))')
    ws["D1"] = "=ROW(F9)+COLUMN(F9)"  # F9 is blank; positional use, not a value read
    ws["D2"] = "=F9*2"  # control: a real value read of the blank cell
    for row in (1, 2, 3, 4, 5, 6, 7, 8, 10, 11):
        ws.cell(row=row, column=6, value=row)  # column F is a filled column with a gap at F9
    ws["E9"] = 1
    path = tmp_path / "positional.xlsx"
    wb.save(path)
    by_rule = _audit(path)
    assert "CIRCULAR_REFERENCE" not in by_rule
    assert [f["location"] for f in by_rule.get("BLANK_PRECEDENT", [])] == ["S!D2"]


def test_blank_precedent_skips_formulas_that_test_the_cell_for_blank(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    ws.append(["Old", "Keep", "New"])
    ws.append(["Alex", '=IF(C2="",A2,C2)', "a"])
    ws.append(["Sam", "=IF(ISBLANK(C3),A3,C3)", None])  # tests C3 for blank: handled
    ws.append(["Jo", '=IF(C4<>"",C4,A4)', None])  # handled
    ws.append(["Kim", "=C5*2", None])  # control: reads the blank without checking
    ws.append(["Pat", "=A6&C6", "b"])
    ws.append(["Lee", '=IF(C7="",A7,C7)', "c"])
    for row in range(8, 16):
        ws.append([f"Row {row}", f"=C{row}*2", "x"])  # column C is mostly filled, so the blanks above are gaps
    path = tmp_path / "blank_tests.xlsx"
    wb.save(path)
    by_rule = _audit(path)
    assert [f["location"] for f in by_rule.get("BLANK_PRECEDENT", [])] == ["S!B5"]


def test_hardcode_ignores_an_input_column_between_running_sums(tmp_path):
    """Columns D and F are running sums with the same relative pattern; the
    input column E between them is not a plug. A lone constant inside a
    column of formulas still is."""
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    ws.append(["Item", "Jan", "Feb", "YTD Feb", "Mar", "YTD Mar", "Double"])
    for row in range(2, 6):
        ws.append([f"Item {row}", 10 * row, 11 * row, f"=B{row}+C{row}", 12 * row, f"=D{row}+E{row}", f"=F{row}*2"])
    ws["G4"] = 999  # the plug: formulas above and below in column G, formula left and nothing right
    path = tmp_path / "running_sums.xlsx"
    wb.save(path)
    by_rule = _audit(path)
    assert [f["location"] for f in by_rule.get("HARDCODE_IN_FORMULA_BLOCK", [])] == ["S!G4"]


def test_external_links_are_reported_once_per_source_per_sheet(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    for row in range(1, 6):
        ws.cell(row=row, column=1, value=f"=+[1]BA!$AI${row}")
        ws.cell(row=row, column=2, value=f"=[2]Rates!$B${row}")
    ws["D1"] = "=SUM(#REF!)"  # a genuinely broken reference stays a separate finding
    path = tmp_path / "links.xlsx"
    wb.save(path)
    by_rule = _audit(path)
    external = [f for f in by_rule.get("BROKEN_REFERENCE", []) if "external" in f["title"].lower()]
    assert len(external) == 2, [f["location"] for f in external]
    assert all("5 cell" in " ".join(f["evidence"]) for f in external)
    assert all(f["severity"] == "Low" for f in external)
    assert [f["location"] for f in by_rule["BROKEN_REFERENCE"] if "deleted" in f["title"].lower()] == ["S!D1"]


def test_deliberate_self_reference_is_still_a_cycle(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    ws["A1"] = "go"
    ws["B1"] = '=IF(A1="go",IF(IFERROR(1/(1/B1),"")="",TODAY(),B1),"")'  # freezes a date once set
    path = tmp_path / "selfref.xlsx"
    wb.save(path)
    assert [f["location"] for f in _audit(path).get("CIRCULAR_REFERENCE", [])] == ["S!B1"]


def test_duplicate_keys_count_only_in_first_match_key_columns(tmp_path):
    """SUMIFS criteria columns repeat by design and VLOOKUP's return column is
    not a key; the first column of a VLOOKUP table and the column a MATCH
    searches are."""
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    ws.append(["Shift", "Machine", "Hours", "Postcode", "School"])
    ws.append(["Night", "M1", 8, "AB1", "North"])
    ws.append(["Night", "M2", 7, "AB2", "South"])
    ws.append(["Early", "M1", 6, "AB1", "East"])
    ws.append(["Total night hours", '=SUMIFS(C2:C4,A2:A4,"Night")'])
    ws.append(["Machine hours", '=VLOOKUP("M2",B2:C4,2,FALSE)'])
    ws.append(["School by postcode", '=INDEX(E2:E4,MATCH("AB1",D2:D4,0))'])
    path = tmp_path / "keys.xlsx"
    wb.save(path)
    assert [f["location"] for f in _audit(path).get("DUPLICATE_KEY", [])] == ["S!B2, S!B4", "S!D2, S!D4"]


def test_drift_ignores_rows_of_unrelated_columns_and_chain_seeds(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    ws.append(["Name", "Initial", "Mirror 1", "Mirror 2"])
    for row in range(2, 6):
        # column B derives from A while C and D mirror another sheet's columns:
        # three formulas per row, two sharing a pattern, none of them drift.
        ws.append([f"Name {row}", f"=LEFT(A{row},1)", f"=Data!A{row}", f"=Data!B{row}"])
    ws["F1"], ws["G1"], ws["H1"], ws["I1"] = "Year 0", "Year 1", "Year 2", "Year 3"
    ws["F2"], ws["G2"], ws["H2"], ws["I2"] = "=Data!C2", "=F2*1.05", "=G2*1.05", "=H2*1.05"  # seed then chain
    ws["F4"], ws["G4"], ws["H4"], ws["I4"] = "=Data!C4", "=F4*1.05", "=G4*1.05", "=G4*1.05"  # I4 really drifts
    data = wb.create_sheet("Data")
    for row in range(1, 6):
        data.append([f"a{row}", f"b{row}", row])
    path = tmp_path / "drift.xlsx"
    wb.save(path)
    assert [f["location"] for f in _audit(path).get("FORMULA_DRIFT", [])] == ["S!I4"]


def test_merged_labels_and_direct_reads_are_quiet(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    ws.merge_cells("A1:C1")
    ws["A1"] = "Section title"  # text merge inside SUM(A1:C3): presentation
    ws.append([1, 2, 3])
    ws.merge_cells("B3:C3")
    ws["A3"], ws["B3"] = 4, 5  # numeric merge inside the summed block: real
    ws["E1"] = "=SUM(A1:C3)"
    ws.merge_cells("A5:B5")
    ws["A5"] = 449
    ws["C5"] = "=A5*2"  # addresses the top-left cell directly: fine
    path = tmp_path / "merges.xlsx"
    wb.save(path)
    assert [f["location"] for f in _audit(path).get("MERGED_CELL_IN_DATA_RANGE", [])] == ["S!B3:C3"]


def test_text_identifiers_are_not_numbers_stored_as_text(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    ws.append(["Serial", "Order", "Amount"])
    ws.append(["0013", "223143", "1,250"])
    ws.append(["0014", "223144", 200])
    ws.append(["0015", "223145", 300])
    ws.append(["0016", "223146", 400])
    ws["E1"] = '=COUNTIF(B2:B5,"223144")'  # matches text against text: fine
    ws["E2"] = '=VLOOKUP("0014",A2:C5,3,FALSE)'  # a text key: fine
    ws["E3"] = "=SUM(C2:C5)"  # consumes C2 as a number: the text drops out
    path = tmp_path / "identifiers.xlsx"
    wb.save(path)
    found = {f["location"]: f["severity"] for f in _audit(path).get("NUMBERS_STORED_AS_TEXT", [])}
    assert found == {"S!C2": "High"}


def test_whitespace_reports_the_stray_padded_key_not_the_export(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    ws.append(["Name ", "Status"])  # a padded header over data: presentation
    ws.append(["Andrew", "Fully linked      "])
    ws.append(["Courtney ", "To be actioned    "])  # the one padded name; a uniformly padded status column
    ws.append(["Donita", "Fully linked      "])
    ws.append(["   Indented label", "Fully linked      "])  # leading spaces indent a label
    ws["D1"] = '=COUNTIF(A2:A4,"Courtney")'
    ws["D2"] = '=COUNTIF(B2:B5,"Fully linked")'
    path = tmp_path / "padding.xlsx"
    wb.save(path)
    found = {f["location"]: f["severity"] for f in _audit(path).get("WHITESPACE_KEY", [])}
    assert found == {"S!A3": "Medium", "S!B2": "Low"}


def test_whole_column_reference_only_in_array_contexts(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    for row in range(1, 6):
        ws.append([row, row * 10])
    ws["D1"] = '=COUNTIF(A:A,">2")'  # bounded internally: fine
    ws["D2"] = "=VLOOKUP(3,A:B,2,FALSE)"  # fine
    ws["D3"] = "=INDEX(B:B,MATCH(4,A:A,0))"  # fine
    ws["D4"] = "=SUMPRODUCT((A:A>2)*B:B)"  # array over a million rows
    ws["D5"] = '=LOOKUP(2,1/(B:B<>""),B:B)'  # array over a million rows
    path = tmp_path / "columns.xlsx"
    wb.save(path)
    assert [f["location"] for f in _audit(path).get("WHOLE_COLUMN_REFERENCE", [])] == ["S!D4", "S!D5"]


def test_today_is_one_note_per_sheet_and_rand_is_per_pattern(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    ws.append(["Due", "Days left", "Sample"])
    for row in range(2, 6):
        ws.append([f"=DATE(2026,1,{row})", f"=A{row}-TODAY()", "=RAND()"])
    path = tmp_path / "volatile.xlsx"
    wb.save(path)
    found = [(f["location"], f["severity"]) for f in _audit(path).get("VOLATILE_FUNCTION", [])]
    assert found == [("S!C2", "Medium"), ("S!B2", "Low")]


def test_hidden_rows_skipped_by_aggregate_are_not_inputs(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    for row in range(1, 6):
        ws.append([row])
    ws.row_dimensions[2].hidden = True
    ws.row_dimensions[4].hidden = True
    ws["C1"] = "=AGGREGATE(9,3,A1:A5)"  # option 3 ignores hidden rows
    ws["C2"] = "=SUBTOTAL(109,A1:A5)"  # 109 ignores hidden rows
    ws["C3"] = "=SUM(A1:A5)"  # counts them
    path = tmp_path / "aggregate.xlsx"
    wb.save(path)
    hidden = _audit(path).get("HIDDEN_STRUCTURE_IN_TOTAL", [])
    assert [f["location"] for f in hidden] == ["S!C3"]
    assert hidden[0]["evidence"][0].startswith("Hidden rows 2, 4 on S feed 1 visible formula(s)")


def test_running_count_is_not_an_exclusion_and_carried_total_is_not_a_subtotal(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    ws.append(["Index", "Month", "Count"])
    ws.append([1, 1, "=COUNT($A$2:A2)"])  # expands as it is filled; B2 is not excluded data
    ws.append([2, 2, "=COUNT($A$2:A3)"])
    ws["E1"], ws["F1"], ws["G1"], ws["H1"] = "This week", None, "Last weeks total", "Total overall"
    ws["E2"], ws["G2"], ws["H2"] = 8, 142, "=SUM(E2:G2)"  # carried forward, not double counted
    path = tmp_path / "running.xlsx"
    wb.save(path)
    by_rule = _audit(path)
    assert "RANGE_EXCLUSION" not in by_rule
    assert "RANGE_INCLUDES_SUBTOTAL" not in by_rule


def test_range_length_ignores_links_and_disjoint_groups(tmp_path):
    """Two relative patterns in the totals column keep FORMULA_DRIFT out of the
    way (as in the seeded corpus), so the length check is exercised alone."""
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    ws.append(["Row", "In 1", "In 2", "In 3", "In 4", None, "Total", "Link", "Debit A", "Debit B", "Credit A", "Credit B", "Debits", "Credits"])
    for row in range(2, 6):
        ws.append([f"r{row}", 1, 2, 3, 4, None, f"=SUM(B{row}:D{row})", f"=SUM(G{row}:G{row})", 5, 6, 7, 8, f"=SUM(I{row}:J{row})", f"=SUM(K{row}:L{row})"])
    ws["G3"] = "=SUM(C3:E3)"  # same length, different pattern
    ws["G4"] = "=SUM(B4:C4)"  # the real short range
    path = tmp_path / "lengths.xlsx"
    wb.save(path)
    assert [f["location"] for f in _audit(path).get("RANGE_LENGTH_MISMATCH", [])] == ["S!G4"]


def test_literals_in_structural_slots_are_not_assumptions(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    ws.append(["Key", "Text", "Hours", "Date"])
    ws.append(["k1", "12-31 10:00", 37, "=DATE(2026,1,1)"])
    ws.append(["k2", "01-15 08:30", 41, "=DATE(2026,2,1)"])
    ws["F1"] = "=VLOOKUP(A2,A2:C3,3,FALSE)"  # column index
    ws["F2"] = '=MID(B2,FIND(" ",B2,1),6)'  # string positions
    ws["F3"] = "=WEEKDAY(D2,11)+HOUR(D2)/60"  # type code and a unit constant
    ws["F4"] = '=IF(C2<37.5,"part","full")'  # a threshold: this one is an assumption
    ws["F5"] = '=IF(C3<37.5,"part","full")'  # same pattern one row down: reported once
    path = tmp_path / "literals.xlsx"
    wb.save(path)
    literals = _audit(path).get("LITERAL_CONSTANT", [])
    assert [f["location"] for f in literals] == ["S!F4"]
    assert "2 cells" in literals[0]["evidence"][1]


def test_iferror_is_reported_once_per_formula_block(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    for row in range(1, 6):
        ws.append([row, f"=IFERROR(1/A{row},0)"])
    ws["D1"] = '=IFERROR(VLOOKUP(1,A1:B5,2,FALSE),"")'
    path = tmp_path / "masks.xlsx"
    wb.save(path)
    masks = _audit(path).get("IFERROR_MASK", [])
    assert [f["location"] for f in masks] == ["S!B1", "S!D1"]
    assert "5 cells" in masks[0]["evidence"][1]


# --- second round: patterns the re-audited corpus still showed ---------------


def test_cell_filename_idiom_is_not_a_cycle(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "BRE"
    ws["Q2"] = '=MID(CELL("filename",Q2),FIND("]",CELL("filename",Q2))+1,255)'
    path = tmp_path / "cellname.xlsx"
    wb.save(path)
    assert "CIRCULAR_REFERENCE" not in _audit(path)


def test_compared_text_codes_inside_sumproduct_are_not_numbers(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Raw"
    ws.append(["Serial", "Code", "Dup"])
    for row, code in enumerate(("9940", "9941", "9940", "9942"), start=2):
        ws.append([row, code, f"=SUMPRODUCT(($B$2:$B$5=$B{row})*1)"])
    ws["E1"] = "=SUM(B2:B5)"  # the control: a SUM does consume the codes as numbers
    path = tmp_path / "codes.xlsx"
    wb.save(path)
    found = {f["location"]: f["severity"] for f in _audit(path).get("NUMBERS_STORED_AS_TEXT", [])}
    assert found == {"Raw!B2": "High", "Raw!B3": "High", "Raw!B4": "High", "Raw!B5": "High"}
    ws["E1"] = None
    wb.save(path)
    assert "NUMBERS_STORED_AS_TEXT" not in _audit(path)


def test_input_column_with_a_total_row_above_and_a_plug_beside_is_still_an_input(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "DRP"
    ws.append(["Buy", "Sell", "Net", "Buy 2", "Sell 2", "Net 2"])
    for row in range(2, 6):
        ws.append([10 * row, row, f"=A{row}-B{row}", 20 * row, 2 * row, f"=C{row}+D{row}-E{row}"])
    ws.append(["=SUM(A2:A5)", "=SUM(B2:B5)", "=SUM(C2:C5)", "=SUM(D2:D5)", "=SUM(E2:E5)", "=SUM(F2:F5)"])
    for row in range(7, 11):
        ws.append([10 * row, row, f"=A{row}-B{row}", 20 * row, 2 * row, f"=C{row}+D{row}-E{row}"])
    ws["C4"] = -1  # a plug in the Net column; the inputs beside it stay inputs
    path = tmp_path / "drp.xlsx"
    wb.save(path)
    assert [f["location"] for f in _audit(path).get("HARDCODE_IN_FORMULA_BLOCK", [])] == ["DRP!C4"]


def test_drift_skips_label_links_that_step_and_totals_rows_that_mix_aggregates(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Total"
    ws.append(["Name", "Hours", "Pay"])
    for i, row in enumerate(range(2, 8)):
        src = 6 + 2 * i  # the label column links every second row of the source sheet
        ws.append([f"=Jan!A{src}", f"=Jan!B{src}+Feb!B{src}", f"=Jan!C{src}+Feb!C{src}"])
    ws.append(["Total", "=SUM(B2:B7)", "=AVERAGE(C2:C7)"])  # one sums, one averages the same block
    for title in ("Jan", "Feb"):
        sheet = wb.create_sheet(title)
        for row in range(1, 20):
            sheet.append([f"n{row}", row, row * 10])
    path = tmp_path / "labels.xlsx"
    wb.save(path)
    assert "FORMULA_DRIFT" not in _audit(path)


def test_function_swapped_inside_a_run_of_totals_is_still_drift(tmp_path):
    """The EUSES 'function replaced' fault: one SUM in a row of five became an
    AVERAGE. A three-cell totals row mixing SUM and AVERAGE stays quiet."""
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    ws.append(["Item", "Q1", "Q2", "Q3", "Q4", "Q5"])
    for row in range(2, 6):
        ws.append([f"r{row}", row, row * 2, row * 3, row * 4, row * 5])
    ws.append(["Total", "=SUM(B2:B5)", "=SUM(C2:C5)", "=AVERAGE(D2:D5)", "=SUM(E2:E5)", "=SUM(F2:F5)"])
    pivot = wb.create_sheet("Pivot")  # headers announce the mix: design, not drift
    pivot.append(["FY", "Sum of Gen.", "Average of PLF", "Average of Avail.", "Average of Grid"])
    for row in range(2, 6):
        pivot.append([f"y{row}", row * 100, row, row * 2, row * 3])
    pivot.append([None, "=SUM(B2:B5)", "=AVERAGE(C2:C5)", "=AVERAGE(D2:D5)", "=AVERAGE(E2:E5)"])
    path = tmp_path / "ffr.xlsx"
    wb.save(path)
    assert [f["location"] for f in _audit(path).get("FORMULA_DRIFT", [])] == ["S!D6"]


def test_drift_skips_a_link_heading_a_label_column(tmp_path):
    """A summary box: M2 and N2 total their columns, L2 = M14 points at another
    total. At L2 the totals' formula would sum L4:L11, a column of names, so L2
    is not part of that run. When column L holds amounts, L2 = M14 is a total
    that stopped totalling and stays reported."""
    wb = Workbook()
    for title, first_column in (("Box", lambda row: f"Payee {row}"), ("Amounts", lambda row: row * 3)):
        ws = wb.create_sheet(title)
        ws["L1"], ws["M1"], ws["N1"] = "Total Tax", "Monthly Income", "Total Income"
        ws["L2"], ws["M2"], ws["N2"] = "=M14", "=SUM(M4:M11)", "=SUM(N4:N11)"
        for row in range(4, 12):
            ws[f"L{row}"], ws[f"M{row}"], ws[f"N{row}"] = first_column(row), row * 100, row * 10
        ws["L14"], ws["M14"] = "Total Property Tax Paid", 1200
    wb.remove(wb["Sheet"])
    path = tmp_path / "box.xlsx"
    wb.save(path)
    assert [f["location"] for f in _audit(path).get("FORMULA_DRIFT", [])] == ["Amounts!L2"]


def test_drift_skips_one_of_a_set_of_counts_over_a_block(tmp_path):
    """Row 5 counts the codes in F5:F11 one criterion per cell; C5 also heads a
    column that extracts text with MID. It is one of the counts, not a broken
    extraction. A fill that reads a shared table through SUMIF, with one member
    pointing at the wrong row, stays reported: its neighbours read the same
    table, so the table is not a range only the flagged cell summarises."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Patterns"
    ws["B4"], ws["C4"], ws["D4"] = "count of 1", "count of 2", "count of 3"
    ws["B5"], ws["C5"], ws["D5"] = "=COUNTIF(F5:F11,1)", "=COUNTIF(F5:F11,2)", "=COUNTIF(F5:F11,3)"
    for row in range(5, 12):
        ws[f"A{row}"], ws[f"F{row}"] = f"13c/15c ({row})", row % 3 + 1
    for row in range(6, 12):
        ws[f"C{row}"] = f'=MID(A{row},FIND("(",A{row})+1,FIND(")",A{row})-FIND("(",A{row})-1)'
    table = wb.create_sheet("Totals")
    for row in range(2, 21):
        table[f"H{row}"], table[f"I{row}"], table[f"J{row}"] = f"k{row}", row, row * 2
    for row in range(5, 12):
        table[f"A{row}"] = f"k{row}"
        key_row = row + 1 if row == 8 else row  # D8 looks up the wrong row
        table[f"D{row}"] = f"=SUMIF($H$2:$H$20,A{key_row},$I$2:$I$20)"
        table[f"E{row}"] = f"=SUMIF($H$2:$H$20,A{row},$J$2:$J$20)"
    path = tmp_path / "counts.xlsx"
    wb.save(path)
    assert [f["location"] for f in _audit(path).get("FORMULA_DRIFT", [])] == ["Totals!D8"]


def test_merged_formulas_in_totals_rows_are_presentation(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    for row in range(1, 5):
        ws.append([row, row * 2])
    ws.append(["=SUM(A1:A4)", "=SUM(B1:B4)"])
    ws.merge_cells("A6:B6")
    ws["A6"] = "=A5-B5"  # a net line merged across the two columns it nets
    ws["D1"] = "=SUM(A1:B6)"  # a range reads through the merge
    path = tmp_path / "net.xlsx"
    wb.save(path)
    assert "MERGED_CELL_IN_DATA_RANGE" not in _audit(path)


def test_whitespace_skips_headers_under_titles_and_deep_indentation_but_keeps_a_stray_space(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    ws.append(["Table 1"])
    ws.append(["ID ", "Value"])  # a header under a title
    ws.append(["a", 1])
    ws.append([" CA13 9BH", 2])  # one stray leading space on a key
    ws.append(["        STUFF 3", 3])  # indentation
    ws["D1"] = '=VLOOKUP("CA13 9BH",A3:B5,2,FALSE)'
    path = tmp_path / "keys.xlsx"
    wb.save(path)
    assert [f["location"] for f in _audit(path).get("WHITESPACE_KEY", [])] == ["S!A4"]


def test_array_formula_that_only_indexes_a_whole_column_is_quiet(tmp_path):
    from openpyxl.worksheet.formula import ArrayFormula

    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    for row in range(1, 6):
        ws.append([f"k{row}", row * 10])
    ws["D1"] = ArrayFormula("D1", '=IFERROR(INDEX($B:$B,SMALL(IF($A$1:$A$5="k3",ROW($A$1:$A$5)),1)),"")')
    ws["D2"] = ArrayFormula("D2", "=MAX(IF($A:$A=\"k2\",$B:$B,0))")  # the whole column really is the array
    path = tmp_path / "arrays.xlsx"
    wb.save(path)
    assert [f["location"] for f in _audit(path).get("WHOLE_COLUMN_REFERENCE", [])] == ["S!D2"]


def test_google_sheets_placeholders_are_flagged_as_a_coverage_gap(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "PAYMENT"
    ws["A1"] = "Size"
    ws["B1"] = '=IFERROR(__xludf.DUMMYFUNCTION("""COMPUTED_VALUE"""),"60*90")'
    path = tmp_path / "gsheets.xlsx"
    wb.save(path)
    result = subprocess.run(
        [sys.executable, "-m", "spreadsheet_auditor", str(path), "--json", "-", "--fail-on", "None"],
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
    )
    payload = json.loads(result.stdout)
    assert "google_sheets_placeholders" in payload["coverage"]["unsupported_features"]
    assert any("Google Sheets" in note for note in payload["coverage"]["limitations"])


def test_blank_precedent_ignores_rows_with_no_inputs_and_half_filled_ledgers(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    ws.append(["Day", "In", "Out", "Hours", "Debit", "Credit", "Balance"])
    times = [("08:00", "16:00"), (None, None), ("09:00", "17:00"), ("08:30", "16:30"), (None, None), ("09:00", "17:00")]
    ledger = [(100, None), (None, 40), (None, 60), (200, None), (None, 50), (None, 70)]
    for i, ((t_in, t_out), (debit, credit)) in enumerate(zip(times, ledger), start=2):
        prev = f"G{i - 1}" if i > 2 else "0"
        ws.append([i, t_in, t_out, f"=C{i}-B{i}", debit, credit, f"={prev}+E{i}-F{i}"])
    ws["I2"] = 5
    ws["I3"] = 7
    ws["I4"] = None  # the one real gap in a filled column
    ws["I5"] = 6
    ws["I6"] = 8
    ws["I7"] = 9
    ws["J4"] = "=I4*2"
    path = tmp_path / "ledger.xlsx"
    wb.save(path)
    assert [f["location"] for f in _audit(path).get("BLANK_PRECEDENT", [])] == ["S!J4"]


def test_countif_criteria_literals_are_not_assumptions(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    for row in range(1, 8):
        ws.append([row % 3])
    ws["C1"] = "=COUNTIF(A1:A7,2)"
    ws["C2"] = '=COUNTIFS(A1:A7,2,A1:A7,"<>1")'
    ws["C3"] = "=SUMIFS(A1:A7,A1:A7,2)"
    ws["C4"] = "=SUMIF(A1:A7,2)*1.3"  # the multiplier is the assumption
    path = tmp_path / "criteria.xlsx"
    wb.save(path)
    literals = _audit(path).get("LITERAL_CONSTANT", [])
    assert [f["location"] for f in literals] == ["S!C4"]
    assert "1.3" in literals[0]["evidence"][0]
