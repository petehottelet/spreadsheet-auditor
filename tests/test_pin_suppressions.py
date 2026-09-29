"""--pin-suppressions rewrites one-cell lines to the pinned form and keeps the rest of the file."""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook

from spreadsheet_auditor import audit


def _model(path: Path, offset: int = 0, new_error_row: int | None = None) -> Path:
    """Nine labeled lines with an accepted #REF! on line 5, shifted down by ``offset`` rows."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Model"
    for line in range(1, 10):
        ws.cell(line + offset, 1, f"Line {line}")
        ws.cell(line + offset, 2, line * 100)
    ws.cell(5 + offset, 3, "=SUM(#REF!)")
    if new_error_row:
        ws.cell(new_error_row, 3, "=B1/#REF!")
    wb.save(path)
    return path


def _run(workbook: Path, ignore: Path, *extra: str) -> dict:
    out = workbook.with_suffix(".json")
    audit.main([str(workbook), "--json", str(out), "--quiet", "--fail-on", "None", "--ignore", str(ignore), *extra])
    return json.loads(out.read_text(encoding="utf-8"))


def _live(payload: dict) -> dict[str, dict]:
    return {f["location"]: f for f in payload["findings"] if f["rule_id"] == "LIVE_ERROR"}


def test_pinning_rewrites_one_cell_lines_and_keeps_everything_else(tmp_path, capsys):
    ignore = tmp_path / ".audit-ignore"
    ignore.write_text(
        "# accepted items, reviewed 2026-09\n"
        "LIVE_ERROR Model!C5 accepted: legacy link\n"
        "BROKEN_REFERENCE Model!C1:C9 legacy block\n"
        "LIVE_ERROR Model!C7 fixed long ago\n"
        "\n",
        encoding="utf-8",
    )
    payload = _run(_model(tmp_path / "v1.xlsx"), ignore, "--pin-suppressions")
    fingerprint = _live(payload)["Model!C5"]["fingerprint"]
    assert ignore.read_text(encoding="utf-8") == (
        "# accepted items, reviewed 2026-09\n"
        f"LIVE_ERROR Model!C5 fingerprint:{fingerprint} accepted: legacy link\n"
        "BROKEN_REFERENCE Model!C1:C9 legacy block\n"
        "LIVE_ERROR Model!C7 fixed long ago\n"
        "\n"
    )
    stderr = capsys.readouterr().err
    assert "Pinned 1 line(s)" in stderr and "line 2: LIVE_ERROR Model!C5 accepted: legacy link" in stderr
    assert "(Model!C5: =SUM(#REF!))" in stderr
    notes = payload["coverage"]["limitations"]
    assert not [note for note in notes if "follows the address" in note]  # fixed by the rewrite
    assert [note for note in notes if "Model!C7 matched no finding" in note]  # stale lines are left to the user


def test_pinned_file_survives_the_row_insert_and_the_next_run_updates_the_address(tmp_path, capsys):
    ignore = tmp_path / ".audit-ignore"
    ignore.write_text("LIVE_ERROR Model!C5 accepted: legacy link\n", encoding="utf-8")
    fingerprint = _live(_run(_model(tmp_path / "v1.xlsx"), ignore, "--pin-suppressions"))["Model!C5"]["fingerprint"]

    # A title row goes in and a new #REF! is typed at C5.
    v2 = _model(tmp_path / "v2.xlsx", offset=1, new_error_row=5)
    live = _live(_run(v2, ignore))
    assert (live["Model!C6"]["suppressed"], live["Model!C5"]["suppressed"]) == (True, False)

    capsys.readouterr()
    _run(v2, ignore, "--pin-suppressions")
    assert ignore.read_text(encoding="utf-8") == f"LIVE_ERROR Model!C6 fingerprint:{fingerprint} accepted: legacy link\n"
    assert "line 1: LIVE_ERROR Model!C5 fingerprint:" in capsys.readouterr().err

    _run(v2, ignore, "--pin-suppressions")
    assert "already pinned and current" in capsys.readouterr().err


def test_pinning_keeps_crlf_line_endings(tmp_path):
    ignore = tmp_path / ".audit-ignore"
    ignore.write_bytes(b"# windows file\r\nLIVE_ERROR Model!C5 accepted\r\n")
    _run(_model(tmp_path / "v1.xlsx"), ignore, "--pin-suppressions")
    raw = ignore.read_bytes()
    assert raw.startswith(b"# windows file\r\nLIVE_ERROR Model!C5 fingerprint:") and raw.endswith(b" accepted\r\n")
    assert raw.count(b"\n") == raw.count(b"\r\n") == 2


def test_pinning_without_a_file_says_so(tmp_path, capsys):
    _run(_model(tmp_path / "v1.xlsx"), tmp_path / "missing-ignore", "--pin-suppressions")
    assert "nothing to pin" in capsys.readouterr().err
    assert not (tmp_path / "missing-ignore").exists()
