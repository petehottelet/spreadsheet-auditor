from __future__ import annotations

import re
from bisect import bisect_left
from collections import Counter, defaultdict

from openpyxl.utils.cell import column_index_from_string, get_column_letter

from .budget import tick
from .finding import Finding
from .formula_parser import formula_text, is_formula, normalize_formula, parse_formula
from .range_checks import is_total_of_segment
from .reference_resolver import boundaries, cell_value
from .workbook_inventory import iter_existing_cells, location

# A plug is a few constants interrupting a formula run, not a block of inputs.
MAX_PLUG_GAP = 3


def _is_plain_number(value) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def detect_formula_drift(formula_cells: list[dict], budget=None, names=None, formula_wb=None) -> list[Finding]:
    findings: list[Finding] = []
    seen: set[str] = set()

    by_row: dict[tuple[str, int], list[dict]] = defaultdict(list)
    by_col: dict[tuple[str, int], list[dict]] = defaultdict(list)
    pattern_at: dict[tuple[str, int, int], str] = {}
    cell_at: dict[tuple[str, int, int], dict] = {}
    for cell in formula_cells:
        by_row[(cell["sheet"], cell["row"])].append(cell)
        by_col[(cell["sheet"], cell["col"])].append(cell)
        pattern_at[(cell["sheet"], cell["row"], cell["col"])] = normalize_formula(
            cell["formula"], cell["row"], cell["col"]
        )
        cell_at[(cell["sheet"], cell["row"], cell["col"])] = cell

    groups = [(values, "row") for values in by_row.values()]
    groups += [(values, "column") for values in by_col.values()]
    for cells, axis in groups:
        key = "col" if axis == "row" else "row"
        for segment in _contiguous_segments(sorted(cells, key=lambda c: c[key]), key):
            tick(budget)
            if len(segment) < 3:
                continue
            fixed = segment[0]["row"] if axis == "row" else segment[0]["col"]
            low, high = segment[0][key], segment[-1][key]
            members = [c for c in segment if not is_total_of_segment(c, axis, fixed, low, high, names)]
            if len(members) < 3:
                continue
            patterns = [pattern_at[(c["sheet"], c["row"], c["col"])] for c in members]
            counts = Counter(patterns)
            majority, count = counts.most_common(1)[0]
            if "#REF!" in majority:
                continue
            pairs = list(zip(members, patterns))
            # The seed of a chain (an opening balance, year 0 of a projection)
            # links to its source while every later cell builds on the
            # previous one; it is expected to differ and is not a member.
            if pairs[0][1] != majority and _chain_step(majority, axis, -1):
                pairs = pairs[1:]
            if pairs and pairs[-1][1] != majority and _chain_step(majority, axis, 1):
                pairs = pairs[:-1]
            if len(pairs) < 3 or count < len(pairs) - 1:
                continue
            for cell, pattern in pairs:
                if pattern == majority or cell["location"] in seen:
                    continue
                if _AGG_RE.sub("AGG(", pattern) == _AGG_RE.sub("AGG(", majority) and (
                    len(pairs) <= 3 or _header_names_function(formula_wb, cell, pattern, axis)
                ):
                    # A totals row where one column sums and the next averages
                    # the same block is a design when the row is short or the
                    # header says so ("Sum of ...", "Average of ..."); a SUM
                    # among a run of SUMs that turned into an AVERAGE is the
                    # classic function slip and stays reported.
                    continue
                if _conforms_across(cell, pattern, majority, axis, pattern_at):
                    continue
                peer = min(
                    (member for member, member_pattern in pairs if member_pattern == majority),
                    key=lambda member: abs(member[key] - cell[key]),
                )
                if _fill_adds_up_labels(cell, peer, formula_wb, names):
                    continue
                if _in_summary_set(cell, pairs, majority, axis, cell_at, names):
                    continue
                seen.add(cell["location"])
                neighbors = _drift_neighbors(members, members.index(cell))
                evidence = [
                    f"Formula differs from the dominant relative pattern in this {axis}.",
                    f"Dominant pattern (n={count}): {majority}",
                    f"This cell: {pattern}",
                ]
                if neighbors:
                    evidence.append("Neighboring formulas: " + "; ".join(neighbors))
                suggested_fix = (
                    "If this cell is the same kind of line as its neighbors, restore their pattern. "
                    "A total, net or summary line differs by design: compare it with the other totals instead."
                )
                own = _own_line_total(
                    cell, by_row[(cell["sheet"], cell["row"])], by_col[(cell["sheet"], cell["col"])]
                )
                if own is not None:
                    total, unit, own_formula = own
                    line = "row" if unit == "column" else "column"
                    evidence.append(
                        f"The other totals on this {line} add up their own {unit} "
                        f"({total['location']}={total['formula']}); this one adds a different {unit}."
                    )
                    suggested_fix = f"Point the total at its own {unit}: {own_formula}"
                findings.append(
                    Finding(
                        rule_id="FORMULA_DRIFT",
                        severity="High",
                        error_confidence="Likely defect",
                        detection_mode="DET",
                        location=cell["location"],
                        title="Formula breaks neighboring pattern",
                        formula=cell["formula"],
                        evidence=evidence,
                        suggested_fix=suggested_fix,
                    )
                )
    return findings


_OFFSET_RE = re.compile(r"\[-?\d+\]")
_AGG_RE = re.compile(r"\b(?:SUM|AVERAGE|AVERAGEA|COUNT|COUNTA|MIN|MAX|MEDIAN|SUBTOTAL)\(")
# A formula that is one aggregate over one same-sheet range, such as =SUM(G65:G122).
_SIMPLE_AGG_RE = re.compile(
    r"^=\s*(SUM|AVERAGE|AVERAGEA|COUNT|COUNTA|MIN|MAX|MEDIAN)\(\s*"
    r"\$?([A-Z]{1,3})\$?(\d+)\s*:\s*\$?([A-Z]{1,3})\$?(\d+)\s*\)\s*$",
    re.IGNORECASE,
)


# Moving the fill's formula to a flagged cell reads each cell of its ranges;
# larger ranges are not judged (the cell stays a drift candidate).
MAX_MOVED_RANGE_CELLS = 10_000
_A1_POINT_RE = re.compile(r"^(\$?)([A-Za-z]{1,3})(\$?)(\d+)$")
_NUMERIC_TEXT_RE = re.compile(r"^[\s$\u20ac\u00a3(+-]*\d[\d,.\s]*%?\)?\s*$")
# Functions that summarise a block: a row of them over one range, each with
# its own criterion, is a set of summaries rather than a fill.
_SUMMARY_FUNCTIONS = {
    "AVERAGE", "AVERAGEIF", "AVERAGEIFS", "COUNT", "COUNTA", "COUNTBLANK", "COUNTIF", "COUNTIFS",
    "LARGE", "MAX", "MAXIFS", "MEDIAN", "MIN", "MINIFS", "SMALL", "SUM", "SUMIF", "SUMIFS", "SUMPRODUCT",
}


def _moved_point(text: str, drow: int, dcol: int) -> tuple[int, int] | None:
    """``(row, col)`` of an A1 endpoint moved by the offset, honouring ``$`` markers."""
    match = _A1_POINT_RE.match(text)
    if match is None:
        return None
    col_absolute, letters, row_absolute, digits = match.groups()
    col = column_index_from_string(letters.upper()) + (0 if col_absolute else dcol)
    row = int(digits) + (0 if row_absolute else drow)
    return (row, col) if row >= 1 and col >= 1 else None


def _fill_adds_up_labels(cell: dict, peer: dict, formula_wb, names) -> bool:
    """True when the fill's formula, moved to this cell, would aggregate nothing but text labels.

    A summary box heads columns L, M and N: M2 = SUM(M4:M11) and
    N2 = SUM(N4:N11) total their columns, while L2 = M14 points at another
    total. At L2 the fill's formula would be SUM(L4:L11), a column of names,
    so L2 cannot belong to that fill. A total aimed at the wrong column stays
    reported: there the moved formula lands on numbers or on empty cells.
    """
    if formula_wb is None:
        return False
    try:
        ws = formula_wb[cell["sheet"]]
    except KeyError:
        return False
    parsed = parse_formula(peer["formula"], names=names, origin=(peer["sheet"], peer["row"], peer["col"]))
    drow, dcol = cell["row"] - peer["row"], cell["col"] - peer["col"]
    ranges = 0
    for ref in parsed.references:
        if not ref.is_range:
            continue
        if not ref.bounded or (ref.sheet is not None and ref.sheet.casefold() != cell["sheet"].casefold()):
            return False
        ends = ref.raw.rsplit("!", 1)[-1].split(":")
        if len(ends) != 2:
            return False
        first, last = _moved_point(ends[0], drow, dcol), _moved_point(ends[1], drow, dcol)
        if first is None or last is None:
            return False
        rows = range(min(first[0], last[0]), max(first[0], last[0]) + 1)
        cols = range(min(first[1], last[1]), max(first[1], last[1]) + 1)
        if len(rows) * len(cols) > MAX_MOVED_RANGE_CELLS:
            return False
        labels = 0
        for row in rows:
            for col in cols:
                value = cell_value(ws, row, col)
                if value is None or (isinstance(value, str) and not value.strip()):
                    continue
                if isinstance(value, str) and not is_formula(value) and not _NUMERIC_TEXT_RE.match(value):
                    labels += 1
                    continue
                return False  # a number, a date, a formula or numeric text: the fill has data here
        if labels == 0:
            return False
        ranges += 1
    return ranges > 0


def _range_boxes(cell: dict, names) -> tuple[set[tuple[int, int, int, int]], set[str]]:
    """The bounded same-sheet multi-cell ranges a formula reads, and the functions it calls."""
    parsed = parse_formula(cell["formula"], names=names, origin=(cell["sheet"], cell["row"], cell["col"]))
    boxes = set()
    for ref in parsed.references:
        if not ref.is_range or not ref.bounded:
            continue
        if ref.sheet is not None and ref.sheet.casefold() != cell["sheet"].casefold():
            continue
        box = boundaries(ref.ref)
        if box is not None and (box[0], box[1]) != (box[2], box[3]):
            boxes.add(box)
    return boxes, set(parsed.functions)


def _in_summary_set(cell: dict, pairs: list, majority: str, axis: str, cell_at: dict, names) -> bool:
    """True when the cell summarises a block the fill never reads, beside a cell that summarises it the same way.

    Row 5 holds COUNTIF(F5:F11,1), COUNTIF(F5:F11,2) and COUNTIF(F5:F11,3),
    a set of counts over one block. C5 also heads a column whose rows extract
    text with MID, so the column scan sees it as the odd one out; it is one of
    the counts. A fill that reads a shared lookup table is not a set: its
    members read the same range, so a broken member stays reported.
    """
    boxes, functions = _range_boxes(cell, names)
    if not boxes or not functions & _SUMMARY_FUNCTIONS:
        return False
    fill_boxes: set[tuple[int, int, int, int]] = set()
    for member, pattern in pairs:
        if pattern == majority:
            fill_boxes |= _range_boxes(member, names)[0]
    own = boxes - fill_boxes
    if not own:
        return False
    if axis == "column":
        across = [(cell["row"], cell["col"] - 1), (cell["row"], cell["col"] + 1)]
    else:
        across = [(cell["row"] - 1, cell["col"]), (cell["row"] + 1, cell["col"])]
    for row, col in across:
        other = cell_at.get((cell["sheet"], row, col))
        if other is None:
            continue
        other_boxes, other_functions = _range_boxes(other, names)
        if other_functions == functions and own & other_boxes:
            return True
    return False


def _simple_aggregate(formula: str) -> tuple[str, int, int, int, int] | None:
    """``(function, min_col, min_row, max_col, max_row)`` for ``=AGG(A1:B2)``, else None."""
    match = _SIMPLE_AGG_RE.match(formula or "")
    if not match:
        return None
    function, col1, row1, col2, row2 = match.groups()
    c1, c2 = column_index_from_string(col1.upper()), column_index_from_string(col2.upper())
    r1, r2 = int(row1), int(row2)
    return function.upper(), min(c1, c2), min(r1, r2), max(c1, c2), max(r1, r2)


def _own_line_total(cell: dict, same_row: list[dict], same_col: list[dict]) -> tuple[dict, str, str] | None:
    """``(peer, unit, formula)`` when a drifted total adds another column (or row) than its peers do.

    A totals row where E123 = SUM(E65:E122) adds its own column but
    F123 = SUM(G65:G122) adds column G is one total pointed at the wrong
    column; the fix is SUM(F65:F122), not the row sum that column F's
    majority pattern would suggest. The same holds across a totals column.
    """
    agg = _simple_aggregate(cell["formula"])
    if agg is None:
        return None
    function, c1, r1, c2, r2 = agg
    if c1 == c2 != cell["col"] and r2 < cell["row"]:
        for peer in same_row:
            other = _simple_aggregate(peer["formula"]) if peer is not cell else None
            if other and other[1] == other[3] == peer["col"] and (other[2], other[4]) == (r1, r2):
                letter = get_column_letter(cell["col"])
                return peer, "column", f"={function}({letter}{r1}:{letter}{r2})"
    if r1 == r2 != cell["row"] and c2 < cell["col"]:
        for peer in same_col:
            other = _simple_aggregate(peer["formula"]) if peer is not cell else None
            if other and other[2] == other[4] == peer["row"] and (other[1], other[3]) == (c1, c2):
                first, last = get_column_letter(c1), get_column_letter(c2)
                return peer, "row", f"={function}({first}{cell['row']}:{last}{cell['row']})"
    return None
_AGG_NAME_RE = re.compile(r"\b(SUM|AVERAGE|AVERAGEA|COUNT|COUNTA|MIN|MAX|MEDIAN)\(")
# Words in a header or label that announce which aggregate a cell holds.
_AGG_WORDS = {
    "SUM": ("sum", "total"),
    "AVERAGE": ("average", "avg", "mean"),
    "AVERAGEA": ("average", "avg", "mean"),
    "COUNT": ("count", "number of", "no. of", "no of"),
    "COUNTA": ("count", "number of", "no. of", "no of"),
    "MIN": ("min", "lowest", "smallest"),
    "MAX": ("max", "highest", "largest", "peak"),
    "MEDIAN": ("median",),
}


def _header_names_function(formula_wb, cell: dict, pattern: str, axis: str) -> bool:
    """True when the header above (row scan) or the label beside (column scan) names the cell's aggregate.

    A pivot-style totals row headed "Sum of Gen." beside "Average of %PLF"
    mixes SUM and AVERAGE by design; the headers say which is which.
    """
    if formula_wb is None:
        return False
    functions = {name for name in _AGG_NAME_RE.findall(pattern)}
    if not functions:
        return False
    try:
        ws = formula_wb[cell["sheet"]]
    except KeyError:
        return False
    texts: list[str] = []
    if axis == "row":  # the deviant sits in a row of totals; its column header names it
        # Walk to the top: stacked pivot blocks share one header row far above.
        for row in range(cell["row"] - 1, 0, -1):
            value = cell_value(ws, row, cell["col"])
            if isinstance(value, str) and not is_formula(value) and value.strip():
                texts.append(value)
                break
    else:  # the deviant sits in a column of totals; its row label names it
        for col in range(cell["col"] - 1, 0, -1):
            value = cell_value(ws, cell["row"], col)
            if isinstance(value, str) and not is_formula(value) and value.strip():
                texts.append(value)
                break
    label = " ".join(texts).lower()
    return any(word in label for name in functions for word in _AGG_WORDS.get(name, ()))


def _shape(pattern: str) -> str:
    """A pattern with its relative offsets blanked: what the formula does, not where it points."""
    return _OFFSET_RE.sub("[]", pattern)


def _conforms_across(
    cell: dict, pattern: str, majority: str, axis: str, pattern_at: dict[tuple[str, int, int], str]
) -> bool:
    """True when the cell's formula matches a formula neighbour on the other axis.

    A row of unrelated columns (a label link, a derived column, a mirror
    column) is not a fill. The cell that differs from the two beside it but
    agrees with the cell above or below it is a different column by design,
    not a broken pattern. When the cell is a structurally different formula
    from the majority (not the same formula with a slipped offset), agreeing
    in shape with its neighbour is enough: a label column that links every
    second row of another sheet never repeats an exact pattern. The price is
    a column filled down wrongly in every row, which then reads as
    consistent; the column scan still sees it when its own pattern breaks.
    """
    if axis == "row":
        keys = [(cell["sheet"], cell["row"] - 1, cell["col"]), (cell["sheet"], cell["row"] + 1, cell["col"])]
    else:
        keys = [(cell["sheet"], cell["row"], cell["col"] - 1), (cell["sheet"], cell["row"], cell["col"] + 1)]
    found = [pattern_at[key] for key in keys if key in pattern_at]
    if any(other == pattern for other in found):
        return True
    shape = _shape(pattern)
    return shape != _shape(majority) and any(_shape(other) == shape for other in found)


def _chain_step(majority: str, axis: str, direction: int) -> bool:
    """True when the dominant pattern reads the neighbouring cell one step back along the axis.

    ``direction`` -1 means each cell builds on its predecessor (a running
    balance, a projection), so the first cell of the run is the seed; 1 means
    the run chains the other way and the last cell is the seed.
    """
    if axis == "row":
        token = "RC[-1]" if direction < 0 else "RC[1]"
    else:
        token = "R[-1]C" if direction < 0 else "R[1]C"
    return token in majority


def _drift_neighbors(segment: list[dict], idx: int, radius: int = 1) -> list[str]:
    """Return up to 2*radius adjacent formulas (with their locations) for evidence."""
    neighbors: list[str] = []
    start = max(0, idx - radius)
    end = min(len(segment), idx + radius + 1)
    for j in range(start, end):
        if j == idx:
            continue
        cell = segment[j]
        neighbors.append(f"{cell['location']}={cell['formula']}")
    return neighbors


def _aggregates_over(text: str, origin: tuple[int, int], sheet: str, target: tuple[int, int], names) -> bool:
    """True when the formula at ``origin`` sums a range that contains ``target``."""
    parsed = parse_formula(text, names=names, origin=(sheet, origin[0], origin[1]))
    for ref in parsed.references:
        if not ref.is_range or not ref.bounded:
            continue
        if (ref.sheet or sheet).casefold() != sheet.casefold():
            continue
        box = boundaries(ref.ref)
        if box is None:
            continue
        min_col, min_row, max_col, max_row = box
        if min_row <= target[0] <= max_row and min_col <= target[1] <= max_col:
            return True
    return False


def detect_hardcode_breaks(
    formula_wb,
    allowed_sheet_names: set[str] | None = None,
    budget=None,
    names=None,
) -> list[Finding]:
    """Flag constants that interrupt a run of formulas sharing one relative pattern.

    A plug replaces a formula: the formulas on either side of it (within a
    short gap) normalize to the same pattern. Inputs that an adjacent subtotal
    sums are not plugs, however many formulas surround them.
    """
    findings: list[Finding] = []
    seen: set[str] = set()
    for ws in formula_wb.worksheets:
        if allowed_sheet_names is not None and ws.title not in allowed_sheet_names:
            continue
        by_row: dict[int, dict[int, object]] = defaultdict(dict)
        by_col: dict[int, dict[int, object]] = defaultdict(dict)
        for cell in iter_existing_cells(ws):
            if cell.value is None:
                continue
            by_row[cell.row][cell.column] = cell.value
            by_col[cell.column][cell.row] = cell.value
        for row, values in by_row.items():
            tick(budget)
            _hardcode_breaks_in_line(values, ws.title, row, None, seen, findings, names, by_row)
        for col, values in by_col.items():
            tick(budget)
            _hardcode_breaks_in_line(values, ws.title, None, col, seen, findings, names, by_row)
    return findings


def _is_input_line(
    by_row: dict[int, dict[int, object]],
    cell_pos: tuple[int, int],
    left_pos: tuple[int, int],
    right_pos: tuple[int, int],
) -> bool:
    """True when the constant belongs to a line of inputs between two lines of formulas.

    An input column between two running-sum columns, or an input row between
    two ratio rows, puts a constant between two formulas that share a pattern
    on every line. The tell is the other axis: the constant's neighbours
    there are constants too, while the formulas' neighbours are formulas.
    """

    def at(row: int, col: int):
        return by_row.get(row, {}).get(col)

    row, col = cell_pos
    across = left_pos[0] == row  # the formula pair lies in this row
    evaluated = 0
    formulas = 0
    for step in (-1, 1):
        if across:  # look up and down
            beside = at(row + step, col)
            corners = [at(row + step, left_pos[1]), at(row + step, right_pos[1])]
        else:  # look left and right
            beside = at(row, col + step)
            corners = [at(left_pos[0], col + step), at(right_pos[0], col + step)]
        # A header, a label, a total row or the edge of the block beside the
        # constant says nothing either way; a constant beside it does, when
        # the formula lines continue there too.
        if _is_plain_number(beside):
            evaluated += len(corners)
            formulas += sum(1 for v in corners if is_formula(v))
    if evaluated == 0:
        return False
    # One of four corners may itself hold a plug in the neighbouring formula line.
    return formulas == evaluated or (evaluated >= 4 and formulas >= evaluated - 1)


def _hardcode_breaks_in_line(
    values: dict[int, object],
    sheet: str,
    row: int | None,
    col: int | None,
    seen: set[str],
    findings: list[Finding],
    names,
    by_row: dict[int, dict[int, object]],
) -> None:
    keys = sorted(values)
    formula_keys = [key for key in keys if is_formula(values[key])]
    if len(formula_keys) < 2:
        return

    def position(key: int) -> tuple[int, int]:
        return (row, key) if row is not None else (key, col)  # type: ignore[return-value]

    for key in keys:
        value = values[key]
        if not _is_plain_number(value):
            continue
        idx = bisect_left(formula_keys, key)
        if idx == 0 or idx >= len(formula_keys):
            continue
        left, right = formula_keys[idx - 1], formula_keys[idx]
        if right - left - 1 > MAX_PLUG_GAP:
            continue
        if any(isinstance(values.get(k), str) and not is_formula(values[k]) for k in range(left + 1, right)):
            continue  # a text label between the two formulas ends one block and starts another
        left_pos, right_pos, cell_pos = position(left), position(right), position(key)
        left_text = formula_text(values[left])
        right_text = formula_text(values[right])
        if left_text is None or right_text is None:
            continue
        pattern = normalize_formula(left_text, *left_pos)
        if pattern != normalize_formula(right_text, *right_pos):
            continue
        if _aggregates_over(left_text, left_pos, sheet, cell_pos, names) or _aggregates_over(
            right_text, right_pos, sheet, cell_pos, names
        ):
            continue
        if _is_input_line(by_row, cell_pos, left_pos, right_pos):
            continue
        loc = location(sheet, *cell_pos)
        if loc in seen:
            continue
        seen.add(loc)
        findings.append(
            Finding(
                rule_id="HARDCODE_IN_FORMULA_BLOCK",
                severity="High",
                error_confidence="Likely defect",
                detection_mode="DET",
                location=loc,
                title="Hardcoded value inside formula block",
                formula=None,
                evidence=[
                    f"Value {value!r} sits between {location(sheet, *left_pos)} and {location(sheet, *right_pos)}, "
                    f"which share the pattern {pattern}."
                ],
                suggested_fix="Confirm whether this is an intentional plug. If not, restore the formula pattern.",
            )
        )


def _contiguous_segments(cells: list[dict], key: str) -> list[list[dict]]:
    if not cells:
        return []
    segments = [[cells[0]]]
    for cell in cells[1:]:
        if cell[key] == segments[-1][-1][key] + 1:
            segments[-1].append(cell)
        else:
            segments.append([cell])
    return segments
