from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook
from openpyxl.comments import Comment


def _anchor_location(location: str) -> tuple[str, str] | None:
    """Resolve a finding location to a single (sheet, coordinate) anchor cell.

    Handles comma-joined multi-locations (uses the first) and range locations
    (uses the top-left cell), e.g. "Sheet1!A1:B2" -> ("Sheet1", "A1").
    """
    if not location:
        return None
    first = location.split(",", 1)[0].strip()
    if "!" not in first:
        return None
    sheet_name, coord = first.split("!", 1)
    coord = coord.split(":", 1)[0].replace("$", "").strip()
    if not sheet_name or not coord:
        return None
    return sheet_name, coord


def annotate_workbook(source_path: str | Path, output_path: str | Path, findings: list[dict]) -> None:
    source = Path(source_path)
    keep_vba = source.suffix.lower() == ".xlsm"
    wb = load_workbook(source, keep_vba=keep_vba)
    # One comment per cell listing every finding there, in payload order
    # (most severe first); assigning a comment per finding kept only the last.
    notes: dict[tuple[str, str], list[str]] = {}
    for finding in findings:
        if finding.get("suppressed"):
            continue
        location = finding.get("location", "")
        anchor = _anchor_location(location)
        if anchor is None:
            continue
        sheet_name, coord = anchor
        if sheet_name not in wb.sheetnames:
            continue
        notes.setdefault((sheet_name, coord), []).append(
            f"{finding.get('severity')} {finding.get('rule_id')}\n"
            f"{finding.get('title')}\n"
            f"Fix: {finding.get('suggested_fix')}"
        )
    for (sheet_name, coord), texts in notes.items():
        wb[sheet_name][coord].comment = Comment("\n\n".join(texts), "Spreadsheet Auditor")
    wb.save(output_path)
