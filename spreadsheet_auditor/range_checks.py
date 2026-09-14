from __future__ import annotations

from collections import Counter, defaultdict

from openpyxl.formula.tokenizer import Token
from openpyxl.utils.cell import column_index_from_string, get_column_letter

from .budget import tick
from .finding import Finding
from .formula_parser import (
    FUNCTION_PREFIXES,
    ParsedReference,
    formula_text,
    is_formula,
    parse_formula,
    split_sheet,
    tokens,
)
from .grouping import group_by_pattern, pattern_note
from .reference_resolver import (
    EXCEL_MAX_COL,
    EXCEL_MAX_ROW,
    boundaries,
    cell_value,
    find_sheet,
    raw_boundaries,
)


AGG_FUNCS = {"SUM", "AVERAGE", "COUNT", "COUNTA"}
SUBTOTAL_WORDS = {"total", "subtotal", "grand total", "sum"}
HEADER_SCAN_ROWS = 10
MAX_DEPENDENTS_SHOWN = 8

Box = tuple[int, int, int, int]


def _origin(cell: dict) -> tuple[str, int, int]:
    return (cell["sheet"], cell["row"], cell["col"])


def aggregate_ranges(formula: str, names=None, origin: tuple[str, int, int] | None = None) -> list[ParsedReference]:
    parsed = parse_formula(formula, names=names, origin=origin)
    if not parsed.functions.intersection(AGG_FUNCS):
        return []
    return [ref for ref in parsed.references if ref.is_range]


def _single_aggregate_range_size(formula: str, names=None, origin=None) -> int | None:
    ranges = aggregate_ranges(formula, names=names, origin=origin)
    if len(ranges) != 1:
        return None
    box = boundaries(ranges[0].ref)
    if box is None:
        return None
    min_col, min_row, max_col, max_row = box
    return (max_col - min_col + 1) * (max_row - min_row + 1)


def is_total_of_segment(cell: dict, axis: str, fixed: int, low: int, high: int, names=None) -> bool:
    """True when ``cell`` aggregates at least two other cells of its own segment.

    A row of column subtotals normally ends in a row total whose formula sums
    the very cells beside it. That total is a different kind of formula: it is
    not drift, its range length legitimately differs from its peers, and it
    must not dilute the majority pattern. ``axis`` is ``"row"`` (segment cells
    share ``row == fixed`` across columns ``low..high``) or ``"column"``.
    """
    parsed = parse_formula(cell["formula"], names=names, origin=_origin(cell))
    hits = 0
    for ref in parsed.references:
        if (ref.sheet or cell["sheet"]).casefold() != cell["sheet"].casefold():
            continue
        box = boundaries(ref.ref)
        if box is None:
            continue
        min_col, min_row, max_col, max_row = box
        if axis == "row":
            if not min_row <= fixed <= max_row:
                continue
            overlap = min(max_col, high) - max(min_col, low) + 1
            if min_col <= cell["col"] <= max_col:
                overlap -= 1
        else:
            if not min_col <= fixed <= max_col:
                continue
            overlap = min(max_row, high) - max(min_row, low) + 1
            if min_row <= cell["row"] <= max_row:
                overlap -= 1
        hits += max(0, overlap)
        if hits >= 2:
            return True
    return False


def _contiguous_by(items: list[tuple[dict, int]], key: str) -> list[list[tuple[dict, int]]]:
    if not items:
        return []
    ordered = sorted(items, key=lambda it: it[0][key])
    segments = [[ordered[0]]]
    for item in ordered[1:]:
        if item[0][key] == segments[-1][-1][0][key] + 1:
            segments[-1].append(item)
        else:
            segments.append([item])
    return segments


def _single_aggregate_box(formula: str, names=None, origin=None) -> Box | None:
    ranges = aggregate_ranges(formula, names=names, origin=origin)
    if len(ranges) != 1:
        return None
    return boundaries(ranges[0].ref)


def _boxes_overlap(a: Box, b: Box) -> bool:
    return a[0] <= b[2] and a[2] >= b[0] and a[1] <= b[3] and a[3] >= b[1]


def _peers_overlap(a: Box, b: Box) -> bool:
    """Two aggregate ranges cover the same group when they overlap along the axis they run on.

    Row sums in one column each cover their own row, so they are compared on
    columns; column sums in one row are compared on rows. A debit sum over
    C:I beside credit sums over J:AI shares no columns and is not a peer.
    """
    a_across = (a[2] - a[0]) >= (a[3] - a[1])
    b_across = (b[2] - b[0]) >= (b[3] - b[1])
    if a_across and b_across:
        return a[0] <= b[2] and a[2] >= b[0]
    if not a_across and not b_across:
        return a[1] <= b[3] and a[3] >= b[1]
    return _boxes_overlap(a, b)


def detect_range_length_mismatch(formula_cells: list[dict], names=None, budget=None) -> list[Finding]:
    findings: list[Finding] = []
    seen: set[str] = set()

    sized: list[tuple[dict, int, Box]] = []
    for cell in formula_cells:
        tick(budget)
        box = _single_aggregate_box(cell["formula"], names=names, origin=_origin(cell))
        if box is None:
            continue
        size = (box[2] - box[0] + 1) * (box[3] - box[1] + 1)
        if size > 1:  # SUM(G213:G213) is a link to one cell, not an aggregate with peers
            sized.append((cell, size, box))

    by_row: dict[tuple[str, int], list[tuple[dict, int]]] = defaultdict(list)
    by_col: dict[tuple[str, int], list[tuple[dict, int]]] = defaultdict(list)
    box_of: dict[str, Box] = {}
    for cell, size, box in sized:
        box_of[cell["location"]] = box
        by_row[(cell["sheet"], cell["row"])].append((cell, size))
        by_col[(cell["sheet"], cell["col"])].append((cell, size))

    groups = [(items, "col", "row") for items in by_row.values()]
    groups += [(items, "row", "column") for items in by_col.values()]

    for items, sort_key, axis_name in groups:
        for segment in _contiguous_by(items, sort_key):
            if len(segment) < 3:
                continue
            first, last = segment[0][0], segment[-1][0]
            fixed = first["row"] if axis_name == "row" else first["col"]
            members = [
                (cell, size)
                for cell, size in segment
                if not is_total_of_segment(cell, axis_name, fixed, first[sort_key], last[sort_key], names)
            ]
            if len(members) < 3:
                continue
            sizes = [size for _, size in members]
            counts = Counter(sizes)
            majority, count = counts.most_common(1)[0]
            if count < len(members) - 1:
                continue
            majority_boxes = [box_of[cell["location"]] for cell, size in members if size == majority]
            for cell, size in members:
                if size == majority or cell["location"] in seen:
                    continue
                if not any(_peers_overlap(box_of[cell["location"]], other) for other in majority_boxes):
                    continue  # sums over disjoint column groups (debits beside credits) are not peers
                seen.add(cell["location"])
                findings.append(
                    Finding(
                        rule_id="RANGE_LENGTH_MISMATCH",
                        severity="High",
                        error_confidence="Likely defect",
                        detection_mode="DET",
                        location=cell["location"],
                        title="Aggregate range length differs from peer formulas",
                        formula=cell["formula"],
                        evidence=[
                            f"This aggregate spans {size} cell(s) while peer formulas in this {axis_name} span {majority}."
                        ],
                        suggested_fix="Compare this aggregate range to adjacent peers and align the range boundaries unless the difference is intentional.",
                    )
                )
    return findings


# --- exclusion and subtotal inclusion ----------------------------------------


def _is_plain_number(value) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def _has_subtotal_word(label: str) -> bool:
    return any(word in label for word in SUBTOTAL_WORDS)


def _row_label(ws, row: int, min_col: int) -> str:
    """Text labels to the left of a range on ``row`` (formulas and numbers ignored)."""
    labels = []
    for col in range(1, min_col):
        value = cell_value(ws, row, col)
        if isinstance(value, str) and not is_formula(value) and value.strip():
            labels.append(value.strip())
    return " ".join(labels).lower()


def _column_header(ws, col: int, above_row: int) -> str:
    """Nearest text cell above ``above_row`` in ``col``; empty if the first non-blank is not text."""
    for row in range(above_row - 1, max(0, above_row - HEADER_SCAN_ROWS - 1), -1):
        value = cell_value(ws, row, col)
        if value is None:
            continue
        if isinstance(value, str) and not is_formula(value):
            return value.strip().lower()
        return ""
    return ""


def _formula_overlaps(sheet_title: str, value, origin_row: int, origin_col: int, outer: Box, names) -> bool:
    """True when the formula aggregates cells that also lie inside ``outer`` on the same sheet."""
    text = formula_text(value)
    if text is None:
        return False
    parsed = parse_formula(text, names=names, origin=(sheet_title, origin_row, origin_col))
    for ref in parsed.references:
        if (ref.sheet or sheet_title).casefold() != sheet_title.casefold():
            continue
        box = boundaries(ref.ref)
        if box is None:
            continue
        if box[0] <= outer[2] and box[2] >= outer[0] and box[1] <= outer[3] and box[3] >= outer[1]:
            return True
    return False


def _is_self(ws, formula_cell: dict, row: int, col: int) -> bool:
    return (
        ws.title.casefold() == formula_cell["sheet"].casefold()
        and row == formula_cell["row"]
        and col == formula_cell["col"]
    )


def detect_range_issues(formula_wb, value_wb, formula_cells: list[dict], names=None, budget=None) -> list[Finding]:
    findings: list[Finding] = []
    for cell in formula_cells:
        tick(budget)
        for ref in aggregate_ranges(cell["formula"], names=names, origin=_origin(cell)):
            if not ref.bounded:
                continue
            title = find_sheet(formula_wb, ref.sheet or cell["sheet"])
            if title is None:
                continue
            ws = formula_wb[title]
            box = boundaries(ref.ref)
            if box is None:
                continue
            if not _is_expanding(ref.raw):
                findings.extend(_detect_exclusion(ws, cell, ref.ref, box))
            findings.extend(_detect_subtotal_inclusion(ws, cell, ref.ref, box, names))
    return findings


def _is_expanding(raw: str) -> bool:
    """``$A$9:A9``: an anchored start and a relative end, the running-total idiom.

    Filled down, the range grows one cell per row; the single cell it covers
    in its first row is not a total that stopped short of its neighbour.
    """
    _, rest = split_sheet(raw)
    if ":" not in rest:
        return False
    start, end = rest.split(":", 1)
    return start.count("$") == 2 and "$" not in end


def _detect_exclusion(ws, formula_cell: dict, ref_text: str, box: Box) -> list[Finding]:
    """Flag a plain number that sits between a total and the range it aggregates.

    The classic off-by-one leaves the last input just outside the range, right
    between the block and the total below (or beside) it. Only that geometry is
    reported: text labels, headers, dates, and formula rows are never treated
    as excluded data, and totals that live elsewhere (another sheet, another
    column) do not produce exclusion findings.
    """
    min_col, min_row, max_col, max_row = box
    findings: list[Finding] = []
    same_sheet = ws.title.casefold() == formula_cell["sheet"].casefold()
    if not same_sheet:
        return findings
    loc = formula_cell["location"]
    formula = formula_cell["formula"]
    frow, fcol = formula_cell["row"], formula_cell["col"]

    if min_col == max_col and fcol == min_col:
        candidates: list[tuple[int, int, str]] = []
        if frow > max_row + 1:
            candidates.append((max_row + 1, min_col, "below"))
        elif frow < min_row - 1:
            candidates.append((min_row - 1, min_col, "above"))
        for row, col, direction in candidates:
            if _is_self(ws, formula_cell, row, col):
                continue
            value = cell_value(ws, row, col)
            if not _is_plain_number(value):
                continue
            if _has_subtotal_word(_row_label(ws, row, min_col)):
                continue
            coord = f"{get_column_letter(col)}{row}"
            findings.append(
                Finding(
                    rule_id="RANGE_EXCLUSION",
                    severity="Critical",
                    error_confidence="Likely defect",
                    detection_mode="DET",
                    location=loc,
                    title="Aggregation range appears to exclude adjacent data row",
                    formula=formula,
                    evidence=[
                        f"{ws.title}!{ref_text} stops short of {ws.title}!{coord} (value {value!r}), "
                        f"which sits {direction} the range, between it and the total."
                    ],
                    suggested_fix=f"Confirm whether {ws.title}!{coord} belongs in the aggregate, then extend the range if appropriate.",
                )
            )
    if min_row == max_row and frow == min_row:
        candidates = []
        if fcol > max_col + 1:
            candidates.append((min_row, max_col + 1, "right of"))
        elif fcol < min_col - 1:
            candidates.append((min_row, min_col - 1, "left of"))
        for row, col, direction in candidates:
            if _is_self(ws, formula_cell, row, col):
                continue
            value = cell_value(ws, row, col)
            if not _is_plain_number(value):
                continue
            if _has_subtotal_word(_column_header(ws, col, row)):
                continue
            coord = f"{get_column_letter(col)}{row}"
            findings.append(
                Finding(
                    rule_id="RANGE_EXCLUSION",
                    severity="Critical",
                    error_confidence="Likely defect",
                    detection_mode="DET",
                    location=loc,
                    title="Aggregation range appears to exclude adjacent data column",
                    formula=formula,
                    evidence=[
                        f"{ws.title}!{ref_text} stops short of {ws.title}!{coord} (value {value!r}), "
                        f"which sits {direction} the range, between it and the total."
                    ],
                    suggested_fix=f"Confirm whether {ws.title}!{coord} belongs in the aggregate, then extend the range if appropriate.",
                )
            )
    return findings


def _subtotal_finding(formula_cell: dict, evidence: str, overlap: bool) -> Finding:
    if overlap:
        return Finding(
            rule_id="RANGE_INCLUDES_SUBTOTAL",
            severity="High",
            error_confidence="Likely defect",
            detection_mode="DET",
            location=formula_cell["location"],
            title="Aggregation range includes a subtotal of its own components",
            formula=formula_cell["formula"],
            evidence=[evidence],
            suggested_fix="Exclude the subtotal from the range, or sum the subtotals instead of the components; as written the components are counted twice.",
        )
    return Finding(
        rule_id="RANGE_INCLUDES_SUBTOTAL",
        severity="Medium",
        error_confidence="Review",
        detection_mode="DET",
        location=formula_cell["location"],
        title="Aggregation range includes a row or column labeled as a total",
        formula=formula_cell["formula"],
        evidence=[evidence],
        suggested_fix="Confirm the labeled total is a genuine component of this aggregate and not a hardcoded subtotal.",
    )


def _detect_subtotal_inclusion(ws, formula_cell: dict, ref_text: str, box: Box, names) -> list[Finding]:
    """Flag ranges that sum a subtotal together with the cells that subtotal already sums.

    A range that includes a row labeled "Subtotal" is only a double count when
    that subtotal's own formula aggregates cells inside the range. "Subtotal +
    Other revenue" is a normal total and stays quiet, and so does a labeled
    row or column of constants: "last week's total" carried forward into
    "total overall" is a design, not a double count.
    """
    min_col, min_row, max_col, max_row = box
    findings: list[Finding] = []

    if min_row < max_row:
        for row in range(min_row, max_row + 1):
            label = _row_label(ws, row, min_col)
            if not _has_subtotal_word(label):
                continue
            cells = [(col, cell_value(ws, row, col)) for col in range(min_col, max_col + 1)]
            formulas = [(col, value) for col, value in cells if is_formula(value)]
            if any(_formula_overlaps(ws.title, value, row, col, box, names) for col, value in formulas):
                findings.append(
                    _subtotal_finding(
                        formula_cell,
                        f"{ws.title}!{ref_text} includes row {row}, labeled {label!r}, whose formula aggregates cells that are also inside the range.",
                        overlap=True,
                    )
                )

    if min_col < max_col:
        for col in range(min_col, max_col + 1):
            header = _column_header(ws, col, min_row)
            if not _has_subtotal_word(header):
                continue
            cells = [(row, cell_value(ws, row, col)) for row in range(min_row, max_row + 1)]
            formulas = [(row, value) for row, value in cells if is_formula(value)]
            letter = get_column_letter(col)
            if any(_formula_overlaps(ws.title, value, row, col, box, names) for row, value in formulas):
                findings.append(
                    _subtotal_finding(
                        formula_cell,
                        f"{ws.title}!{ref_text} includes column {letter}, headed {header!r}, whose formula aggregates cells that are also inside the range.",
                        overlap=True,
                    )
                )
    return findings


# --- hidden structure --------------------------------------------------------


def _hidden_rows(ws) -> set[int]:
    hidden: set[int] = set()
    for key, dim in ws.row_dimensions.items():
        if getattr(dim, "hidden", False):
            try:
                hidden.add(int(key))
            except (TypeError, ValueError):
                continue
    return hidden


def _hidden_cols(ws) -> set[int]:
    hidden: set[int] = set()
    for key, dim in ws.column_dimensions.items():
        if not getattr(dim, "hidden", False):
            continue
        low = getattr(dim, "min", None)
        high = getattr(dim, "max", None)
        if low and high:
            hidden.update(range(int(low), int(high) + 1))
            continue
        try:
            hidden.add(column_index_from_string(str(key)))
        except ValueError:
            continue
    return hidden


def detect_hidden_structure(
    formula_wb,
    formula_cells: list[dict],
    names=None,
    budget=None,
    allowed_sheet_names: set[str] | None = None,
) -> list[Finding]:
    """One finding per hidden row, column, or sheet that feeds visible formulas.

    Formulas that are themselves hidden are not reported as dependents; their
    visible consumers are.
    """
    sheets = {ws.title: ws for ws in formula_wb.worksheets}
    hidden_sheets = {
        title for title, ws in sheets.items() if getattr(ws, "sheet_state", "visible") != "visible"
    }
    rows_cache: dict[str, set[int]] = {}
    cols_cache: dict[str, set[int]] = {}

    def hidden_rows_for(title: str) -> set[int]:
        if title not in rows_cache:
            rows_cache[title] = _hidden_rows(sheets[title])
        return rows_cache[title]

    def hidden_cols_for(title: str) -> set[int]:
        if title not in cols_cache:
            cols_cache[title] = _hidden_cols(sheets[title])
        return cols_cache[title]

    groups: dict[tuple[str, str, int], list[dict]] = defaultdict(list)
    for cell in formula_cells:
        tick(budget)
        if allowed_sheet_names is not None and cell["sheet"] not in allowed_sheet_names:
            continue
        if cell["sheet"] in hidden_sheets or cell["sheet"] not in sheets:
            continue
        if cell["row"] in hidden_rows_for(cell["sheet"]) or cell["col"] in hidden_cols_for(cell["sheet"]):
            continue
        parsed = parse_formula(cell["formula"], names=names, origin=_origin(cell))
        skips_hidden_rows = _ignores_hidden_rows(cell["formula"])
        seen_keys: set[tuple[str, str, int]] = set()
        for ref in parsed.references:
            if ref.positional:
                continue
            target = find_sheet(formula_wb, ref.sheet or cell["sheet"])
            if target is None:
                continue
            if target in hidden_sheets and target != cell["sheet"]:
                key = (target, "sheet", 0)
                if key not in seen_keys:
                    seen_keys.add(key)
                    groups[key].append(cell)
                continue
            raw = raw_boundaries(ref.ref)
            if raw is None:
                continue
            ws = sheets[target]
            min_col, min_row, max_col, max_row = raw
            min_col = min_col or 1
            min_row = min_row or 1
            max_col = max_col if max_col is not None else int(ws.max_column or 1)
            max_row = max_row if max_row is not None else int(ws.max_row or 1)
            if not skips_hidden_rows:
                for hidden_row in hidden_rows_for(target):
                    if min_row <= hidden_row <= max_row:
                        key = (target, "row", hidden_row)
                        if key not in seen_keys:
                            seen_keys.add(key)
                            groups[key].append(cell)
            for hidden_col in hidden_cols_for(target):
                if min_col <= hidden_col <= max_col:
                    key = (target, "col", hidden_col)
                    if key not in seen_keys:
                        seen_keys.add(key)
                        groups[key].append(cell)

    # Hidden rows that feed the same visible formulas are one fact: "rows 46,
    # 48 and 60 of the receipts sheet are hidden and still counted".
    merged: dict[tuple[str, str, tuple[str, ...]], list[int]] = defaultdict(list)
    dependents_of: dict[tuple[str, str, tuple[str, ...]], list[dict]] = {}
    for key in sorted(groups):
        dependents = sorted(groups[key], key=lambda c: (c["sheet"], c["row"], c["col"]))
        target, kind, number = key
        merged_key = (target, kind, tuple(c["location"] for c in dependents))
        merged[merged_key].append(number)
        dependents_of[merged_key] = dependents

    findings: list[Finding] = []
    for merged_key in sorted(merged):
        target, kind, _ = merged_key
        numbers = merged[merged_key]
        dependents = dependents_of[merged_key]
        first = dependents[0]
        if kind == "row":
            what = f"Hidden row {numbers[0]} on {target}" if len(numbers) == 1 else f"Hidden rows {_number_list(numbers)} on {target}"
        elif kind == "col":
            letters = [get_column_letter(number) for number in numbers]
            what = f"Hidden column {letters[0]} on {target}" if len(letters) == 1 else f"Hidden columns {_number_list(letters)} on {target}"
        else:
            what = f"Hidden sheet {target!r}"
        verb = "feeds" if len(numbers) == 1 else "feed"
        shown = ", ".join(c["location"] for c in dependents[:MAX_DEPENDENTS_SHOWN])
        if len(dependents) > MAX_DEPENDENTS_SHOWN:
            shown += ", ..."
        findings.append(
            Finding(
                rule_id="HIDDEN_STRUCTURE_IN_TOTAL",
                severity="Medium",
                error_confidence="Review",
                detection_mode="DET",
                location=first["location"],
                title="Formula inputs include hidden structure",
                formula=first["formula"],
                evidence=[f"{what} {verb} {len(dependents)} visible formula(s): {shown}."],
                suggested_fix="Confirm hidden inputs are intentional and disclosed in visible workbook documentation.",
            )
        )
    return findings


def _number_list(items: list) -> str:
    shown = ", ".join(str(item) for item in items[:MAX_DEPENDENTS_SHOWN])
    if len(items) > MAX_DEPENDENTS_SHOWN:
        shown += f" and {len(items) - MAX_DEPENDENTS_SHOWN} more"
    return shown


def _ignores_hidden_rows(formula: str) -> bool:
    """AGGREGATE with option 1, 3, 5 or 7 and SUBTOTAL 101-111 skip hidden rows by definition."""
    toks = tokens(formula)
    if not toks:
        return False
    for idx, (value, ttype, subtype) in enumerate(toks):
        if ttype != Token.FUNC or subtype != Token.OPEN:
            continue
        name = value[:-1].upper()
        for prefix in FUNCTION_PREFIXES:
            if name.startswith(prefix):
                name = name[len(prefix) :]
                break
        if name not in ("AGGREGATE", "SUBTOTAL"):
            continue
        numbers = _leading_numbers(toks, idx + 1)
        if name == "SUBTOTAL" and numbers and numbers[0] >= 101:
            return True
        if name == "AGGREGATE" and len(numbers) >= 2 and numbers[1] in (1, 3, 5, 7):
            return True
    return False


def _leading_numbers(toks: tuple, start: int) -> list[int]:
    """The run of plain integer arguments that opens a call: ``AGGREGATE(9,3,`` -> [9, 3]."""
    numbers: list[int] = []
    expect_number = True
    for value, ttype, subtype in toks[start:]:
        if ttype == Token.WSPACE:
            continue
        if expect_number:
            if ttype == Token.OPERAND and subtype == Token.NUMBER:
                try:
                    numbers.append(int(float(value)))
                except ValueError:
                    break
                expect_number = False
                continue
            break
        if ttype == Token.SEP and subtype == Token.ARG:
            expect_number = True
            continue
        break
    return numbers


# --- literals and fragile functions ------------------------------------------


def detect_literal_constants(formula_cells: list[dict], budget=None) -> list[Finding]:
    """One finding per formula pattern that embeds an assumption-like literal.

    Literals in structural argument slots (a VLOOKUP column index, a MID
    length, a date part, a rounding digit) are not reported; see
    ``LITERAL_SLOT_FUNCS`` in the parser.
    """
    hits: list[dict] = []
    literals_of: dict[int, list[str]] = {}
    for cell in formula_cells:
        tick(budget)
        literals = parse_formula(cell["formula"]).numeric_literals
        if literals:
            hits.append(cell)
            literals_of[id(cell)] = literals
    findings: list[Finding] = []
    for members in group_by_pattern(hits):
        lead = members[0]
        evidence = [f"Non-trivial numeric literal(s) found: {', '.join(literals_of[id(lead)])}."]
        note = pattern_note(members)
        if note:
            evidence.append(note)
        findings.append(
            Finding(
                rule_id="LITERAL_CONSTANT",
                severity="Medium",
                error_confidence="Review",
                detection_mode="DET",
                location=lead["location"],
                title="Formula contains embedded numeric literal",
                formula=lead["formula"],
                evidence=evidence,
                suggested_fix="Move the assumption to a labeled input cell and reference it from the formula.",
            )
        )
    return findings


def detect_fragile_functions(formula_cells: list[dict], names=None, budget=None) -> list[Finding]:
    """Volatile functions and whole-column references, once per formula pattern.

    RAND, RANDBETWEEN, OFFSET and INDIRECT make results irreproducible or
    fragile and are reported per pattern. TODAY and NOW are the normal way to
    age a date; they get one Info note per sheet. A whole-column reference is
    only reported where it is evaluated as an array (a comparison, arithmetic,
    IF, SUMPRODUCT, an array formula): COUNTIF or VLOOKUP over ``A:A`` bound
    themselves to the used range.
    """
    volatile = {"OFFSET", "INDIRECT", "RAND", "RANDBETWEEN"}
    clock = {"TODAY", "NOW"}
    volatile_cells: list[dict] = []
    found_of: dict[int, set[str]] = {}
    clock_cells: dict[str, list[dict]] = defaultdict(list)
    whole_cells: list[dict] = []
    for cell in formula_cells:
        tick(budget)
        parsed = parse_formula(cell["formula"], names=names, origin=_origin(cell))
        found = parsed.functions.intersection(volatile)
        if found:
            volatile_cells.append(cell)
            found_of[id(cell)] = found
        elif parsed.functions.intersection(clock):
            clock_cells[cell["sheet"]].append(cell)
        unbounded = [ref for ref in parsed.references if not ref.bounded and not ref.positional]
        if unbounded and (cell.get("array") or any(ref.array_context for ref in unbounded)):
            whole_cells.append(cell)

    findings: list[Finding] = []
    for members in group_by_pattern(volatile_cells):
        lead = members[0]
        evidence = [f"Function(s) found: {', '.join(sorted(found_of[id(lead)]))}."]
        note = pattern_note(members)
        if note:
            evidence.append(note)
        findings.append(
            Finding(
                rule_id="VOLATILE_FUNCTION",
                severity="Medium",
                error_confidence="Review",
                detection_mode="DET",
                location=lead["location"],
                title="Formula uses volatile or fragile function",
                formula=lead["formula"],
                evidence=evidence,
                suggested_fix="Confirm the volatility is intentional; prefer stable bounded references when possible.",
            )
        )
    for sheet in sorted(clock_cells):
        cells = clock_cells[sheet]
        shown = ", ".join(c["location"] for c in cells[:MAX_DEPENDENTS_SHOWN])
        if len(cells) > MAX_DEPENDENTS_SHOWN:
            shown += ", ..."
        findings.append(
            Finding(
                rule_id="VOLATILE_FUNCTION",
                severity="Low",
                error_confidence="Info",
                detection_mode="DET",
                location=cells[0]["location"],
                title="Formulas depend on the current date or time",
                formula=cells[0]["formula"],
                evidence=[f"{len(cells)} formula(s) on {sheet} use TODAY or NOW, so their results change with the clock: {shown}."],
                suggested_fix="Nothing to fix if ageing against today is the intent; freeze the date in an input cell for a reproducible snapshot.",
            )
        )
    for members in group_by_pattern(whole_cells):
        lead = members[0]
        evidence = [
            "A whole-column reference is evaluated as an array here, so every recalculation scans the full column height."
        ]
        note = pattern_note(members)
        if note:
            evidence.append(note)
        findings.append(
            Finding(
                rule_id="WHOLE_COLUMN_REFERENCE",
                severity="Medium",
                error_confidence="Review",
                detection_mode="DET",
                location=lead["location"],
                title="Formula evaluates an entire column as an array",
                formula=lead["formula"],
                evidence=evidence,
                suggested_fix="Use a bounded range sized to the actual table.",
            )
        )
    return findings


__all__ = [
    "AGG_FUNCS",
    "SUBTOTAL_WORDS",
    "EXCEL_MAX_COL",
    "EXCEL_MAX_ROW",
    "aggregate_ranges",
    "is_total_of_segment",
    "detect_range_length_mismatch",
    "detect_range_issues",
    "detect_hidden_structure",
    "detect_literal_constants",
    "detect_fragile_functions",
]
