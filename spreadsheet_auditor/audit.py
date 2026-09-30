from __future__ import annotations

import argparse
import contextlib
import csv
import importlib.resources
import io
import json
import platform
import sys
import tempfile
import time
import traceback
from collections import Counter, defaultdict
from pathlib import Path

from . import __version__
from .budget import AuditTimeout, Budget
from .config_loader import ConfigError, allowed_sheets, apply_check_settings, load_config, sheet_is_allowed
from .finding import Finding, assign_ids, sort_findings
from .locations import location_matches
from .names import NameTable
from .preflight import PreflightError, preflight
from .recalc import recalc_if_available, soffice_path
from .report import (
    build_payload,
    render_html,
    render_markdown,
    write_html,
    write_json,
    write_markdown,
)

AUDIT_VERSION = __version__
FAIL_ORDER = {"Critical": 0, "High": 1, "Medium": 2, "Low": 3, "None": 99}


REQUIRED_PACKAGES = ["openpyxl"]
OPTIONAL_PACKAGES = ["defusedxml", "yaml"]


def _probe_package(package: str) -> dict:
    try:
        module = __import__(package)
        return {
            "name": package,
            "status": "installed",
            "version": getattr(module, "__version__", "unknown"),
        }
    except Exception:
        return {"name": package, "status": "missing", "version": None}


def collect_healthcheck() -> dict:
    required = [_probe_package(pkg) for pkg in REQUIRED_PACKAGES]
    optional = [_probe_package(pkg) for pkg in OPTIONAL_PACKAGES]
    libreoffice_path = soffice_path()
    libreoffice = {
        "path": libreoffice_path,
        "status": "found" if libreoffice_path else "missing",
    }
    recalculation_mode = "available" if libreoffice_path else "cached_values_only"
    missing_required = [pkg["name"] for pkg in required if pkg["status"] == "missing"]
    return {
        "tool_version": AUDIT_VERSION,
        "python": sys.version.split()[0],
        "platform": platform.platform(),
        "required_packages": required,
        "optional_packages": optional,
        "libreoffice": libreoffice,
        "recalculation_mode": recalculation_mode,
        "macro_execution": "disabled",
        "external_links": "inventoried_only",
        "supported_file_types": [".xlsx", ".xlsm", ".csv"],
        "status": "ready" if not missing_required else "missing_required_dependencies",
        "missing_required": missing_required,
    }


def _render_healthcheck(info: dict) -> str:
    lines = [
        f"Spreadsheet Auditor Healthcheck (v{info['tool_version']})",
        "",
        f"Python: {info['python']} OK",
    ]
    for pkg in info["required_packages"]:
        status = "OK" if pkg["status"] == "installed" else "MISSING"
        version = pkg["version"] or "n/a"
        lines.append(f"{pkg['name']}: {pkg['status']} ({version}) {status}")
    for pkg in info["optional_packages"]:
        if pkg["status"] == "installed":
            lines.append(f"{pkg['name']}: installed ({pkg['version']}) OK")
        else:
            lines.append(f"{pkg['name']}: not installed (optional)")
    if info["libreoffice"]["status"] == "found":
        lines.append(f"LibreOffice: found at {info['libreoffice']['path']} OK")
    else:
        lines.append("LibreOffice: not installed (optional; value-dependent checks limited)")
    lines.append(f"Recalculation mode: {info['recalculation_mode']}")
    lines.append(f"Macro execution: {info['macro_execution']}")
    lines.append(f"External links: {info['external_links']}")
    lines.append("")
    if info["missing_required"]:
        lines.append(f"Status: missing required dependencies: {', '.join(info['missing_required'])}")
    else:
        lines.append("Status: ready")
    return "\n".join(lines)


def healthcheck(as_json: bool = False) -> int:
    info = collect_healthcheck()
    if as_json:
        print(json.dumps(info, indent=2))
    else:
        print(_render_healthcheck(info))
    return 0 if not info["missing_required"] else 3


def _package_available(package: str) -> bool:
    try:
        __import__(package)
        return True
    except Exception:
        return False


def _suppression_notes(
    suppressions: list[dict],
    config: dict,
    incomplete: list[dict],
    findings: list[Finding],
    workbook_sheets: set[str] | None = None,
    allowed_sheets: set[str] | None = None,
) -> list[str]:
    """Notes on the suppressions that need the user's attention.

    * A pinned line (a location plus the fingerprint it accepted) whose
      finding has moved away from that location: still suppressed, but the
      address in the file is out of date.
    * An unpinned one-cell line that matched: it follows the address, so a
      different finding that later lands there would be hidden too; the note
      gives the pinned line to replace it with.
    * A location line that covers some but not all cells of a finding that
      stands for several: the finding stays reported, and the note says how
      to cover it whole.
    * A suppression that matched nothing although it could have: its finding
      was fixed, or its cells moved. Suppressions that could not have matched
      this run stay quiet: their rule is turned off, their rule was not
      checked because the audit is incomplete, or their sheet is excluded by
      scope.
    """
    from .config_loader import rule_enabled
    from .locations import location_matches, target_sheet
    from .suppressions import describe, is_one_cell, pinned_line

    unchecked = {rule for entry in incomplete for rule in entry["rules"]}
    everything_unchecked = any(not entry["rules"] for entry in incomplete)
    existing = {name.casefold() for name in workbook_sheets or ()}
    allowed = {name.casefold() for name in allowed_sheets or ()}
    notes = []
    for suppression in suppressions:
        rule = str(suppression["rule_id"]).upper() if suppression.get("rule_id") else None
        target = str(suppression.get("range") or "")
        pin = suppression.get("fingerprint")
        head = f"Suppression ({suppression.get('source', 'suppression')}) {describe(suppression)}"
        matched = suppression.get("matched_findings") or []
        partial = [entry for entry in suppression.get("partial_findings") or [] if not entry[0].suppressed]
        for finding, inside, total in partial:
            span = _span(finding.covered())
            cover = f"cover all of them ({span})" if span else "cover all of them"
            notes.append(
                f"{head} covers {inside} of the {total} cells of the {finding.rule_id} finding at "
                f"{finding.location}, so the finding stays reported. To accept it, {cover} or pin the line "
                f"to it: {pinned_line(finding.rule_id, finding.location, finding.fingerprint, suppression.get('reason') or '<reason>')}"
            )
        if partial and not matched:
            continue
        if matched:
            if suppression.get("rewritten"):
                continue  # --pin-suppressions just brought this line up to date
            if pin and target and not any(location_matches(f.location, target) for f in matched):
                where = ", ".join(f.location for f in matched)
                notes.append(f"{head} follows its finding, which is now at {where}; update the address when convenient.")
            elif not pin and is_one_cell(target):
                lines = "; ".join(
                    pinned_line(f.rule_id, f.location, f.fingerprint, suppression.get("reason") or "<reason>")
                    for f in matched
                )
                notes.append(
                    f"{head} follows the address, so a different finding that later lands on it would be "
                    f"hidden too. Pin it to the finding it accepts (--pin-suppressions rewrites it): {lines}"
                )
            continue
        if rule is None:  # a bare fingerprint: its rule is unknown until it matches
            if incomplete:
                continue
            why = (
                "The finding was fixed, or the flagged cell's formula, value or row label changed; "
                "copy its new fingerprint from the report or remove the suppression."
            )
        else:
            if not rule_enabled(config, rule):
                continue
            if everything_unchecked or rule in unchecked:
                continue
            sheet = (target_sheet(target) or "").casefold()
            if sheet in existing and sheet not in allowed:
                continue
            if pin:
                now = [f for f in findings if f.rule_id == rule and location_matches(f.location, target)]
                why = "The accepted finding was fixed, or its cell's formula, value or row label changed."
                if now:
                    why += (
                        f" {target} now holds a different {rule} finding, which is reported "
                        f"(fingerprint {now[0].fingerprint})."
                    )
                why += " Pin the line to the current finding or remove it."
            else:
                why = (
                    "The finding was fixed, or its cells moved or the target is misspelled; update or remove the "
                    "suppression, since a stale target can later hide a different finding on that address."
                )
        notes.append(f"{head} matched no finding. {why}")
    return notes


def _span(cells: list[str]) -> str | None:
    """The smallest target covering ``cells`` when they share a sheet, quoted as the file needs."""
    from openpyxl.utils.cell import get_column_letter

    from .locations import anchor, format_target

    spots = [anchor(cell) for cell in cells]
    if not spots or any(spot is None for spot in spots) or len({spot[0] for spot in spots}) != 1:
        return None
    rows = [spot[1] for spot in spots]
    cols = [spot[2] for spot in spots]
    top_left = f"{get_column_letter(min(cols))}{min(rows)}"
    bottom_right = f"{get_column_letter(max(cols))}{max(rows)}"
    return format_target(f"{spots[0][0]}!{top_left}:{bottom_right}")


def _finding_counts(findings: list[Finding]) -> dict:
    """Counts of the active findings by severity and by rule, for summaries that outlive the cap."""
    active = [finding for finding in findings if not finding.suppressed]
    by_rule = Counter((finding.severity, finding.error_confidence, finding.rule_id) for finding in active)
    return {
        "total": len(active),
        "shown": len(active),
        "suppressed": len(findings) - len(active),
        "by_severity": dict(Counter(finding.severity for finding in active)),
        "by_rule": [
            {"rule_id": rule, "severity": severity, "error_confidence": confidence, "count": count}
            for (severity, confidence, rule), count in sorted(
                by_rule.items(), key=lambda item: (FAIL_ORDER.get(item[0][0], 4), -item[1], item[0][2])
            )
        ],
    }


def _cap_findings(
    findings: list[Finding], max_reported: int, truncated: dict, limitations: list[str]
) -> tuple[list[Finding], dict]:
    """Keep at most ``max_reported`` findings without letting one rule crowd out the rest.

    Active findings go before suppressed ones. Every rule first keeps its most
    severe active finding, so a workbook with hundreds of one error still shows
    the other rules' findings; the remaining places go in severity order. The
    counts are taken before the cap, and the limitation note says what each
    rule lost.
    """
    counts = _finding_counts(findings)
    if max_reported <= 0 or len(findings) <= max_reported:
        return findings, counts
    ordered = sorted(findings, key=lambda finding: finding.suppressed)  # stable: severity order within
    kept: list[Finding] = []
    kept_ids: set[int] = set()
    rules_seen: set[str] = set()
    for finding in ordered:
        if not finding.suppressed and finding.rule_id not in rules_seen:
            rules_seen.add(finding.rule_id)
            kept.append(finding)
            kept_ids.add(id(finding))
    for finding in ordered:
        if len(kept) >= max_reported:
            break
        if id(finding) not in kept_ids:
            kept.append(finding)
            kept_ids.add(id(finding))
    kept = sort_findings(kept[:max_reported])
    kept_ids = {id(finding) for finding in kept}
    left_out = Counter(
        finding.rule_id for finding in findings if id(finding) not in kept_ids and not finding.suppressed
    )
    truncated["findings"] = True
    counts["shown"] = sum(1 for finding in kept if not finding.suppressed)
    note = f"Findings output capped at {max_reported} of {len(findings)} findings by config"
    if left_out:
        note += "; not shown: " + ", ".join(f"{count} {rule}" for rule, count in left_out.most_common())
    limitations.append(note + ". Raise limits.max_reported_findings to see them all.")
    return kept, counts


def _pin_suppressions(ignore_path: str, suppressions: list[dict]) -> None:
    """Run --pin-suppressions and say on stderr what changed, so it can be reviewed."""
    from .suppressions import pin_suppression_file

    if not Path(ignore_path).exists():
        print(f"No suppression file at {ignore_path}; nothing to pin.", file=sys.stderr)
        return
    changes = pin_suppression_file(ignore_path, suppressions)
    if not changes:
        print(f"Every one-cell line in {ignore_path} is already pinned and current.", file=sys.stderr)
        return
    print(f"Pinned {len(changes)} line(s) in {ignore_path}; review them before committing:", file=sys.stderr)
    for change in changes:
        print(f"  line {change['line']}: {change['old']}", file=sys.stderr)
        for line, finding in zip(change["new"], change["findings"]):
            print(f"    -> {line}", file=sys.stderr)
            print(f"       ({finding.location}: {finding.formula or finding.title})", file=sys.stderr)


def _known_rules() -> set[str]:
    """Rule IDs the registered checks report, which a config's checks section may name."""
    from .checks import checks as registered_checks

    return {rule for check_cls in registered_checks() for rule in check_cls.rule_ids}


def audit_workbook(args: argparse.Namespace) -> tuple[dict, int]:
    from .identity import assign_identities
    from .suppressions import apply_suppressions, load_suppressions
    from .workbook_inventory import (
        formula_cells,
        inventory,
        load_workbook_formulas,
        load_workbook_values,
    )

    input_path = Path(args.workbook)
    preflight_info = preflight(input_path)
    config = load_config(args.config, known_rules=_known_rules())
    limitations: list[str] = []
    # Why part of the workbook went unchecked (a check failed or ran out of
    # time, or a size cap skipped formulas or cells). Any entry makes the audit
    # incomplete, which exits 6 whatever the findings: a partial audit must
    # never pass a CI gate as if it were whole. Each message is also a
    # limitation, so the reports that list limitations need no change.
    incomplete: list[dict] = []
    unsupported_features: set[str] = set()
    findings: list[Finding] = []
    truncated = {"formulas": False, "findings": False, "cells": False, "timeout": False}
    start_time = time.monotonic()
    limits_config = config.get("limits") or {}
    timeout_seconds = int(limits_config.get("timeout_seconds", 0) or 0)
    budget = Budget(timeout_seconds, start=start_time)

    def _note_incomplete(reason: str, message: str, rules: list[str] | None = None) -> None:
        limitations.append(message)
        incomplete.append({"reason": reason, "message": message, "rules": sorted(set(rules or []))})

    def _note_timeout(detail: str, skipped: list) -> None:
        if not truncated["timeout"]:
            truncated["timeout"] = True
            rules = [rule for check_cls in skipped for rule in check_cls.rule_ids]
            _note_incomplete("timeout", f"Audit timeout of {timeout_seconds}s exceeded; {detail}", rules)

    if preflight_info["extension"] == ".csv":
        payload = audit_csv(input_path, preflight_info, config, args.ignore)
        return payload, exit_code(
            payload["findings"], args.fail_on, payload["coverage"]["limitations"], strict=args.strict
        )

    with tempfile.TemporaryDirectory(prefix="spreadsheet-auditor-audit-") as temp_dir:
        recalc_config = config.get("recalc") or {}
        recalc_enabled = recalc_config.get("enabled", True)
        recalc_timeout = args.recalc_timeout if args.recalc_timeout is not None else int(recalc_config.get("timeout_seconds", 60))
        if recalc_enabled:
            recalc = recalc_if_available(input_path, timeout_seconds=recalc_timeout, work_dir=temp_dir)
        else:
            recalc = {
                "status": "disabled",
                "path": str(input_path),
                "limitations": ["Recalculation disabled by config; using static analysis and cached values only."],
            }
        limitations.extend(recalc.get("limitations", []))
        if not _package_available("defusedxml"):
            limitations.append("defusedxml is not available; XML parsing relies on workbook library defaults in this runtime.")
        analysis_path = Path(recalc.get("path", input_path))
        recalc_status = recalc.get("status", "unknown")

        # Static checks always read the original workbook so formula text is
        # exactly what the author wrote; LibreOffice re-serializes formulas
        # (for example ``=SUM(#ref!)``) in its converted copy. That copy, when
        # recalculation ran, only supplies cached values.
        try:
            formula_wb = load_workbook_formulas(input_path)
            value_wb = load_workbook_values(analysis_path)
        except Exception as exc:
            # The file passed preflight as a zip but its contents do not parse.
            raise PreflightError(
                f"The workbook could not be read ({type(exc).__name__}: {exc}); open it in Excel, "
                "save a copy, and audit the copy."
            ) from exc
        inv = inventory(input_path, formula_wb, value_wb, preflight_info)
        include, exclude = allowed_sheets(config)
        allowed_sheet_names = {
            ws.title for ws in formula_wb.worksheets if sheet_is_allowed(ws.title, include, exclude)
        }
        if "max_range_expansion_cells" in limits_config:
            limitations.append(
                "limits.max_range_expansion_cells is deprecated and ignored; ranges are resolved "
                "exactly against the formula index, so no reference is dropped for being large."
            )
        if getattr(args, "annotated", None) and preflight_info.get("drawing_parts"):
            limitations.append(
                f"The annotated copy is written by openpyxl, which does not preserve drawings, charts, "
                f"images, or form controls; this workbook contains {preflight_info['drawing_parts']} "
                "such part(s) that the copy will drop. Keep the original workbook as the master."
            )
        # Cell-count guardrail. The grid checks walk only the cells a sheet holds,
        # so count those, not the used-range rectangle: one stray cell at row
        # 48,000 must not switch the grid checks off.
        max_cells = int(limits_config.get("max_cells", 1_000_000) or 0)
        total_cells = 0
        for ws in formula_wb.worksheets:
            if ws.title not in allowed_sheet_names:
                continue
            held = getattr(ws, "_cells", None)
            total_cells += len(held) if held is not None else int(ws.max_row or 0) * int(ws.max_column or 0)
        grid_scan_allowed = True
        if max_cells > 0 and total_cells > max_cells:
            truncated["cells"] = True
            grid_scan_allowed = False
            _note_incomplete(
                "cell_cap",
                f"Cell scan capped: workbook holds {total_cells} cells, exceeding the configured "
                f"max_cells={max_cells}; cell-grid checks (HARDCODE_IN_FORMULA_BLOCK, CROSS_FOOT_FAILURE, "
                "data hygiene) were skipped. Formula-based checks still ran. Raise limits.max_cells to check them.",
                ["HARDCODE_IN_FORMULA_BLOCK", "CROSS_FOOT_FAILURE", "NUMBERS_STORED_AS_TEXT",
                 "WHITESPACE_KEY", "DUPLICATE_KEY", "MERGED_CELL_IN_DATA_RANGE"],
            )

        all_formulas = formula_cells(formula_wb)
        formulas = [cell for cell in all_formulas if cell["sheet"] in allowed_sheet_names]
        max_formulas = int(limits_config.get("max_formulas", 50000))
        if max_formulas > 0 and len(formulas) > max_formulas:
            truncated["formulas"] = True
            _note_incomplete(
                "formula_cap",
                f"Formula scan capped at {max_formulas} formulas by config; {len(formulas) - max_formulas} of "
                f"{len(formulas)} formulas, from {formulas[max_formulas]['location']} on, were not checked. "
                "Raise limits.max_formulas to check them.",
            )
            formulas = formulas[:max_formulas]

        from .checks import CheckContext, checks as registered_checks

        names = NameTable.from_workbook(formula_wb)
        array_formulas = int(inv.get("array_formulas", 0) or 0)
        if array_formulas:
            unsupported_features.add("array_formulas")
            limitations.append(
                f"{array_formulas} array formula(s) were parsed for references only; "
                "array and spill semantics are not evaluated."
            )
        ctx = CheckContext(
            workbook_path=input_path,
            formula_wb=formula_wb,
            value_wb=value_wb,
            allowed_sheet_names=allowed_sheet_names,
            formulas=formulas,
            config=config,
            inventory=inv,
            unsupported_features=unsupported_features,
            names=names,
            budget=budget,
            grid_scan_allowed=grid_scan_allowed,
        )
        ordered = registered_checks()
        for index, check_cls in enumerate(ordered):
            check_name = getattr(check_cls, "name", "") or check_cls.__name__
            if budget.expired():
                _note_timeout(f"'{check_name}' and remaining checks were skipped.", ordered[index:])
                break
            check = check_cls()
            try:
                findings.extend(check.run(ctx))
            except AuditTimeout:
                _note_timeout(
                    f"'{check_name}' was interrupted and remaining checks were skipped.", ordered[index:]
                )
                break
            except Exception as exc:
                _note_incomplete(
                    "check_failed",
                    f"Check '{check_name}' raised an exception and was skipped: {exc!r}",
                    list(check_cls.rule_ids),
                )

        findings = _dedupe_range_length_with_drift(findings)
        if "google_sheets_placeholders" in unsupported_features:
            limitations.append(
                "Some formulas are Google Sheets placeholders (IFERROR(__xludf.DUMMYFUNCTION(...), cached value)); "
                "those cells are frozen values, not live calculations."
            )
        if recalc_status != "completed":
            limitations.append(
                "Recalculation did not run; value-dependent checks (TOTAL_MISMATCH, CROSS_FOOT_FAILURE) "
                "rely on cached values and may be incomplete."
            )

        findings = apply_impact_escalation(findings, config)
        findings = apply_check_settings(findings, config)
        assign_identities(findings, formula_wb)
        suppressions = load_suppressions(config, args.ignore, warnings=limitations)
        findings = apply_suppressions(findings, suppressions)
        if getattr(args, "pin_suppressions", False) and args.ignore:
            _pin_suppressions(args.ignore, suppressions)
        limitations.extend(
            _suppression_notes(
                suppressions,
                config,
                incomplete,
                findings,
                workbook_sheets={ws.title for ws in formula_wb.worksheets},
                allowed_sheets=allowed_sheet_names,
            )
        )
        findings = sort_findings(findings)
        all_findings = findings
        findings, finding_counts = _cap_findings(
            findings, int(limits_config.get("max_reported_findings", 200)), truncated, limitations
        )
        assign_ids(findings)

        coverage = {
            "macros_present": inv["macros_present"],
            "macros_executed": False,
            "external_links_present": bool(inv["external_links"]) or "external_workbook_links" in unsupported_features,
            "unsupported_features": sorted(unsupported_features),
            "limitations": limitations,
            "complete": not incomplete,
            "incomplete": incomplete,
            "finding_counts": finding_counts,
            "truncated": truncated,
            "elapsed_seconds": round(time.monotonic() - start_time, 3),
        }
        workbook_meta = {
            "path": input_path.as_posix(),
            "sha256": inv["sha256"],
            "sheets_analyzed": len(allowed_sheet_names),
            "formulas_scanned": len(formulas),
            "recalc_status": recalc_status,
        }
        payload = build_payload(AUDIT_VERSION, workbook_meta, coverage, findings)
        return payload, exit_code(
            [finding.to_dict() for finding in all_findings],
            args.fail_on,
            limitations,
            strict=args.strict,
            incomplete=bool(incomplete),
        )


def audit_csv(path: Path, preflight_info: dict, config: dict | None = None, ignore_path: str | None = None) -> dict:
    from .suppressions import apply_suppressions, load_suppressions

    # One finding per column: a padded notes column is one thing to trim, and
    # a column padded throughout came out of a fixed-width export (as in the
    # workbook rule, it gets an Info note). Whitespace-only cells are spacers.
    padded: dict[int, list[tuple[int, str]]] = defaultdict(list)
    filled: Counter = Counter()
    # Excel's "CSV (Comma delimited)" is Windows-1252 on Western Windows
    # systems; only "CSV UTF-8" is UTF-8.
    raw = path.read_bytes()
    encoding_notes: list[str] = []
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        try:
            text = raw.decode("cp1252")
        except UnicodeDecodeError as exc:
            raise PreflightError(
                "The CSV is neither UTF-8 nor Windows-1252 text; save it as 'CSV UTF-8' and audit that."
            ) from exc
        encoding_notes.append(
            "The CSV is not UTF-8, so it was read as Windows-1252, the encoding of Excel's 'CSV (Comma delimited)'."
        )
    for row_idx, row in enumerate(csv.reader(io.StringIO(text, newline="")), start=1):
        for col_idx, value in enumerate(row, start=1):
            if not value.strip():
                continue
            filled[col_idx] += 1
            if value != value.strip():
                padded[col_idx].append((row_idx, value))
    findings: list[Finding] = []
    for col_idx, cells in sorted(padded.items()):
        row_idx, value = cells[0]
        export = len(cells) >= 3 and len(cells) * 5 >= filled[col_idx] * 4
        if export:
            evidence = [
                f"{len(cells)} of {filled[col_idx]} values in column {col_idx} carry leading or trailing whitespace, "
                f"for example {value!r}; this looks like fixed-width padding from an export."
            ]
        else:
            evidence = [f"Raw value is {value!r}."]
            if len(cells) > 1:
                shown = ", ".join(f"R{row}C{col_idx} {text!r}" for row, text in cells[1:9])
                if len(cells) > 9:
                    shown += f", and {len(cells) - 9} more"
                evidence.append(
                    f"{len(cells)} of {filled[col_idx]} values in column {col_idx} carry leading or trailing "
                    f"whitespace; the others are {shown}."
                )
        findings.append(
            Finding(
                rule_id="WHITESPACE_KEY",
                severity="Low" if export else "Medium",
                error_confidence="Info" if export else "Review",
                detection_mode="DET",
                location=f"CSV!R{row_idx}C{col_idx}",
                title="Column of padded CSV values" if export else "CSV value has leading or trailing whitespace",
                evidence=evidence,
                suggested_fix="Trim if this field is used as a key, label, or numeric input.",
            )
        )
    csv_limitations = ["CSV audit is limited to data-hygiene checks.", *encoding_notes]
    if config:
        findings = apply_check_settings(findings, config)
    suppressions = load_suppressions(config, ignore_path, warnings=csv_limitations)
    findings = apply_suppressions(findings, suppressions)
    csv_limitations.extend(_suppression_notes(suppressions, config or {}, [], findings))
    findings = sort_findings(findings)
    truncated = {"formulas": False, "findings": False, "cells": False, "timeout": False}
    findings, finding_counts = _cap_findings(
        findings,
        int(((config or {}).get("limits") or {}).get("max_reported_findings", 200)),
        truncated,
        csv_limitations,
    )
    assign_ids(findings)
    return build_payload(
        AUDIT_VERSION,
        {
            "path": Path(path).as_posix(),
            "sha256": preflight_info["sha256"],
            "sheets_analyzed": 1,
            "formulas_scanned": 0,
            "recalc_status": "not_applicable",
        },
        {
            "macros_present": False,
            "macros_executed": False,
            "external_links_present": False,
            "unsupported_features": [],
            "limitations": csv_limitations,
            "complete": True,
            "incomplete": [],
            "finding_counts": finding_counts,
            "truncated": truncated,
        },
        findings,
    )


def detect_reference_issues(
    formula_wb,
    value_wb,
    formulas: list[dict],
    unsupported_features: set[str],
    names=None,
    budget=None,
    blank_precedents: bool = True,
) -> list[Finding]:
    """Broken references, blank precedents, external links, and error masks.

    ``blank_precedents`` is False when BLANK_PRECEDENT is turned off (its
    default), which skips the blank-cell analysis altogether.
    """
    from bisect import bisect_left

    from .error_masks import masks as error_masks
    from .formula_parser import external_source, is_formula, parse_formula
    from .grouping import group_by_pattern, locations, pattern_note
    from .reference_resolver import boundaries, existing_cell, find_sheet, reference_in_bounds
    from .workbook_inventory import iter_existing_cells

    # Per sheet: rows holding at least one constant, and per column the sorted
    # rows holding constants. Formulas are not constants: a template row whose
    # only content is formulas is unused space.
    layout: dict[str, tuple[set[int], dict[int, list[int]]]] = {}

    def constants_layout(title: str) -> tuple[set[int], dict[int, list[int]]]:
        if title not in layout:
            rows: set[int] = set()
            cols: dict[int, list[int]] = {}
            for existing in iter_existing_cells(formula_wb[title]):
                value = existing.value
                if value is None or is_formula(value):
                    continue
                rows.add(existing.row)
                cols.setdefault(existing.column, []).append(existing.row)
            layout[title] = (rows, {col: sorted(found) for col, found in cols.items()})
        return layout[title]

    def anomalous_blank(title: str, row: int, col: int, window: int = 10) -> bool:
        """A blank is anomalous when it is a gap in a filled column of a filled row.

        Its row must hold a constant somewhere, its column must hold constants
        both above and below it, and the column must be mostly filled around
        it; a sparse ledger column or the tail of a template is not a gap.
        """
        rows, cols = constants_layout(title)
        if row not in rows:
            return False
        column_rows = cols.get(col)
        if not column_rows:
            return False
        idx = bisect_left(column_rows, row)
        if idx == 0 or idx >= len(column_rows):
            return False
        low = max(column_rows[0], row - window)
        high = min(column_rows[-1], row + window)
        filled = bisect_left(column_rows, high + 1) - bisect_left(column_rows, low)
        # Three in four cells around the blank must be filled: a ledger with a
        # debit and a credit column is half blank by design.
        return filled * 4 >= (high - low + 1) * 3

    findings: list[Finding] = []
    masked: list[tuple[dict, set[str]]] = []
    # Broken formulas are reported once per relative pattern: a #REF! filled
    # down a column is one repair, not one finding per row.
    deleted: list[dict] = []
    unresolved: dict[str, tuple[dict, list[str]]] = {}
    external: dict[tuple[str, str], list[dict]] = {}
    for cell in formulas:
        if budget is not None:
            budget.tick()
        formula = cell["formula"]
        parsed = parse_formula(formula, names=names, origin=(cell["sheet"], cell["row"], cell["col"]))
        if parsed.parse_error:
            unsupported_features.add("unparseable_formulas")
            continue
        if parsed.deleted_reference:
            deleted.append(cell)
        if parsed.external_references:
            unsupported_features.add("external_workbook_links")
            for source in sorted({external_source(raw) for raw in parsed.external_references}):
                external.setdefault((cell["sheet"], source), []).append(cell)
        if parsed.unresolved_structured:
            unsupported_features.add("structured_references")
        if parsed.three_d_references:
            unsupported_features.add("3d_references")
        if parsed.implicit_intersection or parsed.spill:
            unsupported_features.add("structured_or_dynamic_references")
        if parsed.unresolved_names and not parsed.functions.intersection({"LET", "LAMBDA"}):
            unsupported_features.add("unresolved_defined_names")
        if "DUMMYFUNCTION" in parsed.functions:
            # Google Sheets exports functions Excel lacks as
            # IFERROR(__xludf.DUMMYFUNCTION("..."), <cached value>).
            unsupported_features.add("google_sheets_placeholders")

        if parsed.functions & {"IFERROR", "IFNA"}:
            found = error_masks(formula)
            if found:
                masked.append((cell, found))

        # A cell the formula tests for blank anywhere (=""; ISBLANK; a
        # comparison; inside IFERROR) is handled wherever else it appears.
        tested = {
            ((ref.sheet or cell["sheet"]).casefold(), ref.ref) for ref in parsed.references if ref.guarded
        }
        # A row whose inputs are all blank (a day off on a timesheet, an
        # unused template row) is not a broken link; only a lone blank among
        # filled inputs is.
        row_inputs_blank = _row_inputs_all_blank(formula_wb, cell, parsed)
        for ref in parsed.references:
            ok, reason = reference_in_bounds(ref, cell["sheet"], formula_wb)
            if not ok:
                unresolved.setdefault(cell["location"], (cell, []))[1].append(reason or f"Could not resolve {ref.raw}.")
                continue
            if not blank_precedents:
                continue
            if ref.is_range or not ref.bounded or ref.positional or not ref.sensitive or row_inputs_blank:
                continue
            if ((ref.sheet or cell["sheet"]).casefold(), ref.ref) in tested:
                continue
            box = boundaries(ref.ref)
            if box is None:
                continue
            title = find_sheet(formula_wb, ref.sheet or cell["sheet"])
            if title is None:
                continue
            target = existing_cell(formula_wb[title], box[1], box[0])
            if target is None or target.value is None:
                if not anomalous_blank(title, box[1], box[0]):
                    continue
                findings.append(
                    Finding(
                        rule_id="BLANK_PRECEDENT",
                        severity="Medium",
                        error_confidence="Review",
                        detection_mode="DET",
                        location=cell["location"],
                        title="Formula references a blank precedent",
                        formula=formula,
                        evidence=[
                            f"Referenced cell {title}!{ref.ref} is blank although the cells around it in its column hold data, "
                            "and the formula uses it without testing for a blank."
                        ],
                        suggested_fix="Confirm the blank precedent is intentional or update the formula to the correct input.",
                    )
                )

    for members in group_by_pattern(deleted):
        lead = members[0]
        evidence = ["Formula text contains #REF!."]
        note = pattern_note(members, "contains #REF!")
        if note:
            evidence.append(note)
        findings.append(
            Finding(
                rule_id="BROKEN_REFERENCE",
                severity="High",
                error_confidence="Defect",
                detection_mode="DET",
                location=lead["location"],
                members=locations(members),
                title="Formula contains deleted reference",
                formula=lead["formula"],
                evidence=evidence,
                suggested_fix="Restore the deleted reference or rebuild the formula from intended source cells.",
            )
        )
    reasons_at = {loc: reasons for loc, (_cell, reasons) in unresolved.items()}
    for members in group_by_pattern([cell for cell, _reasons in unresolved.values()]):
        lead = members[0]
        evidence = list(dict.fromkeys(reasons_at[lead["location"]]))
        note = pattern_note(members, "has the same unresolved reference")
        if note:
            evidence.append(note)
        findings.append(
            Finding(
                rule_id="BROKEN_REFERENCE",
                severity="High",
                error_confidence="Likely defect",
                detection_mode="DET",
                location=lead["location"],
                members=locations(members),
                title="Formula reference cannot be resolved cleanly",
                formula=lead["formula"],
                evidence=evidence,
                suggested_fix="Review the reference target and restore the intended sheet/range.",
            )
        )

    for (sheet, source), cells in sorted(external.items()):
        shown = ", ".join(c["location"] for c in cells[:8]) + (", ..." if len(cells) > 8 else "")
        findings.append(
            Finding(
                rule_id="BROKEN_REFERENCE",
                severity="Low",
                error_confidence="Info",
                detection_mode="DET",
                location=cells[0]["location"],
                members=locations(cells),
                title="Formulas reference an external workbook",
                formula=cells[0]["formula"],
                evidence=[
                    f"{len(cells)} cell(s) on {sheet} link to {source}; external links are inventoried but not "
                    f"followed, so their cached values are taken as given: {shown}."
                ],
                suggested_fix="Provide the linked workbook or confirm the cached linked values are current.",
            )
        )

    masks_of = {id(cell): found for cell, found in masked}
    for members in group_by_pattern([cell for cell, _ in masked]):
        lead = members[0]
        mask = masks_of[id(lead)][0]
        if mask.sources == ("a Google Sheets placeholder",):
            evidence = [
                (
                    f"{mask.function} wraps a Google Sheets placeholder and returns its cached value, "
                    f"{mask.fallback}: the cell holds a frozen value, not a live calculation."
                )
            ]
        else:
            evidence = [
                (
                    f"{mask.function} replaces any error from {', '.join(mask.sources)} with {mask.fallback}, "
                    "which reads as a real value to anyone looking at the cell and to every formula that uses it."
                )
            ]
        note = pattern_note(members)
        if note:
            evidence.append(note)
        findings.append(
            Finding(
                rule_id="IFERROR_MASK",
                severity="Medium",
                error_confidence="Review",
                detection_mode="HEUR",
                location=lead["location"],
                members=locations(members),
                title="Error replaced by a value that reads as data",
                formula=lead["formula"],
                evidence=evidence,
                suggested_fix=(
                    'Return a visible marker for the failure ("" or NA()) or fix the input the error comes from; '
                    "keep the replacement only if it is the right result when the calculation fails."
                ),
            )
        )
    return findings


def _row_inputs_all_blank(formula_wb, cell: dict, parsed) -> bool:
    """True when the formula reads two or more single cells from its own row and all are blank.

    A timesheet day with neither an in nor an out time, or a template row
    with every input empty, is unused space; a lone blank input among filled
    ones is what BLANK_PRECEDENT is for.
    """
    from .reference_resolver import boundaries, existing_cell, find_sheet

    inputs = 0
    for ref in parsed.references:
        if ref.is_range or not ref.bounded or ref.positional:
            continue
        if ref.sheet is not None and ref.sheet.casefold() != cell["sheet"].casefold():
            continue
        box = boundaries(ref.ref)
        if box is None or box[1] != cell["row"] or box[0] == cell["col"]:
            continue
        title = find_sheet(formula_wb, cell["sheet"])
        if title is None:
            return False
        target = existing_cell(formula_wb[title], box[1], box[0])
        if target is not None and target.value is not None:
            return False
        inputs += 1
    return inputs >= 2


def _dedupe_range_length_with_drift(findings: list[Finding]) -> list[Finding]:
    drift_locations = {f.location for f in findings if f.rule_id == "FORMULA_DRIFT"}
    return [
        f
        for f in findings
        if not (f.rule_id == "RANGE_LENGTH_MISMATCH" and f.location in drift_locations)
    ]


def apply_impact_escalation(findings: list[Finding], config: dict) -> list[Finding]:
    """Escalate findings that feed configured headline outputs or exceed materiality.

    Activates scope.headline_outputs and materiality config and records impact
    flags on each affected finding.
    """
    from .materiality import escalate, exceeds_materiality

    scope = config.get("scope") or {}
    headline_outputs = list(scope.get("headline_outputs") or [])
    materiality = config.get("materiality") or {}
    absolute = float(materiality.get("absolute", 1000.0))
    relative = float(materiality.get("relative", 0.001))

    for finding in findings:
        should_escalate = False

        # Any cell of a repeated mistake can be the headline output.
        if any(location_matches(cell, target) for cell in finding.covered() for target in headline_outputs):
            finding.impact["feeds_headline_output"] = True
            should_escalate = True

        delta = finding.impact.get("estimated_delta") if isinstance(finding.impact, dict) else None
        if delta is not None and exceeds_materiality(delta, absolute=absolute, relative=relative):
            finding.impact["materiality_exceeded"] = True
            should_escalate = True

        if should_escalate:
            finding.severity = escalate(finding.severity)

    return findings


def detect_cycles(
    formulas: list[dict],
    expansion_limit: int | None = None,
    names=None,
    extents: dict[str, tuple[int, int]] | None = None,
    budget=None,
) -> list[Finding]:
    """Report one CIRCULAR_REFERENCE finding per cyclic dependency component.

    Cells that each depend only on themselves and share a relative formula
    (a self-referencing formula filled down a column) are one finding.

    ``expansion_limit`` is accepted for backward compatibility and ignored; see
    :func:`spreadsheet_auditor.dependency_graph.build_dependency_graph`.
    """
    from .dependency_graph import build_dependency_graph, find_cycles
    from .grouping import group_by_pattern, locations, pattern_note

    graph = build_dependency_graph(formulas, names=names, extents=extents, budget=budget)
    cell_at = {item["location"]: item for item in formulas}
    findings: list[Finding] = []
    self_references: list[dict] = []
    for members in find_cycles(graph):
        if len(members) == 1 and members[0] in cell_at:
            self_references.append(cell_at[members[0]])
            continue
        if len(members) == 1:
            title = "Formula references its own cell"
            evidence = [
                f"{members[0]} depends on itself, for example a total whose range includes the total cell."
            ]
        else:
            title = "Formula dependency cycle detected"
            shown = " -> ".join(members[:8]) + (" -> ..." if len(members) > 8 else "")
            evidence = [f"{len(members)} cells depend on each other in a cycle: {shown}"]
        findings.append(
            Finding(
                rule_id="CIRCULAR_REFERENCE",
                severity="High",
                error_confidence="Likely defect",
                detection_mode="DET",
                location=members[0],
                members=list(members),
                title=title,
                evidence=evidence,
                suggested_fix="Confirm whether iterative calculation is intentional; otherwise break the circular dependency.",
            )
        )
    self_references.sort(key=lambda cell: (cell["sheet"], cell["row"], cell["col"]))
    for group in group_by_pattern(self_references):
        lead = group[0]
        evidence = [f"{lead['location']} depends on itself, for example a total whose range includes the total cell."]
        note = pattern_note(group, "references its own cell")
        if note:
            evidence.append(note)
        findings.append(
            Finding(
                rule_id="CIRCULAR_REFERENCE",
                severity="High",
                error_confidence="Likely defect",
                detection_mode="DET",
                location=lead["location"],
                members=locations(group),
                title="Formula references its own cell",
                formula=lead["formula"],
                evidence=evidence,
                suggested_fix="Confirm whether iterative calculation is intentional; otherwise break the circular dependency.",
            )
        )
    return findings


def _infer_format(path: str | None, explicit: str | None) -> str:
    if explicit:
        return explicit
    if not path:
        return "markdown"
    lower = path.lower()
    if lower.endswith((".html", ".htm")):
        return "html"
    if lower.endswith(".json"):
        return "json"
    if lower.endswith(".sarif") or lower.endswith(".sarif.json"):
        return "sarif"
    return "markdown"


def _write_report(
    payload: dict,
    path: str,
    explicit_format: str | None,
    show_suppressed: bool = False,
) -> None:
    fmt = _infer_format(path, explicit_format)
    if fmt == "html":
        write_html(payload, path, show_suppressed=show_suppressed)
    elif fmt == "json":
        write_json(payload, path)
    elif fmt == "sarif":
        # SARIF is a strict, validated format. If rendering fails we must NOT
        # silently write generic findings JSON to a .sarif path; that would make
        # the run look successful while emitting an invalid SARIF file. Let the
        # exception propagate so the CLI exits with code 5.
        from .sarif import write_sarif

        write_sarif(payload, path)
    else:
        write_markdown(payload, path, show_suppressed=show_suppressed)


def exit_code(
    findings: list[dict], fail_on: str, limitations: list[str], strict: bool = False, incomplete: bool = False
) -> int:
    # An audit that skipped part of the workbook outranks its findings: the
    # findings it has are real, but a clean-looking partial run is exactly
    # what must not pass, so it fails whatever --fail-on says.
    if incomplete:
        return 6
    threshold = FAIL_ORDER.get(fail_on, 99)
    if fail_on != "None" and any(
        FAIL_ORDER.get(finding["severity"], 99) <= threshold and not finding.get("suppressed")
        for finding in findings
    ):
        return 1
    # Benign limitations (no recalc, missing optional packages) are expected on a
    # clean run and must not fail CI. Only surface them as exit code 2 when the
    # caller explicitly opts in via --strict or --fail-on None.
    if (strict or fail_on == "None") and limitations:
        return 2
    return 0


EPILOG = """\
Exit codes:
  0  No findings at or above --fail-on (default Critical), or --fail-on None.
  1  Findings at or above --fail-on were detected.
  2  --strict (or --fail-on None) and coverage limitations were present
     (e.g. recalculation unavailable, defusedxml missing).
  3  Healthcheck failed: required dependencies missing.
  4  Invalid input: bad command-line arguments or config, a workbook, CSV or
     suppression file that cannot be read, or an output path that cannot be
     written. stderr says what to fix.
  5  Internal auditor error, a bug: stderr has the traceback to attach to a
     bug report.
  6  Audit incomplete: a check failed or ran out of time, or a size limit
     (limits.max_formulas, limits.max_cells) skipped part of the workbook.
     The findings reported are real but partial; the limitations say what
     was not checked. Takes precedence over 1.

Examples:
  spreadsheet-auditor model.xlsx --out report.md --json findings.json
  spreadsheet-auditor model.xlsx --format html --out report.html
  spreadsheet-auditor model.xlsx --summary
  spreadsheet-auditor model.xlsx --summary --pin-suppressions
  spreadsheet-auditor --demo
  spreadsheet-auditor --healthcheck --json
"""


def _summary_lines(payload: dict, fail_on: str) -> list[str]:
    findings = [f for f in payload["findings"] if not f.get("suppressed")]
    coverage = payload.get("coverage", {})
    workbook = payload.get("workbook", {})
    # Counted before the report cap, so a capped report still says how many
    # findings each rule has.
    totals = coverage.get("finding_counts")
    if totals:
        counts = Counter(totals["by_severity"])
        by_rule = Counter(
            {(entry["severity"], entry["error_confidence"], entry["rule_id"]): entry["count"] for entry in totals["by_rule"]}
        )
        shown = "" if totals["shown"] == totals["total"] else f" ({totals['shown']} of {totals['total']} shown)"
    else:
        counts = Counter(f["severity"] for f in findings)
        by_rule = Counter((f["severity"], f["error_confidence"], f["rule_id"]) for f in findings)
        shown = ""
    if coverage.get("complete", True):
        status = "complete"
    else:
        status = "INCOMPLETE: part of the workbook was not checked; findings are partial (see limitations)"
    lines = [
        f"workbook   : {workbook.get('path')}",
        f"audit      : {status}",
        f"sheets     : {workbook.get('sheets_analyzed')}",
        f"formulas   : {workbook.get('formulas_scanned')}",
        f"recalc     : {workbook.get('recalc_status')}",
        "findings   : "
        + ", ".join(
            f"{counts.get(sev, 0)} {sev}" for sev in ("Critical", "High", "Medium", "Low", "Info")
        )
        + shown,
        f"fail_on    : {fail_on}",
    ]
    if by_rule:
        lines.append("by rule    :")
        for (severity, confidence, rule), count in sorted(
            by_rule.items(), key=lambda item: (FAIL_ORDER.get(item[0][0], 4), -item[1], item[0][2])
        ):
            lines.append(f"  {count:4d}  {rule} ({severity}, {confidence})")
    limitations = coverage.get("limitations") or []
    if limitations:
        lines.append("limitations:")
        for note in limitations:
            lines.append(f"  - {note}")
    return lines


def _configure_streams() -> None:
    """Never let a non-ASCII cell label crash report output on a legacy console encoding."""
    for stream in (sys.stdout, sys.stderr):
        reconfigure = getattr(stream, "reconfigure", None)
        if reconfigure is None:
            continue
        encoding = (getattr(stream, "encoding", "") or "").replace("-", "").lower()
        try:
            if encoding == "utf8":
                reconfigure(errors="replace")
            else:
                reconfigure(encoding="utf-8", errors="replace")
        except (ValueError, OSError):  # pragma: no cover - exotic streams
            pass


def _report_internal_error(stage: str) -> None:
    """Explain an exit 5 on stderr with the exception type and traceback, for a bug report."""
    error_type, error, _ = sys.exc_info()
    summary = "".join(traceback.format_exception_only(error_type, error)).strip()
    print(
        f"Internal audit error {stage}: {summary}\n"
        f"This is a bug in spreadsheet-auditor {__version__}, not a problem with the workbook. "
        "Please report it with the traceback below.",
        file=sys.stderr,
    )
    traceback.print_exc(file=sys.stderr)


class _ArgumentParser(argparse.ArgumentParser):
    """Exit 4 on a usage error instead of argparse's 2.

    Exit 2 means a completed audit with coverage limitations, so a mistyped
    flag in a CI script that accepts 2 would otherwise pass as an audit.
    """

    def error(self, message: str):  # type: ignore[override]
        self.print_usage(sys.stderr)
        self.exit(4, f"{self.prog}: error: {message}\n")


def main(argv: list[str] | None = None) -> int:
    _configure_streams()
    parser = _ArgumentParser(
        prog="spreadsheet-auditor",
        description="Audit an existing spreadsheet for correctness defects.",
        epilog=EPILOG,
        formatter_class=argparse.RawDescriptionHelpFormatter,
    )
    parser.add_argument("workbook", nargs="?", help="Path to .xlsx, .xlsm, or .csv workbook (omit when using --healthcheck or --demo)")
    parser.add_argument("--out", help="Report output path. Format inferred from extension or --format.")
    parser.add_argument(
        "--json",
        dest="json_out",
        nargs="?",
        const="-",
        default=None,
        help=(
            "Findings JSON output path. Pass '-' (or no value) to write JSON to stdout. "
            "With --healthcheck, --json (no value) emits the machine-readable healthcheck report."
        ),
    )
    parser.add_argument("--annotated", help="Optional annotated workbook output path. The source workbook is never modified.")
    parser.add_argument(
        "--format",
        choices=["markdown", "json", "html", "sarif"],
        default=None,
        dest="report_format",
        help="Report format. Inferred from --out extension when omitted.",
    )
    parser.add_argument("--config", help="Optional JSON, YAML, or YML config path")
    parser.add_argument(
        "--ignore",
        default=None,
        help="Suppression file path. Default '.audit-ignore' in the working folder, used when it exists.",
    )
    parser.add_argument(
        "--fail-on",
        default="Critical",
        choices=["Critical", "High", "Medium", "Low", "None"],
        help="Lowest severity that causes a non-zero exit code (default: Critical).",
    )
    parser.add_argument(
        "--strict",
        action="store_true",
        help="Return exit code 2 when coverage limitations are present (e.g. recalculation unavailable).",
    )
    parser.add_argument(
        "--show-suppressed",
        action="store_true",
        help="Render suppressed findings in the report instead of hiding them.",
    )
    parser.add_argument(
        "--summary",
        action="store_true",
        help="Print a one-screen summary (counts, coverage, limitations) instead of the full report.",
    )
    parser.add_argument(
        "--quiet",
        "-q",
        action="store_true",
        help="Suppress all stdout output. Useful in CI when only the exit code matters.",
    )
    parser.add_argument(
        "--demo",
        action="store_true",
        help="Audit the bundled demo workbook (examples/demo_bad_budget.xlsx). Useful for a 60-second tour.",
    )
    parser.add_argument(
        "--version",
        action="version",
        version=f"spreadsheet-auditor {__version__}",
    )
    parser.add_argument(
        "--pin-suppressions",
        action="store_true",
        help=(
            "After the audit, rewrite the --ignore file so each one-cell line carries the fingerprint of "
            "the finding it matches now, and each pinned line names where its finding is now. Other lines "
            "are kept as they are, and the workbook is never written. A line is pinned to whatever it "
            "matches in this run, so run it before rows move and review the changes printed on stderr."
        ),
    )
    parser.add_argument("--recalc-timeout", type=int, default=None, help="Override recalculation timeout in seconds.")
    parser.add_argument("--healthcheck", action="store_true", help="Report environment/runtime readiness and exit.")
    args = parser.parse_args(argv)

    if args.healthcheck:
        return healthcheck(as_json=args.json_out is not None)

    with contextlib.ExitStack() as stack:
        if args.demo:
            demo_path = stack.enter_context(_resolve_demo_workbook())
            if demo_path is None:
                print(
                    "Demo workbook not found. Generate it with `python examples/make_demo.py`.",
                    file=sys.stderr,
                )
                return 4
            args.workbook = str(demo_path)

        if not args.workbook:
            parser.error("workbook is required unless --healthcheck or --demo is used")

        # The default file is optional; a file named on the command line is not,
        # or a mistyped path would drop every suppression without a word.
        if args.ignore is None:
            args.ignore = ".audit-ignore"
        elif not Path(args.ignore).is_file():
            print(f"Preflight failed: suppression file not found: {args.ignore}", file=sys.stderr)
            return 4

        source = Path(args.workbook)
        targets = [
            (flag, Path(target))
            for flag, target in (("--out", args.out), ("--json", args.json_out), ("--annotated", args.annotated))
            if target and not (flag == "--json" and target == "-")
        ]
        for index, (flag, target) in enumerate(targets):
            for previous_flag, previous in targets[:index]:
                if target.resolve() == previous.resolve() or (
                    target.exists() and previous.exists() and target.samefile(previous)
                ):
                    print(
                        f"Preflight failed: {flag} and {previous_flag} refer to the same output file; use distinct paths.",
                        file=sys.stderr,
                    )
                    return 4
        for flag, target in (("--out", args.out), ("--json", args.json_out), ("--annotated", args.annotated)):
            if target and target != "-" and source.exists() and Path(target).exists() and source.samefile(target):
                print(f"Preflight failed: {flag} {target} is the workbook being audited; write to a new file.", file=sys.stderr)
                return 4

        try:
            payload, code = audit_workbook(args)
        except PreflightError as exc:
            print(f"Preflight failed: {exc}", file=sys.stderr)
            return 4
        except ConfigError as exc:
            print(f"Config error: {exc}", file=sys.stderr)
            return 4
        except Exception:
            _report_internal_error("while auditing")
            return 5

        try:
            if args.json_out:
                if args.json_out == "-":
                    if not args.quiet:
                        print(json.dumps(payload, indent=2))
                else:
                    write_json(payload, args.json_out)
            if args.out:
                _write_report(payload, args.out, args.report_format, show_suppressed=args.show_suppressed)
            if args.annotated:
                from .annotate import annotate_workbook

                annotate_workbook(args.workbook, args.annotated, payload["findings"])
                for note in payload.get("coverage", {}).get("limitations", []):
                    if note.startswith("The annotated copy"):
                        print(f"Warning: {note}", file=sys.stderr)
        except OSError as exc:
            # A missing folder, a file open in Excel, no permission: fixable by the user.
            print(f"Could not write output: {exc}", file=sys.stderr)
            return 4
        except Exception:
            _report_internal_error("while writing output")
            return 5

        # Even with --quiet: a CI log that shows only the exit code should
        # still say why the audit is partial.
        for entry in payload.get("coverage", {}).get("incomplete", []):
            print(f"Audit incomplete: {entry['message']}", file=sys.stderr)

        if args.quiet:
            return code

        if args.summary:
            for line in _summary_lines(payload, args.fail_on):
                print(line)
            return code

        if not args.out and not args.json_out:
            if args.report_format == "html":
                print(render_html(payload, show_suppressed=args.show_suppressed))
            elif args.report_format == "json":
                print(json.dumps(payload, indent=2))
            else:
                print(render_markdown(payload, show_suppressed=args.show_suppressed))
        elif not args.out and args.json_out == "-":
            # already printed JSON above; nothing else to do
            pass
        else:
            counts = Counter(f["severity"] for f in payload["findings"] if not f.get("suppressed"))
            done = "complete" if payload.get("coverage", {}).get("complete", True) else "incomplete"
            print(
                f"Audit {done}: {counts.get('Critical', 0)} Critical, "
                f"{counts.get('High', 0)} High, {counts.get('Medium', 0)} Medium, "
                f"{counts.get('Low', 0)} Low."
            )
        return code


@contextlib.contextmanager
def _resolve_demo_workbook():
    """Yield a filesystem path to the bundled demo workbook, or ``None``.

    Resolution order:

    1. A repository checkout (``examples/demo_bad_budget.xlsx`` next to the
       package or under the current working directory) so development uses the
       canonical copy.
    2. The packaged resource ``spreadsheet_auditor/demo/demo_bad_budget.xlsx``
       so installed wheels and unpacked skill packages work too.

    Yielded as a context manager because the packaged resource may live inside a
    zip and must be materialized via :func:`importlib.resources.as_file`; the
    real path stays valid for the lifetime of the ``with`` block. Callers (recalc,
    annotation) re-open the path, so the context must wrap the whole audit.
    """
    for candidate in (
        Path(__file__).resolve().parents[1] / "examples" / "demo_bad_budget.xlsx",
        Path.cwd() / "examples" / "demo_bad_budget.xlsx",
    ):
        if candidate.exists():
            yield candidate
            return
    try:
        resource = importlib.resources.files("spreadsheet_auditor.demo") / "demo_bad_budget.xlsx"
        if resource.is_file():
            with importlib.resources.as_file(resource) as path:
                yield path
                return
    except (ModuleNotFoundError, FileNotFoundError):
        pass
    yield None


if __name__ == "__main__":
    raise SystemExit(main())
