"""Defined names and structured table references resolved to A1 ranges.

Built once per audit from the formula workbook and handed to the parser via
``CheckContext.names``. Names whose definition is a constant or a formula
(``=OFFSET(...)``) cannot be reduced to a range; they are recorded in
``unresolvable`` and formulas that use them surface as coverage limitations
rather than as phantom cell references.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field

from openpyxl.utils.cell import get_column_letter, range_boundaries

_SPECIFIERS = {"#ALL", "#DATA", "#HEADERS", "#TOTALS", "#THIS ROW"}
_STRUCTURED_RE = re.compile(r"^(?P<table>[A-Za-z_\\][A-Za-z0-9_.]*)?\[(?P<body>.*)\]$", re.S)
_ESCAPES = (("'[", "["), ("']", "]"), ("'#", "#"), ("'@", "@"), ("''", "'"))


@dataclass
class TableInfo:
    name: str
    sheet: str
    min_col: int
    min_row: int
    max_col: int
    max_row: int
    header_rows: int = 1
    totals_rows: int = 0
    columns: list[str] = field(default_factory=list)

    @property
    def data_rows(self) -> tuple[int, int] | None:
        first = self.min_row + self.header_rows
        last = self.max_row - self.totals_rows
        return (first, last) if first <= last else None

    def contains(self, sheet: str, row: int, col: int) -> bool:
        return (
            sheet.casefold() == self.sheet.casefold()
            and self.min_row <= row <= self.max_row
            and self.min_col <= col <= self.max_col
        )

    def column_index(self, name: str) -> int | None:
        wanted = name.strip().casefold()
        for offset, column in enumerate(self.columns):
            if column.strip().casefold() == wanted:
                return self.min_col + offset
        return None


def _unescape(item: str) -> str:
    text = item.strip()
    for escaped, plain in _ESCAPES:
        text = text.replace(escaped, plain)
    return text


def _split_items(body: str) -> list[tuple[str, str]]:
    """Split a structured-reference body into ``(separator, item)`` pairs."""
    items: list[tuple[str, str]] = []
    i = 0
    sep = ""
    while i < len(body):
        ch = body[i]
        if ch in " \t":
            i += 1
            continue
        if ch in ",:":
            sep = ch
            i += 1
            continue
        if ch == "[":
            depth = 1
            j = i + 1
            while j < len(body) and depth:
                if body[j] == "[":
                    depth += 1
                elif body[j] == "]":
                    depth -= 1
                j += 1
            items.append((sep, body[i + 1 : j - 1]))
            i = j
            sep = ""
            continue
        j = i
        while j < len(body) and body[j] not in ",:[":
            j += 1
        items.append((sep, body[i:j].strip()))
        i = j
        sep = ""
    return items


def _iter_defined_names(container):
    if container is None:
        return []
    values = getattr(container, "values", None)
    if callable(values):
        try:
            return list(values())
        except Exception:  # pragma: no cover - defensive
            return []
    defined = getattr(container, "definedName", None)
    if defined is not None:
        return list(defined)
    try:
        return list(container)
    except TypeError:
        return []


def _iter_tables(ws):
    tables = getattr(ws, "tables", None)
    if not tables:
        return []
    values = getattr(tables, "values", None)
    if callable(values):
        try:
            return [table for table in values() if hasattr(table, "ref")]
        except Exception:  # pragma: no cover - defensive
            return []
    return []


def _table_info(ws, table) -> TableInfo | None:
    ref = getattr(table, "ref", None)
    if not ref:
        return None
    try:
        min_col, min_row, max_col, max_row = range_boundaries(str(ref))
    except (ValueError, TypeError):
        return None
    if None in (min_col, min_row, max_col, max_row):
        return None
    header_rows = getattr(table, "headerRowCount", 1)
    header_rows = 1 if header_rows is None else int(header_rows)
    totals_rows = int(getattr(table, "totalsRowCount", 0) or 0)
    columns = [
        str(getattr(column, "name", "") or "")
        for column in (getattr(table, "tableColumns", None) or [])
    ]
    if not any(columns) and header_rows > 0:
        cells = getattr(ws, "_cells", {})
        columns = []
        for col in range(min_col, max_col + 1):
            cell = cells.get((min_row, col))
            value = getattr(cell, "value", None)
            columns.append("" if value is None else str(value))
    name = getattr(table, "displayName", None) or getattr(table, "name", None) or ""
    if not name:
        return None
    return TableInfo(
        name=str(name),
        sheet=ws.title,
        min_col=min_col,
        min_row=min_row,
        max_col=max_col,
        max_row=max_row,
        header_rows=header_rows,
        totals_rows=totals_rows,
        columns=columns,
    )


class NameTable:
    """Lookup for defined names (workbook- and sheet-scoped) and tables."""

    def __init__(self) -> None:
        self.global_names: dict[str, list[tuple[str, str]]] = {}
        self.sheet_names: dict[tuple[str, str], list[tuple[str, str]]] = {}
        self.unresolvable: set[str] = set()
        self.tables: dict[str, TableInfo] = {}

    @classmethod
    def from_workbook(cls, workbook) -> "NameTable":
        table = cls()
        if workbook is None:
            return table
        sheet_titles = list(getattr(workbook, "sheetnames", []) or [])
        for defined in _iter_defined_names(getattr(workbook, "defined_names", None)):
            table._add_defined_name(defined, sheet_titles)
        for ws in getattr(workbook, "worksheets", []) or []:
            for defined in _iter_defined_names(getattr(ws, "defined_names", None)):
                table._add_defined_name(defined, sheet_titles, scope=ws.title)
            for raw_table in _iter_tables(ws):
                info = _table_info(ws, raw_table)
                if info is not None:
                    table.tables[info.name.upper()] = info
        return table

    def _add_defined_name(self, defined, sheet_titles: list[str], scope: str | None = None) -> None:
        name = getattr(defined, "name", None)
        if not name or str(name).startswith("_xlnm"):
            return
        key = str(name).upper()
        if scope is None:
            local = getattr(defined, "localSheetId", None)
            if local is not None:
                try:
                    scope = sheet_titles[int(local)]
                except (IndexError, ValueError, TypeError):
                    scope = None
        try:
            destinations = list(getattr(defined, "destinations", []) or [])
        except Exception:
            destinations = []
        targets = [
            (str(sheet), str(cells).replace("$", "").upper())
            for sheet, cells in destinations
            if sheet and cells
        ]
        if not targets:
            self.unresolvable.add(key)
            return
        if scope is None:
            self.global_names[key] = targets
        else:
            self.sheet_names[(scope.casefold(), key)] = targets

    def is_empty(self) -> bool:
        return not (self.global_names or self.sheet_names or self.tables)

    def resolve(self, name: str, sheet: str | None = None) -> list[tuple[str, str]]:
        key = name.strip().upper()
        if sheet:
            scoped = self.sheet_names.get((sheet.casefold(), key))
            if scoped:
                return list(scoped)
        return list(self.global_names.get(key, []))

    def _table_containing(self, origin: tuple[str, int, int] | None) -> TableInfo | None:
        if origin is None:
            return None
        sheet, row, col = origin
        for info in self.tables.values():
            if info.contains(sheet, row, col):
                return info
        return None

    def resolve_structured(
        self, text: str, origin: tuple[str, int, int] | None = None
    ) -> list[tuple[str, str]]:
        match = _STRUCTURED_RE.match(text.strip())
        if not match:
            return []
        table_name, body = match.group("table"), match.group("body")
        info = self.tables.get(table_name.upper()) if table_name else self._table_containing(origin)
        if info is None:
            return []

        this_row = False
        specs: set[str] = set()
        columns: list[str] = []
        span = False
        for sep, item in _split_items(body):
            item = item.strip()
            if item.startswith("@"):
                this_row = True
                item = item[1:].strip()
                if item.startswith("[") and item.endswith("]"):
                    item = item[1:-1]
            if not item:
                continue
            upper = item.upper()
            if upper in _SPECIFIERS:
                specs.add(upper)
                continue
            if sep == ":" and columns:
                span = True
            columns.append(_unescape(item))
        if "#THIS ROW" in specs:
            this_row = True

        data = info.data_rows
        if this_row:
            if origin is None or origin[0].casefold() != info.sheet.casefold() or data is None:
                return []
            row = origin[1]
            if not data[0] <= row <= data[1]:
                return []
            rows = (row, row)
        elif "#ALL" in specs:
            rows = (info.min_row, info.max_row)
        elif "#HEADERS" in specs and "#DATA" in specs:
            rows = (info.min_row, info.max_row - info.totals_rows)
        elif "#HEADERS" in specs:
            if info.header_rows == 0:
                return []
            rows = (info.min_row, info.min_row + info.header_rows - 1)
        elif "#TOTALS" in specs and "#DATA" in specs:
            rows = (info.min_row + info.header_rows, info.max_row)
        elif "#TOTALS" in specs:
            if info.totals_rows == 0:
                return []
            rows = (info.max_row - info.totals_rows + 1, info.max_row)
        else:
            if data is None:
                return []
            rows = data

        if not columns:
            cols = (info.min_col, info.max_col)
        else:
            indexes = [info.column_index(column) for column in columns]
            if any(index is None for index in indexes):
                return []
            resolved = [index for index in indexes if index is not None]
            if len(resolved) == 1:
                cols = (resolved[0], resolved[0])
            elif len(resolved) == 2 and span:
                cols = (min(resolved), max(resolved))
            else:
                return []

        first = f"{get_column_letter(cols[0])}{rows[0]}"
        last = f"{get_column_letter(cols[1])}{rows[1]}"
        return [(info.sheet, first if first == last else f"{first}:{last}")]
